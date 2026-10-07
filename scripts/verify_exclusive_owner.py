"""One real API reconstruction; old result is evidence, never an input revision."""
from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

import httpx
import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
from app.services.reconstruction.exclusive_ownership import (
    audit_raster_text,
)
from app.utils import image_io


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--old-project", required=True)
    parser.add_argument("--url", default="http://127.0.0.1:8000")
    args = parser.parse_args()
    old_root = ROOT / "outputs" / args.old_project
    source = old_root / "original.png"
    old_layout = json.loads((old_root / "slides/page_1.json").read_text(encoding="utf-8"))
    report_root = ROOT / "outputs" / "exclusive-ownership-regression"
    report_root.mkdir(exist_ok=True)
    audit_before = audit_raster_text(source, old_layout, old_root)
    (report_root / "audit_before.json").write_text(json.dumps(audit_before, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"beforeDuplicateTextCount": audit_before["rasterNativeDuplicateTextCount"], "beforeMissingAssets": audit_before["missingAssetCount"]}), flush=True)
    with httpx.Client(base_url=args.url, timeout=900, trust_env=False) as client:
        health = client.get("/api/health")
        health.raise_for_status()
        project = client.post("/api/projects", json={"name": "互斥Ownership单页回归（不覆盖原结果）"})
        project.raise_for_status()
        pid = project.json()["id"]
        upload = client.post(f"/api/projects/{pid}/images", files={"files": ("失败页面.png", source.read_bytes(), "image/png")})
        upload.raise_for_status()
        print(json.dumps({"projectId": pid, "phase": "real_qwen_ocr_sam_analysis_single_pass"}), flush=True)
        start = time.monotonic()
        response = client.post(f"/api/projects/{pid}/analyze?mode=maximum&allow_fallback=false")
        (report_root / "api_response.json").write_text(response.text, encoding="utf-8")
        print(json.dumps({"analysisStatus": response.status_code, "seconds": round(time.monotonic()-start, 2)}), flush=True)
        response.raise_for_status()
        root = ROOT / "outputs" / pid
        layout = json.loads((root / "slides/page_1.json").read_text(encoding="utf-8"))
        audit_after = audit_raster_text(source, layout, root)
        graph = layout["metadata"]["ownershipAudit"]
        qa = json.loads((root / "visual_score.json").read_text(encoding="utf-8"))
        # Re-read every per-object mask; no duplicate or unowned source pixels.
        original = image_io.imread(source)
        counts = np.zeros(original.shape[:2], np.uint16)
        for node in graph["nodes"]:
            mask = image_io.imread(root / node["mask"], -1)
            x1,y1,x2,y2 = node["sourceBBox"]
            counts[y1:y2,x1:x2] += mask[:, :, 3] > 0
        masks_ok = int((counts == 0).sum()) == 0 and int((counts > 1).sum()) == 0
        assert masks_ok
        assert audit_after["missingAssetCount"] == audit_after["rasterNativeDuplicateTextCount"] == 0
        paused = client.post(f"/api/projects/{pid}/pages/1/revise")
        assert paused.status_code == 409 and paused.json()["detail"]["code"] == "AUTOMATIC_REVISIONS_PAUSED"
        deck = Presentation(root / "editable.pptx")
        ppt_counts = {"slides": len(deck.slides), "textboxes": sum(s.has_text_frame for slide in deck.slides for s in slide.shapes),
                      "pictures": sum(s.shape_type == 13 for slide in deck.slides for s in slide.shapes)}
        panels = [source, old_root / "reconstructed_preview.png", root / "reconstructed_preview.png"]
        width = 1000
        height = round(original.shape[0]/original.shape[1]*width)
        canvas = Image.new("RGB", (width*3, height+45), "#FFFFFF")
        draw = ImageDraw.Draw(canvas)
        for index, (path, title) in enumerate(zip(panels, ("ORIGINAL", "OLD: duplicated text / fragmented visuals", "NEW: exclusive owners, revisions stopped"))):
            with Image.open(path) as image:
                canvas.paste(image.convert("RGB").resize((width,height)), (index*width,45))
            draw.text((index*width+12,12), title, fill="#17365D")
        canvas.save(report_root / "single_page_compare.png")
        summary = {"projectId": pid, "analysisHttpStatus": response.status_code, "apiUsed": response.json().get("aiUsed"),
            "visionProvider": response.json().get("visionProvider"), "ocrProvider": response.json().get("ocrProvider"),
            "segmentationProvider": layout["metadata"].get("segmentationProvider"), "auditBefore": audit_before, "auditAfter": audit_after,
            "unownedSourcePixels": int((counts == 0).sum()), "multiplyOwnedSourcePixels": int((counts > 1).sum()),
            "qa": qa, "ppt": ppt_counts, "revisionStopHttpStatus": paused.status_code, "automaticRevisionRounds": 0,
            "comparison": str(report_root / "single_page_compare.png"), "layout": str(root / "slides/page_1.json"),
            "pptx": str(root / "editable.pptx"), "originalProjectUntouched": True}
        (report_root / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({key: summary[key] for key in ("projectId", "apiUsed", "ocrProvider", "segmentationProvider", "unownedSourcePixels", "multiplyOwnedSourcePixels", "ppt")}), flush=True)


if __name__ == "__main__":
    main()
