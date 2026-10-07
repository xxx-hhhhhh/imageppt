"""Real OCR/SAM2, four supplied images and two measured local revisions.

Usage: python scripts/verify_owner_integration.py SIMPLE PARTY SCIENCE CHARTS
No mocks, no page-specific reconstruction rules, no paid VLM calls. New test
projects are created under the existing outputs store; originals stay read-only.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

from app.config import OUTPUTS_DIR
from app.models.project_store import ProjectStore
from app.services.reconstruction.pipeline import ReconstructionPipeline
from app.services.reconstruction.revision import revise_problem_regions
from app.services.scene.scene_analyzer import SceneAnalyzer
from app.services.segmentation.segmentation_provider import create_segmentation_provider
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("images", nargs=4, type=Path)
    args = parser.parse_args()
    store = ProjectStore()
    pipeline = ReconstructionPipeline(store)
    pipeline.scene_analyzer = SceneAnalyzer("auto", "none")
    pipeline.segmentation_provider, pipeline.segmentation_warnings = create_segmentation_provider("sam2")
    results = []
    for label, image in zip(("simple", "party", "science", "charts"), args.images):
        assert image.is_file(), image
        record = store.create(f"Object First integration regression: {label}")
        identifier = record["id"]
        store.add_image(identifier, {"id": label, "path": str(image.resolve()), "filename": image.name})
        layouts, provider, _ = pipeline.analyze_project(identifier, mode="standard", allow_fallback=True)
        assert layouts
        root = OUTPUTS_DIR / identifier
        qa = json.loads((root / "visual_score.json").read_text(encoding="utf-8"))
        assert qa["ownerGate"]["unownedPixelCount"] == 0
        assert qa["ownerGate"]["missingAssetCount"] == 0
        deck = Presentation(str(root / "editable.pptx"))
        shapes = list(deck.slides[0].shapes)
        revisions = []
        for _ in range(2):
            before = hashlib.sha256((root / "reconstructed_preview.png").read_bytes()).hexdigest()
            revision = revise_problem_regions(store, identifier, 1)
            after = hashlib.sha256((root / "reconstructed_preview.png").read_bytes()).hexdigest()
            if not revision["accepted"]:
                assert before == after, "Rejected revision changed published preview"
            assert revision["missingAssetCount"] == 0
            revisions.append({key: revision.get(key) for key in (
                "revisionRound", "accepted", "rollbackTriggered", "stagnationReason", "integrityErrors",
                "visualBefore", "visualAfter", "assetsBefore", "assetsAfter", "unownedPixelsBefore", "unownedPixelsAfter")})
        result = {"case": label, "projectId": identifier, "ocrProvider": provider,
                  "segmentationProvider": pipeline.segmentation_provider.name,
                  "nativeTextCount": sum(s.has_text_frame and bool(s.text.strip()) for s in shapes),
                  "nativeShapeCount": sum(s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE for s in shapes),
                  "imageObjectCount": sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in shapes),
                  "ownerGate": qa["ownerGate"], "objectizationAudit": qa.get("objectizationAudit"),
                  "editableTextCoverage": qa.get("editableTextCoverage"), "movableVisualCoverage": qa.get("movableVisualCoverage"),
                  "ghostingCount": qa.get("ghostingCount"), "revisions": revisions}
        results.append(result)
        (OUTPUTS_DIR / "object_first_integration_regression.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({key: result[key] for key in ("case", "projectId", "ocrProvider", "segmentationProvider", "nativeTextCount", "nativeShapeCount", "imageObjectCount")}), flush=True)


if __name__ == "__main__":
    main()
