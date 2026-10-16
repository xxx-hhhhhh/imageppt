"""Verify browser-edited Layout JSON is exported without geometry/text drift."""
import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'backend'))
from app.services.pptx.renderer import PPTXRenderer


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('layout', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    layout = json.loads(args.layout.read_text('utf-8'))
    path, _ = PPTXRenderer().render_project('geometry-regression', [layout], output_dir=args.output)
    deck = Presentation(path)
    shapes = list(deck.slides[0].shapes)
    elements = sorted((e for e in layout['elements'] if e['type'] != 'group' and not any(
        e.get('metadata', {}).get(k) for k in ('suppressed', 'suppressRender', 'ownedBy'))), key=lambda e: e['zIndex'])
    assert len(elements) == len(shapes), 'An object disappeared at export'
    sx, sy = deck.slide_width/layout['slide']['width'], deck.slide_height/layout['slide']['height']
    for element, shape in zip(elements, shapes):
        for key, actual, scale in [('x', shape.left, sx), ('y', shape.top, sy),
                                   ('width', shape.width, sx), ('height', shape.height, sy)]:
            assert abs(actual/scale-element[key]) < .01, (element['id'], key)
        if element['type'] == 'text':
            assert shape.has_text_frame and shape.text == element['text']
            assert abs(shape.text_frame.paragraphs[0].runs[0].font.size.pt-
                       element['style']['fontSize']*sx/914400*72) < .02
    (args.output/'slides').mkdir(exist_ok=True)
    (args.output/'slides/page_1.json').write_text(json.dumps(layout, ensure_ascii=False), encoding='utf-8')
    print(json.dumps({'objects': len(shapes), 'nativeTexts': sum(e['type']=='text' for e in elements),
                      'allPositionsSizesAndTextMatch': True, 'pptx': str(path)}))


if __name__ == '__main__':
    main()
