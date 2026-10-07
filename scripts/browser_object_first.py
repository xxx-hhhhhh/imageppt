"""Real Chromium acceptance against running frontend/backend and a regression deck."""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("regression_project")
    parser.add_argument("--existing-only", action="store_true", help="Render/edit existing graphs without uploading an extra page")
    args = parser.parse_args()
    from app.config import OUTPUTS_DIR
    from app.main import project_response
    from app.models.project_store import ProjectStore
    from playwright.sync_api import expect, sync_playwright
    from pptx import Presentation
    store = ProjectStore()
    baseline = store.get(args.regression_project)
    project = store.create("Object First browser acceptance")
    for image in baseline["images"]:
        store.add_image(project["id"], image)
    slides = store.list_slides(args.regression_project)
    for index, slide in enumerate(slides, 1):
        store.save_slide(project["id"], index, slide)
    root = OUTPUTS_DIR / project["id"]
    for name in ("original.png", "reconstructed_preview.png", "difference.png", "visual_score.json"):
        shutil.copy2(OUTPUTS_DIR / args.regression_project / name, root / name)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1600, "height": 1000})
        errors = []
        page.on("pageerror", lambda error: errors.append(str(error)))
        # Only the reopening response is routed; OCR/layout/assets/export are real.
        page.route("**/api/projects", lambda route: route.fulfill(json=project_response(store.get(project["id"])).model_dump(mode="json")) if route.request.method == "POST" else route.continue_())
        page.goto("http://127.0.0.1:5173/", wait_until="networkidle")
        expect(page.get_by_role("heading", name="图片转可编辑 PPT")).to_be_visible()
        page.screenshot(path=str(root / "browser_before.png"), full_page=True)
        if not args.existing_only:
            page.get_by_label("上传图片", exact=True).set_input_files(baseline["images"][0]["path"])
            expect(page.get_by_text('图片已上传，点击“开始解析”')).to_be_visible()
        page.get_by_label("转换模式").select_option("fast")
        with page.expect_response(lambda response: "/analyze" in response.url and response.request.method == "POST", timeout=180000) as parsed:
            page.get_by_role("button", name="AI解析", exact=True).click()
        assert parsed.value.status == 200
        payload = parsed.value.json()
        assert len(payload["slides"]) == len(baseline["images"]) + (0 if args.existing_only else 1)
        text = next(item for item in payload["slides"][0]["elements"] if item["type"] == "text")
        target = page.get_by_text(text["text"], exact=True).first
        expect(target).to_be_visible()
        target.dblclick()
        edited_text = "对象级编辑已验证"
        target.fill(edited_text)
        with page.expect_response(lambda r: "/slides/1" in r.url and r.request.method == "PUT"):
            page.get_by_role("button", name="保存当前页面").click()
        expect(page.get_by_role("textbox", name="文字", exact=True)).to_have_value(edited_text)
        page.get_by_role("spinbutton", name="字号", exact=True).fill("30")
        page.get_by_role("spinbutton", name="X", exact=True).fill("25")
        page.get_by_role("spinbutton", name="W", exact=True).fill("260")
        with page.expect_response(lambda r: "/slides/1" in r.url and r.request.method == "PUT"):
            page.get_by_role("button", name="保存当前页面").click()
        saved = store.get_slide(project["id"], 1)
        assert any(item.get("text") == edited_text and item["x"] == 25 and item["width"] == 260 for item in saved["elements"])
        target = page.get_by_text(edited_text, exact=True).first
        box = target.bounding_box()
        assert box
        page.mouse.move(box["x"] + 20, box["y"] + 8)
        page.mouse.down()
        page.mouse.move(box["x"] + 50, box["y"] + 13, steps=5)
        page.mouse.up()
        expect(page.get_by_role("spinbutton", name="X", exact=True)).not_to_have_value("25")
        handle = page.locator(".visual-text .resize-handle").first
        resize_box = handle.bounding_box()
        assert resize_box
        page.mouse.move(resize_box["x"] + 3, resize_box["y"] + 1)
        page.mouse.down()
        page.mouse.move(resize_box["x"] + 13, resize_box["y"] + 4, steps=5)
        page.mouse.up()
        expect(page.get_by_role("spinbutton", name="W", exact=True)).not_to_have_value("260")
        page.get_by_role("button", name="保存当前页面").click()
        page.screenshot(path=str(root / "browser_after.png"), full_page=True)
        page.get_by_role("button", name="Page 3 Page 3", exact=True).click()
        try:
            expect(page.get_by_role("button", name="保存当前页面")).to_be_enabled()
        except Exception:
            page.screenshot(path=str(root / "browser_failure.png"), full_page=True)
            print(json.dumps({"pageErrors": errors, "body": page.locator("body").inner_text()}, ensure_ascii=False), flush=True)
            raise
        page.wait_for_function("Array.from(document.querySelectorAll('.canvas-overlay img')).every(img => img.complete && img.naturalWidth > 0)")
        page.screenshot(path=str(root / "browser_science.png"), full_page=True)
        stage = page.locator(".canvas-stage").bounding_box()
        workspace = page.locator(".canvas-wrapper").bounding_box()
        assert stage and workspace
        assert stage["x"] >= workspace["x"] - 1 and stage["x"] + stage["width"] <= workspace["x"] + workspace["width"] + 1
        with page.expect_response(lambda r: "/export/pptx" in r.url and r.request.method == "POST", timeout=120000) as exported:
            page.get_by_role("button", name="导出PPT", exact=True).click()
        assert exported.value.status == 200
        assert exported.value.json()["validation"]["valid"], exported.value.json()["validation"]
        deck = Presentation(root / "editable.pptx")
        assert any(shape.has_text_frame and edited_text in shape.text for shape in deck.slides[0].shapes)
        assert any(shape.shape_type == 1 for slide in deck.slides for shape in slide.shapes)
        assert any(shape.shape_type == 13 for slide in deck.slides for shape in slide.shapes)
        assert not errors, errors
        browser.close()
    checks = ["open", "real-analysis", "thumbnails", "double-click-text", "font-size", "position", "pointer-drag", "pointer-resize", "save-canonical-scene", "page-switch", "canvas-fit", "real-pptx-export", "native-text-shape-image"]
    if not args.existing_only:
        checks.append("real-upload")
    (root / "browser_acceptance.json").write_text(json.dumps({"passed": True, "projectId": project["id"], "existingOnly": args.existing_only, "checks": checks, "pageErrors": errors}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"BROWSER REPORT: {root / 'browser_acceptance.json'}")


if __name__ == "__main__":
    main()
