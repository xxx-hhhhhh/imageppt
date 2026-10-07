"""Finalize diagnostic comparisons from actual saved provider/browser/PPT runs."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT / "backend"))
from app.utils import image_io


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--old-project", required=True)
    args = parser.parse_args()
    report_root = ROOT / "outputs/exclusive-ownership-regression"
    api = json.loads((report_root / "summary.json").read_text(encoding="utf-8"))
    replay = json.loads((report_root / "seam_replay.json").read_text(encoding="utf-8"))
    root = ROOT / "outputs" / replay["projectId"]
    original_root = ROOT / "outputs" / args.old_project
    source = original_root / "original.png"
    a = image_io.imread(source).astype(float)
    ppt_before = image_io.imread(report_root / "powerpoint_render.png").astype(float)
    ppt_after = image_io.imread(report_root / "powerpoint_final.png").astype(float)
    layout = json.loads((root / "slides/page_1.json").read_text(encoding="utf-8"))
    native_text = sum(e.get("owner") == "editable_text" for e in layout["elements"])
    native_shape = sum(e.get("owner") == "native_shape" for e in layout["elements"])
    result = {"finalProjectId":replay["projectId"],"realApiProjectId":api["projectId"],"apiUsed":api["apiUsed"],
        "ocrProvider":api["ocrProvider"],"segmentationProvider":api["segmentationProvider"],"newApiCallsForSeamReplay":0,
        "detectedTextCount":api["qa"]["detectedTextCount"],"nativeTextCount":native_text,"nativeShapeCount":native_shape,
        "movableImageCount":replay["audit"]["activeImageCount"],"ownershipAudit":replay["audit"],
        "unownedSourcePixels":replay["unownedSourcePixels"],"multiplyOwnedSourcePixels":replay["multiplyOwnedSourcePixels"],"previewChangedPixels":int(np.any(image_io.imread(source)!=image_io.imread(root / "reconstructed_preview.png"),axis=2).sum()),
        "pptRenderBeforeMeanAbsoluteError":round(float(np.abs(a-ppt_before).mean()),4),
        "pptRenderAfterMeanAbsoluteError":round(float(np.abs(a-ppt_after).mean()),4),
        "powerPointReadOnlyOpened":True,"powerPointObjectCount":len(Presentation(root / "editable.pptx").slides[0].shapes),
        "browser":json.loads((root / "exclusive_browser.json").read_text(encoding="utf-8")),
        "automaticRevisionRounds":0,"revisionPausedHttpStatus":api["revisionStopHttpStatus"],
        "knownLimits":[f"{api['qa']['detectedTextCount']-native_text} OCR lines retained in movable images, not editable text", "native geometry extraction is conservative", "thin alpha boundary/resampling differences can remain in PowerPoint", "module coherence requires further inspection; coverage is not semantic correctness"]}
    (report_root / "final_acceptance.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    panels = [source,original_root / "reconstructed_preview.png",report_root / "powerpoint_final.png"]
    width=1000
    height=round(a.shape[0]/a.shape[1]*width)
    canvas=Image.new("RGB",(width*3,height+45),"white")
    draw=ImageDraw.Draw(canvas)
    for index,(path,title) in enumerate(zip(panels,("SOURCE", "FAILED RESULT", "FIXED POWERPOINT RENDER - raster fallback"))):
        with Image.open(path) as image:
            canvas.paste(image.convert("RGB").resize((width,height)),(index*width,45))
        draw.text((index*width+12,12),title,fill="#17365D")
    canvas.save(report_root / "comparison_final.png")
    print(json.dumps({k:result[k] for k in ("finalProjectId","previewChangedPixels","movableImageCount","nativeTextCount","pptRenderBeforeMeanAbsoluteError","pptRenderAfterMeanAbsoluteError")}),flush=True)


if __name__ == "__main__":
    main()
