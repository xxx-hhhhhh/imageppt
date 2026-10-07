"""Run real OCR/segmentation/PPT and two revisions on caller-supplied images.

No fixture text is injected. Original projects and files are read-only.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("images", nargs="+", type=Path)
    parser.add_argument("--mode", choices=["fast", "standard", "high_quality"], default="fast")
    parser.add_argument("--crop-first", type=int, nargs=4, help="Real-image simple-page crop: x1 y1 x2 y2, regression only")
    args = parser.parse_args()
    from app.config import OUTPUTS_DIR, TEMP_DIR
    from app.models.project_store import ProjectStore
    from app.services.reconstruction.pipeline import ReconstructionPipeline
    from app.services.scene.ownership import resolve_asset
    from pptx import Presentation
    store = ProjectStore()
    project = store.create("Object First real-image regression")
    if args.crop_first:
        from PIL import Image
        crop_path = TEMP_DIR / f"regression_simple_{project['id']}.png"
        with Image.open(args.images[0]) as source:
            source.crop(args.crop_first).save(crop_path)
        args.images[0] = crop_path
    for source in args.images:
        if not source.is_file():
            raise FileNotFoundError(source)
        store.add_image(project["id"], {"path": str(source.resolve()), "name": source.name})
    pipeline = ReconstructionPipeline(store)
    slides, provider, warnings = pipeline.analyze_project(project["id"], args.mode)
    results = []
    for index, (source, layout) in enumerate(zip(args.images, slides), 1):
        qa = layout["metadata"]["preservationQA"]
        result = {"page": index, "source": str(source), "ocrProvider": provider, "segmentationProvider": layout["metadata"]["segmentationProvider"], "qa": {k: v for k, v in qa.items() if k != "objects"}, "elementCount": len(layout["elements"]), "largeVisualAssetCount": sum(item["type"] == "image" and item["width"] * item["height"] > layout["slide"]["width"] * layout["slide"]["height"] * 0.6 for item in layout["elements"]), "revisions": []}
        results.append(result)
        print(json.dumps({k: v for k, v in result.items() if k != "qa"}, ensure_ascii=False), flush=True)
    baseline_assets = [{item["id"]: item.get("src") for item in slide["elements"]} for slide in slides]
    for round_index in range(2):
        slides, _, _ = pipeline.analyze_project(project["id"], args.mode)
        for index, slide in enumerate(slides):
            qa = slide["metadata"]["preservationQA"]
            results[index]["revisions"].append({"round": round_index + 1, **slide["metadata"]["lastRevision"], "visualAreaPreserved": qa["visualAreaPreserved"], "missingAssetCount": qa["missingAssetCount"], "elementCount": len(slide["elements"])})
            assert {item["id"]: item.get("src") for item in slide["elements"]} == baseline_assets[index]
            for item in slide["elements"]:
                if item.get("src"):
                    assert resolve_asset(item["src"], OUTPUTS_DIR).is_file()
    deck = Presentation(OUTPUTS_DIR / project["id"] / "editable.pptx")
    counts = [{"textboxes": sum(shape.has_text_frame and bool(shape.text.strip()) for shape in slide.shapes), "pictures": sum(shape.shape_type == 13 for shape in slide.shapes), "nativeShapes": sum(shape.shape_type == 1 for shape in slide.shapes)} for slide in deck.slides]
    report = {"projectId": project["id"], "mode": args.mode, "results": results, "pptxObjectCounts": counts, "warnings": warnings}
    output = OUTPUTS_DIR / project["id"] / "object_first_regression.json"
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"REPORT: {output}", flush=True)


if __name__ == "__main__":
    main()
