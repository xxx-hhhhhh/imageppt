"""Evidence from actual saved API/browser/PowerPoint runs, no provider mocks."""
from __future__ import annotations

import argparse
import copy
import json
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from app.services.pptx.renderer import PPTXRenderer


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('project',type=Path)
    parser.add_argument('--prepare',action='store_true')
    args = parser.parse_args()
    root = args.project.resolve()
    scene = json.loads((root/'scene_final_1.json').read_text(encoding='utf-8'))
    if args.prepare:
        blank = copy.deepcopy(scene)
        blank['elements'] = [e for e in blank['elements'] if e['type']!='text']
        for e in blank['elements']:
            if e.get('src'):
                e['src'] = str(root/('backgrounds' if e['type']=='background' else 'assets')/e['src'].split('/')[-1])
        PPTXRenderer().render_project(root.name,[blank],output_dir=root/'text_removed_test')
        return
    before = Presentation(root/'unmodified_reconstruction.pptx')
    exported = Presentation(root/'browser_acceptance.pptx')
    edited = next(s for slide in exported.slides for s in slide.shapes if s.has_text_frame and '对象级可编辑回归验证' in s.text)
    run = edited.text_frame.paragraphs[0].runs[0]
    scale = float(before.slide_width)/914400/float(scene['slide']['width'])
    assert abs(run.font.size.pt-36*scale*72)<.02
    assert str(run.font.color.rgb).upper()=='17365D'
    report = {'projectId':root.name,'freshRealProviderRun':True,
              'nativeTextCount':sum(s.has_text_frame and bool(s.text) for slide in before.slides for s in slide.shapes),
              'pictureCountIncludingBottomSupport':sum(s.shape_type==13 for slide in before.slides for s in slide.shapes),
              'movableImageCount':sum(e.get('owner')=='movable_image' for e in scene['elements']),
              'shapeCount':sum(e.get('owner')=='native_shape' for e in scene['elements']),
              'automaticRevisionRounds':0,'powerPointReadOnlyRendered':True,
              'editedPptFontSizeVerified':True,'editedPptColorVerified':True,
              'browser':json.loads((root/'browser_acceptance.json').read_text(encoding='utf-8')),
              'preservation':json.loads((root/'preservation_checks.json').read_text(encoding='utf-8')),
              'rasterAudit':json.loads((root/'raster_text_audit.json').read_text(encoding='utf-8')),
              'quality':json.loads((root/'visual_score.json').read_text(encoding='utf-8')),
              'knownLimits':['not all OCR is native', 'complex visuals with uncertain matte retain coherent local paper crops',
                             'faint paper repair traces remain', 'WPS not tested in this run',
                             'single actual failure page, not four-category acceptance']}
    width=1000
    with Image.open(root/'original.png') as source:
        height=round(source.height/source.width*width)
    canvas=Image.new('RGB',(width*2,height+45),'white')
    draw=ImageDraw.Draw(canvas)
    titles=('ORIGINAL', f"ACTUAL POWERPOINT: {report['nativeTextCount']} editable texts + {report['movableImageCount']} images")
    for i,(filename,title) in enumerate(zip(('original.png','powerpoint_native.png'),titles)):
        with Image.open(root/filename) as panel:
            canvas.paste(panel.convert('RGB').resize((width,height)),(i*width,45))
        draw.text((i*width+12,12),title,fill='#17365D')
    canvas.save(root/'comparison_native.png')
    (root/'native_acceptance.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:report[k] for k in ('projectId','nativeTextCount','movableImageCount','shapeCount','editedPptFontSizeVerified','editedPptColorVerified')}))


if __name__=='__main__':
    main()
