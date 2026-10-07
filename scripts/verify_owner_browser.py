"""Real upload/analysis/edit/drag/resize/save/export browser acceptance.

Analysis is bounded to standard mode (one critic round), never mocked. Vision
uses the user's existing configuration; no settings or credentials are changed.
"""
import argparse
import json
import re
from pathlib import Path
from urllib.parse import parse_qs, urlencode, urlsplit, urlunsplit

from playwright.sync_api import expect, sync_playwright
from pptx import Presentation


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("image", type=Path)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    errors, missing = [], []
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1600, "height": 1000}, accept_downloads=True)
        page.on("pageerror", lambda error: errors.append(str(error)))
        page.on("response", lambda response: missing.append(response.url) if response.status >= 400 and "/media/" in response.url else None)
        def bounded_analysis(route):
            parts = urlsplit(route.request.url)
            params = parse_qs(parts.query)
            params["mode"] = ["standard"]
            params["allow_fallback"] = ["true"]
            route.continue_(url=urlunsplit(parts._replace(query=urlencode(params, doseq=True))))
        page.route("**/api/projects/*/analyze?*", bounded_analysis)
        page.goto("http://127.0.0.1:5173/", wait_until="domcontentloaded")
        expect(page.get_by_role("button", name="AI设置", exact=True)).to_be_visible()
        page.get_by_role("button", name="AI设置", exact=True).click()
        page.get_by_role("button", name="取消", exact=True).click()
        with page.expect_response(lambda response: "/images" in response.url and response.request.method == "POST", timeout=30000) as uploaded:
            page.locator("input[type=file]").first.set_input_files(str(args.image.resolve()))
        upload = uploaded.value.json()
        identifier = upload["project"]["id"]
        with page.expect_response(lambda response: "/analyze" in response.url and response.request.method == "POST", timeout=240000) as analyzed:
            page.get_by_role("button", name=re.compile("AI解析")).click()
        response = analyzed.value
        assert response.ok, response.status
        result = response.json()
        expect(page.get_by_role("button", name="手动编辑", exact=True).first).to_be_enabled(timeout=30000)
        page.get_by_role("button", name="手动编辑", exact=True).first.click()
        layout = result["slides"][0]
        texts = [item for item in layout["elements"] if item["type"] == "text" and len(item.get("text") or "") >= 4
                 and not any(item.get("metadata", {}).get(k) for k in ("suppressed", "suppressRender", "ownedBy"))]
        assert texts, "Real OCR produced no editable text"
        target = min(texts, key=lambda item: item["y"])
        node = page.locator(f'[data-element-id="{target["id"]}"]')
        expect(node).to_be_visible()
        page.wait_for_function("() => [...document.querySelectorAll('.visual-image img')].every(i => i.complete && i.naturalWidth > 0)", timeout=30000)
        output = root / "outputs" / identifier
        page.screenshot(path=str(output / "browser_before_edit.png"), full_page=True)
        node.dblclick()
        expect(node.locator('[data-text-content]')).to_have_attribute("contenteditable", "true")
        edited = "对象级可编辑回归验证"
        page.keyboard.press("Control+A")
        page.keyboard.insert_text(edited)
        page.get_by_label("字号", exact=True).click()
        expect(page.get_by_role("textbox", name="内容", exact=True)).to_have_value(edited)
        page.get_by_label("字号", exact=True).fill("36")
        page.get_by_label("颜色", exact=True).evaluate("el => { const setter = Object.getOwnPropertyDescriptor(HTMLInputElement.prototype, 'value').set; setter.call(el, '#17365d'); el.dispatchEvent(new Event('input', {bubbles: true})); el.dispatchEvent(new Event('change', {bubbles: true})); }")
        before = node.bounding_box()
        assert before
        page.mouse.move(before["x"]+10, before["y"]+10)
        page.mouse.down()
        page.mouse.move(before["x"]+25, before["y"]+22, steps=8)
        page.mouse.up()
        expect(node).not_to_have_css("left", f'{target["x"]}px')
        handle = node.locator(".resize-handle")
        expect(handle).to_be_visible()
        bounds = handle.bounding_box()
        assert bounds
        page.mouse.move(bounds["x"]+bounds["width"]/2, bounds["y"]+bounds["height"]/2)
        page.mouse.down()
        page.mouse.move(bounds["x"]+20, bounds["y"]+10, steps=6)
        page.mouse.up()
        expect(page.get_by_label("宽度", exact=True)).not_to_have_value(str(target["width"]))
        page.screenshot(path=str(output / "browser_after_edit.png"), full_page=True)
        with page.expect_response(lambda response: "/slides/1" in response.url and response.request.method == "PUT"):
            page.get_by_role("button", name="保存当前页面", exact=True).click()
        page.get_by_role("button", name="返回逐页确认", exact=True).click()
        # These are review controls, not provider mocks.
        if page.get_by_role("button", name="使用当前结果", exact=True).count():
            page.get_by_role("button", name="使用当前结果", exact=True).click()
        expect(page.get_by_role("button", name="通过本页", exact=True)).to_be_visible()
        with page.expect_response(lambda response: "/approve" in response.url and response.request.method == "POST"):
            page.get_by_role("button", name="通过本页", exact=True).click()
        with page.expect_download(timeout=30000) as downloaded:
            page.get_by_role("button", name="导出PPT", exact=True).click()
        output = root / "outputs" / identifier
        pptx = output / "browser_acceptance.pptx"
        downloaded.value.save_as(str(pptx))
        assert any(edited in shape.text for slide in Presentation(pptx).slides for shape in slide.shapes if shape.has_text_frame)
        assert not errors and not missing, {"errors": errors, "missing": missing}
        page.screenshot(path=str(output / "browser_acceptance.png"), full_page=True)
        report = {"projectId": identifier, "realUpload": True, "realAnalysis": True, "mode": "standard", "aiUsed": result.get("aiUsed"),
                  "ocrProvider": result.get("ocrProvider"), "editableText": True, "drag": True, "resize": True, "save": True,
                  "nativeTextInExport": True, "missingAssetCount": len(missing), "browserErrors": errors}
        (output / "browser_acceptance.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(report), flush=True)
        browser.close()


if __name__ == "__main__":
    main()
