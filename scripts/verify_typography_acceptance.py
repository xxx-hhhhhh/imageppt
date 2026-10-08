"""Verify real exported native fonts and complete mixed-size OCR spans."""
import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('project',type=Path)
    args=parser.parse_args()
    root=args.project.resolve()
    scene=json.loads((root/'scene_final_1.json').read_text(encoding='utf-8'))
    native=sorted((e for e in scene['elements'] if e.get('owner')=='editable_text'),key=lambda e:e['zIndex'])
    deck=Presentation(root/'unmodified_reconstruction.pptx')
    runs=[shape.text_frame.paragraphs[0].runs[0] for slide in deck.slides for shape in slide.shapes if shape.has_text_frame and shape.text]
    assert len(runs)==len(native)
    for element,run in zip(native,runs):
        assert run.text==element['text']
        for tag in ('latin','ea','cs'):
            assert run.font._element.find(qn('a:'+tag)).get('typeface')==element['style']['fontFamily']
    decisions=scene['metadata']['ownershipAudit']['textDecisions']
    mixed=[]
    for decision in decisions:
        if decision.get('reason')!='mixed_typography_spans':
            continue
        texts=[e for e in native if e['id'] in decision['editableSpanIds']]
        texts.sort(key=lambda e:e['metadata']['sourceCharRange'][0])
        restored=''.join(e['text'] for e in texts)
        mixed.append({'ocrId':decision['id'],'originalOCRText':decision['text'],'editableText':restored,
                      'complete':restored==decision['text'],'fontSizes':[e['style']['fontSize'] for e in texts]})
    assert mixed and all(item['complete'] for item in mixed)
    complete=sum(d['editableCharacterCount']==d['sourceCharacterCount'] for d in decisions)
    report={'projectId':root.name,'nativeTextBoxCount':len(native),'completeEditableOCRLines':complete,
            'detectedOCRLines':len(decisions),'editableTextCoverage':complete/max(1,len(decisions)),
            'editableCharacterCoverage':sum(d['editableCharacterCount'] for d in decisions)/max(1,sum(d['sourceCharacterCount'] for d in decisions)),
            'eastAsianFontVerifiedForEveryNativeTextbox':True,'mixedLines':mixed,
            'powerPointRenderExists':(root/'powerpoint_native.png').is_file(),
            'automaticRevisionRounds':0,'limits':['Font matching is an installed-font approximation, not font identity proof.',
                                                'Not all OCR lines are native; unverified source text is retained as raster.']}
    (root/'typography_acceptance.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False),flush=True)


if __name__=='__main__':
    main()
