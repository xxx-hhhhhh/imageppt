"""Replay REAL saved OCR/Qwen/SAM evidence, verifying editable text and mattes.

No fake OCR, provider responses, keys or source-page special cases. New records
are diagnostic candidates; the old result is not overwritten.
"""
from __future__ import annotations

import argparse
import copy
import json
import shutil
import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT / "backend"))
from app.models.project_store import ProjectStore
from app.services.pptx.renderer import PPTXRenderer
from app.services.reconstruction.exclusive_ownership import (
    audit_raster_text,
    build_exclusive_scene,
)
from app.services.reconstruction.preservation_checks import (
    inspect_preservation,
)
from app.services.visual_qa.analyzer import render_preview, run_visual_qa


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("recorded_project",type=Path)
    parser.add_argument("--output-root",type=Path)
    parser.add_argument('--refresh-segmentation',action='store_true')
    parser.add_argument('--require-text-id',action='append',default=[],help='Saved OCR ids that must be fully editable')
    args=parser.parse_args()
    previous=args.recorded_project.resolve()
    candidate=next((previous / "candidates").iterdir()) if (previous/'candidates').is_dir() else previous
    store=ProjectStore(args.output_root or ROOT / "outputs")
    record=store.create("Real evidence - native text / coherent visual regression")
    root=store.root / record["id"]
    print(json.dumps({"projectId":record["id"],"output":str(root)}),flush=True)
    shutil.copytree(candidate / "assets",root / "assets")
    ocr=json.loads((candidate / "ocr_layout.json").read_text(encoding="utf-8"))
    scene=json.loads((candidate / "scene_raw.json").read_text(encoding="utf-8"))
    source=previous / "original.png"
    if args.refresh_segmentation:
        from app.services.segmentation.segmentation_provider import (
            create_segmentation_provider,
        )
        provider,_ = create_segmentation_provider('sam2')
        provider.text_regions = [e for e in ocr['elements'] if e.get('type')=='text']
        scene['segmentation'] = provider.segment(source,root/'assets',record['id'])
        print(json.dumps({'segmentationProvider':provider.name,'segments':len(scene['segmentation'])}),flush=True)
    (root/'ocr_layout.json').write_text(json.dumps(ocr,ensure_ascii=False,indent=2),encoding='utf-8')
    (root/'scene_raw.json').write_text(json.dumps(scene,ensure_ascii=False,indent=2),encoding='utf-8')
    layout=build_exclusive_scene(source,ocr,scene,scene.get("segmentation",[]),root,record["id"],1)
    original=json.loads((previous / "slides/page_1.json").read_text(encoding="utf-8"))
    for key in ("ocrProvider","visionProvider","visionModel","segmentationProvider","conversionMode"):
        layout["metadata"][key]=original["metadata"].get(key)
    layout["metadata"]["recordedProviderSourceProject"]=previous.name
    (root/'scene_final_1.json').write_text(json.dumps(layout,ensure_ascii=False,indent=2),encoding='utf-8')
    local=copy.deepcopy(layout)
    for e in local["elements"]:
        if e.get("src"):
            e["src"]=str(root / ("backgrounds" if e["type"]=="background" else "assets") / e["src"].split("/")[-1])
    preview=root / "reconstructed_preview.png"
    render_preview(root / "backgrounds/page_1.png",local,preview)
    audit=audit_raster_text(source,layout,root)
    preservation=inspect_preservation(source,layout,root,preview)
    native=sum(e.get("owner")=="editable_text" for e in layout["elements"])
    images=sum(e.get("owner")=="movable_image" for e in layout["elements"])
    score=run_visual_qa(source,preview,root,local)
    score.update({"revisionRounds":0,"revisionRound":0,"revisionStatus":"automatic_revisions_paused",
                  "editableTextCoverage":native/max(1,len(layout["metadata"]["ownershipAudit"]["textDecisions"])),
                  "visionProvider":layout["metadata"].get("visionProvider"),"visionModel":layout["metadata"].get("visionModel"),
                  "preservationChecks":preservation,"nativeTextCount":native})
    from app.services.validation.reconstruction_quality import reconstruction_quality
    score.update(reconstruction_quality(layout,audit,preservation))
    score['visualSimilarityScore'] = score['overall']
    score['overall'] = min(score['overall'],score['editableCharacterCoverage'])
    (root / "visual_score.json").write_text(json.dumps(score,ensure_ascii=False,indent=2),encoding="utf-8")
    report={"projectId":record["id"],"recordedRealProject":previous.name,"newApiCalls":0,"nativeTextCount":native,
            "movableImageCount":images,"audit":audit,"preservation":preservation,
            "assetPolicies":[e["metadata"].get("visualAssetPolicy") for e in layout["elements"] if e.get("owner")=="movable_image"]}
    (root / "native_recovery.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False),flush=True)
    assert native>0, "Raster-only reconstruction is not editable success"
    assert audit["missingAssetCount"]==audit["rasterNativeDuplicateTextCount"]==0
    assert preservation["unownedSourcePixels"]==preservation["multiplyOwnedSourcePixels"]==preservation["unexpectedVisualChangedPixels"]==0
    assert preservation["unownedGlyphPaperPixels"]==preservation["multiplyOwnedGlyphPaperPixels"]==0
    store.add_image(record["id"],{"id":"source_1","path":str(source),
        "source_url":f"/api/projects/{record['id']}/artifacts/source.png"})
    store.save_slide(record["id"],1,layout)
    shutil.copy2(source,root / "source.png")
    shutil.copy2(source,root / "original.png")
    pptx,_=PPTXRenderer().render_project(record["id"],[local],output_dir=root)
    deck=Presentation(pptx)
    assert sum(s.has_text_frame and bool(s.text) for s in deck.slides[0].shapes)==native
    expected=[e for e in layout['elements'] if e.get('owner')=='editable_text']
    actual=[s.text for s in deck.slides[0].shapes if s.has_text_frame and s.text]
    assert Counter(actual)==Counter(e['text'] for e in expected), 'PPT native text differs from authoritative layout'
    by_id={e['id']:e for e in expected}
    decisions=layout['metadata']['ownershipAudit']['textDecisions']
    complete=[]
    for decision in decisions:
        if decision['decision']!='editable_text':
            continue
        body=''.join(by_id[i]['text'] for i in decision['editableSpanIds'])
        assert body==decision['text'], f"Partial/misordered OCR line: {decision['id']}"
        complete.append(decision['id'])
    assert set(args.require_text_id)<=set(complete), 'Required OCR line still retained as image'
    previous_complete={d['id'] for d in original.get('metadata',{}).get('ownershipAudit',{}).get('textDecisions',[])
                       if d.get('decision')=='editable_text'}
    text_report={'allNativePptStringsMatchLayout':True,'completeOcrLineIds':complete,
                 'requiredOcrLineIds':args.require_text_id,'regressedOcrLineIds':sorted(previous_complete-set(complete)),
                 'editableCharacterCoverage':score['editableCharacterCoverage'],
                 'editableTextCount':score['editableTextCount'],'detectedTextCount':score['detectedTextCount']}
    (root/'text_recovery_verification.json').write_text(json.dumps(text_report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(text_report,ensure_ascii=False),flush=True)
    # Export a genuine blank-text diagnostic: retained old glyphs will show here.
    blank=copy.deepcopy(local)
    blank["elements"]=[e for e in blank["elements"] if e.get("type")!="text"]
    blank_dir=root / "text_removed_test"
    PPTXRenderer().render_project(record["id"],[blank],output_dir=blank_dir)
    print(json.dumps({"projectId":record["id"],"nativeTextBoxesInPpt":native,"pptx":str(pptx)}),flush=True)


if __name__=="__main__":
    main()
