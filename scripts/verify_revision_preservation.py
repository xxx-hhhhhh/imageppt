"""Replay two current regional revisions against real OCR/SAM regression scenes."""
import hashlib
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

import cv2
from app.config import OUTPUTS_DIR
from app.models.project_store import ProjectStore
from app.services.pptx.renderer import PPTXRenderer
from app.services.reconstruction.owner_gate import ownership_evidence
from app.services.reconstruction.revision import revise_problem_regions
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def main():
    path = OUTPUTS_DIR / "object_first_integration_regression.json"
    results = json.loads(path.read_text(encoding="utf-8"))
    store = ProjectStore()
    for result in results:
        identifier = result["projectId"]
        root = OUTPUTS_DIR / identifier
        rounds = []
        for _ in range(2):
            before = hashlib.sha256((root / "reconstructed_preview.png").read_bytes()).hexdigest()
            revision = revise_problem_regions(store, identifier, 1)
            after = hashlib.sha256((root / "reconstructed_preview.png").read_bytes()).hexdigest()
            if not revision["accepted"]:
                assert before == after
            assert revision["missingAssetCount"] == 0
            rounds.append({key: revision.get(key) for key in ("revisionRound", "accepted", "rollbackTriggered", "stagnationReason", "integrityErrors", "assetsBefore", "assetsAfter", "compactedFragmentCount", "unownedPixelsBefore", "unownedPixelsAfter")})
        layout = store.get_slide(identifier, 1)
        _, proof = ownership_evidence(cv2.imread(str(root / "source.png")), layout, root / "assets")
        assert proof["unownedPixelCount"] == 0 and proof["missingAssetCount"] == 0
        deck_path, validation = PPTXRenderer().render_project(identifier, [layout])
        shapes = list(Presentation(deck_path).slides[0].shapes)
        result["currentRevisions"] = rounds
        result["currentOwnerGate"] = proof
        result["currentNativeTextCount"] = sum(s.has_text_frame and bool(s.text.strip()) for s in shapes)
        result["currentNativeShapeCount"] = sum(s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE for s in shapes)
        result["currentImageObjectCount"] = sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in shapes)
        result["pptxValidation"] = validation
        path.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({"case": result["case"], "text": result["currentNativeTextCount"], "shape": result["currentNativeShapeCount"], "image": result["currentImageObjectCount"], "revisions": rounds}), flush=True)


if __name__ == "__main__":
    main()
