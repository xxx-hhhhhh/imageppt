from __future__ import annotations

import cv2
import numpy as np
from pptx import Presentation

from app.services.pptx import renderer as renderer_module
from app.services.reconstruction.residual_objects import _group_components, extract_residual_objects, visual_candidate_mask


def test_disconnected_visual_parts_group_into_movable_ppt_images(tmp_path, monkeypatch):
    image = np.full((300, 600, 3), 255, np.uint8)
    cv2.ellipse(image, (125, 245), (65, 24), 0, 0, 360, (30, 35, 210), -1)
    cv2.ellipse(image, (215, 245), (20, 18), 0, 0, 360, (32, 38, 207), -1)
    cv2.rectangle(image, (370, 75), (450, 185), (65, 65, 65), -1)
    cv2.rectangle(image, (456, 100), (478, 185), (68, 68, 68), -1)
    cv2.circle(image, (530, 40), 9, (190, 115, 20), -1)
    occupied = np.zeros(image.shape[:2], np.uint8)
    project_id = "demo"
    assets, stats = extract_residual_objects(image, occupied, tmp_path / project_id / "assets", project_id, 1)

    assert stats["residualObjectsCount"] == len(assets) == 3
    assert stats["residualCoverageArea"] > 0
    assert stats["residualObjectizationRate"] > 0.95
    assert all(item["width"] * item["height"] < image.shape[0] * image.shape[1] * 0.4 for item in assets)
    monkeypatch.setattr(renderer_module, "OUTPUTS_DIR", tmp_path)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for item in assets:
        renderer_module.PPTXRenderer()._add_element(slide, item, 0.01, 0.01)
    assert len(slide.shapes) == len(assets)
    assert all(shape.shape_type == 13 for shape in slide.shapes)  # separate pictures


def test_many_small_decorations_are_not_silently_dropped_by_asset_limit(tmp_path):
    image = np.full((180, 480, 3), 255, np.uint8)
    centers = []
    for row in range(9):
        for column in range(20):
            x, y = 10 + column * 23, 10 + row * 19
            cv2.rectangle(image, (x, y), (x + 5, y + 5), (40, 90, 190), -1)
            centers.append((x + 2, y + 2))

    assets, stats = extract_residual_objects(
        image, np.zeros(image.shape[:2], np.uint8), tmp_path / "demo" / "assets", "demo", 1,
    )

    assert len(assets) == len(centers)
    assert stats["residualObjectizationRate"] > 0.99
    assert all(any(item["x"] <= x < item["x"] + item["width"]
                   and item["y"] <= y < item["y"] + item["height"] for item in assets)
               for x, y in centers)


def test_spatial_grouping_keeps_same_nearby_color_and_area_rules():
    components = [
        {"box": (20, 20, 30, 30), "color": np.array([30, 90, 190])},
        {"box": (34, 20, 44, 30), "color": np.array([32, 92, 188])},
        {"box": (47, 20, 57, 30), "color": np.array([180, 30, 20])},
        {"box": (140, 20, 150, 30), "color": np.array([30, 90, 190])},
    ]

    groups = _group_components(components, 200, 120)

    assert sorted(sorted(group) for group in groups) == [[0, 1], [2], [3]]


def test_residual_detection_respects_owners_and_colored_background(tmp_path):
    image = np.full((180, 320, 3), (225, 235, 245), np.uint8)
    cv2.rectangle(image, (40, 45), (85, 90), (35, 60, 165), -1)
    cv2.rectangle(image, (190, 50), (245, 110), (25, 90, 160), -1)
    occupied = np.zeros(image.shape[:2], np.uint8)
    occupied[45:91, 40:86] = 255

    assets, stats = extract_residual_objects(image, occupied, tmp_path / "demo" / "assets", "demo", 1)

    assert stats["residualObjectsCount"] == 1
    assert assets[0]["x"] > 180
    assert stats["residualObjectizationRate"] > 0.95


def test_page_sized_environment_is_not_repackaged_as_background_image(tmp_path):
    image = np.full((180, 320, 3), 255, np.uint8)
    cv2.rectangle(image, (5, 5), (314, 174), (20, 70, 100), -1)

    assets, stats = extract_residual_objects(image, np.zeros(image.shape[:2], np.uint8), tmp_path / "demo" / "assets", "demo", 1)

    assert assets == []
    assert stats["residualObjectsCount"] == 0
    assert stats["residualCandidateArea"] == 0  # page environment is excluded from local coverage
    assert stats["residualObjectizationRate"] == 1


def test_large_irregular_map_and_silk_remain_separate_movable_assets(tmp_path):
    image = np.full((400, 700, 3), 255, np.uint8)
    contour = np.array([[45, 55], [300, 18], [650, 48], [675, 170], [620, 270], [490, 300], [270, 320], [65, 275]], np.int32)
    cv2.fillPoly(image, [contour], (105, 145, 175))
    cv2.line(image, (100, 170), (570, 140), (40, 90, 135), 8)
    silk = np.array([[80, 350], [270, 340], [460, 358], [625, 335], [590, 383], [310, 375]], np.int32)
    cv2.fillPoly(image, [silk], (35, 50, 195))
    assets, stats = extract_residual_objects(image, np.zeros(image.shape[:2], np.uint8), tmp_path / "demo" / "assets", "demo", 1)
    assert len(assets) >= 2
    assert any(item["width"] > 500 and item["height"] > 250 for item in assets)
    assert any(item["y"] >= 330 and item["height"] < 60 for item in assets)
    assert stats["residualObjectizationRate"] > 0.95


def test_near_white_page_gradient_keeps_cards_and_ribbon_as_local_assets(tmp_path):
    height, width = 240, 420
    image = np.empty((height, width, 3), np.uint8)
    for y in range(height):
        shade = round(247 + 7 * y / (height - 1))
        image[y, :] = (shade, min(255, shade + 1), min(255, shade + 3))
    cv2.rectangle(image, (35, 45), (175, 145), (244, 244, 253), -1)
    cv2.rectangle(image, (225, 55), (370, 150), (241, 245, 251), -1)
    ribbon = np.array([[20, 190], [140, 178], [280, 198], [400, 180], [400, 239], [20, 239]], np.int32)
    cv2.fillPoly(image, [ribbon], (30, 45, 195))

    mask = visual_candidate_mask(image)
    assert np.mean(mask) < 0.7
    assert mask[90, 100] and mask[100, 300] and mask[220, 200]
    assets, stats = extract_residual_objects(image, np.zeros((height, width), np.uint8), tmp_path / "demo" / "assets", "demo", 1)

    assert stats["residualCoverageArea"] > 25000
    assert any(item["x"] <= 100 <= item["x"] + item["width"] and item["y"] <= 90 <= item["y"] + item["height"] for item in assets)
    assert any(item["x"] <= 300 <= item["x"] + item["width"] and item["y"] <= 100 <= item["y"] + item["height"] for item in assets)
    assert not any(item["x"] <= 100 <= item["x"] + item["width"] and item["x"] <= 300 <= item["x"] + item["width"] and item["y"] <= 90 <= item["y"] + item["height"] for item in assets)
    assert any(item["y"] >= 175 and item["width"] >= 350 for item in assets)
    assert all(item["width"] * item["height"] < width * height * 0.5 for item in assets)


def test_neutral_plate_one_level_darker_than_gradient_is_movable(tmp_path):
    height, width = 240, 420
    image = np.empty((height, width, 3), np.uint8)
    for y in range(height):
        shade = round(247 + 7 * y / (height - 1))
        image[y, :] = (shade, shade, shade)
    assert not np.any(visual_candidate_mask(image))
    cv2.rectangle(image, (70, 60), (190, 150), (248, 248, 248), -1)

    assets, stats = extract_residual_objects(
        image, np.zeros(image.shape[:2], np.uint8), tmp_path / "demo" / "assets", "demo", 1,
    )

    assert len(assets) == 1
    assert assets[0]["x"] <= 70 and assets[0]["y"] <= 60
    assert assets[0]["x"] + assets[0]["width"] >= 190
    assert stats["residualObjectizationRate"] > 0.99


def test_dense_dark_page_does_not_use_near_white_gradient_filter():
    image = np.full((120, 220, 3), (30, 32, 35), np.uint8)
    cv2.rectangle(image, (30, 25), (185, 95), (240, 240, 240), -1)
    mask = visual_candidate_mask(image)
    assert mask[50, 100]


def test_neutral_pale_plate_survives_near_white_gradient_filter(tmp_path):
    height, width = 180, 300
    image = np.empty((height, width, 3), np.uint8)
    for y in range(height):
        shade = round(247 + 7 * y / (height - 1))
        image[y, :] = (shade, shade, shade)
    cv2.rectangle(image, (55, 45), (220, 125), (242, 242, 242), -1)

    mask = visual_candidate_mask(image)
    assert mask[80, 100]
    assert not mask[80, 15]
    assets, _ = extract_residual_objects(image, np.zeros((height, width), np.uint8), tmp_path / "demo" / "assets", "demo", 1)
    assert any(item["x"] <= 55 and item["x"] + item["width"] >= 220 for item in assets)
    assert all(item["width"] * item["height"] < width * height * 0.4 for item in assets)


def test_low_contrast_wide_support_strip_is_movable(tmp_path):
    image = np.full((240, 420, 3), 250, np.uint8)
    cv2.rectangle(image, (25, 75), (395, 110), (248, 248, 248), -1)

    mask = visual_candidate_mask(image)
    assert mask[90, 200]
    assert not mask[30, 200]
    assets, stats = extract_residual_objects(
        image, np.zeros(image.shape[:2], np.uint8), tmp_path / "demo" / "assets", "demo", 1,
    )

    assert len(assets) == 1
    assert assets[0]["x"] <= 25 and assets[0]["x"] + assets[0]["width"] >= 395
    assert assets[0]["height"] < image.shape[0] * 0.25
    assert stats["residualObjectizationRate"] > 0.99


def test_irregular_ring_keeps_transparent_center_when_extracted(tmp_path):
    image = np.full((190, 240, 3), 255, np.uint8)
    cv2.circle(image, (95, 90), 48, (40, 85, 190), 13)
    assets, _ = extract_residual_objects(image, np.zeros(image.shape[:2], np.uint8), tmp_path / "demo" / "assets", "demo", 1)
    ring = next(item for item in assets if item["x"] <= 95 <= item["x"] + item["width"])
    png = cv2.imread(str(tmp_path / "demo" / "assets" / ring["src"].split("/")[-1]), cv2.IMREAD_UNCHANGED)
    assert png.shape[2] == 4
    assert png[90 - ring["y"], 95 - ring["x"], 3] == 0
