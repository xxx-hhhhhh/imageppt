from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation

from app.services.reconstruction.pipeline import ReconstructionPipeline
from app.services.reconstruction.planner import AIReconstructionPlanner
from app.services.vision.schemas import extract_json, validate_json
from app.services.pptx import renderer as renderer_module


def _element(element_id: str, kind: str, box: tuple[int, int, int, int], text: str = "", confidence: float = 0.9) -> dict:
    x, y, width, height = box
    return {"id": element_id, "type": kind, "bbox": {"left": x, "top": y, "width": width, "height": height}, "zIndex": 2, "text": text, "confidence": confidence, "style": {"fontSize": 18}, "metadata": {}}


def test_qwen_plan_accepts_observed_array_boxes_and_missing_bracket() -> None:
    raw = '{"modules":[{"moduleId":"panel","bbox":[0.1,0.2,0.5,0.7],"visualComplexity":"high","editablePriority":2,"ownership":"user","preserveRegions":[{"bbox":[0.2,0.3,0.4,0.6}]}]}'
    payload = extract_json(raw)
    plan = validate_json({"reconstructionPlan": payload}, "scene")["reconstructionPlan"]
    assert len(plan["modules"]) == 1
    panel = plan["modules"][0]
    assert panel["bbox"] == {"left": 0.1, "top": 0.2, "width": 0.4, "height": 0.5}
    assert panel["visualComplexity"] == 0.8
    assert panel["ownership"] == {"owner": "user"}
    assert panel["preserveRegions"][0]["width"] == 0.2


def test_qwen_scene_repairs_malformed_property_separator() -> None:
    payload = extract_json('{"page":{"role":"slide"}, "modules":[{"id":"one" "bbox":{"left":0.1,"top":0.1,"width":0.3,"height":0.3}}]}')
    assert payload["modules"][0]["id"] == "one"


def test_plan_accepts_layered_module_strategies() -> None:
    modules = [
        {"id": strategy, "bbox": {"left": 0.1, "top": 0.1, "width": 0.2, "height": 0.2}, "reconstructionStrategy": strategy}
        for strategy in ("editable_text", "native_shape", "cutout_image", "mixed_component", "background", "ignore")
    ]
    plan = validate_json({"reconstructionPlan": {"modules": modules}}, "scene")["reconstructionPlan"]
    assert [item["reconstructionStrategy"] for item in plan["modules"]] == [item["reconstructionStrategy"] for item in modules]


def test_detected_visual_outside_ai_plan_becomes_editable_image_asset(tmp_path: Path) -> None:
    source = tmp_path / "source.png"
    with Image.new("RGB", (400, 300), "white") as image:
        draw = ImageDraw.Draw(image)
        draw.rectangle((30, 70, 170, 200), fill="#174aa3")
        draw.text((45, 90), "MAP", fill="white")
        image.save(source)
    scene = {
        "canvas": {"width": 400, "height": 300},
        "vision": {"aiUsed": False},
        "regions": [{"id": "map-region", "type": "image", "bbox": {"left": 30, "top": 70, "width": 140, "height": 130}, "confidence": 0.9}],
        "elements": [_element("map-label", "text", (45, 90, 45, 15), "MAP")],
    }
    stats = AIReconstructionPlanner().apply(scene, source, tmp_path / "assets", "test-project", 1)
    assert stats["wholeImageRegions"] == 1
    asset = next(item for item in scene["elements"] if item.get("id", "").startswith("planner_page_1_region_"))
    assert asset["type"] == "image"
    assert asset["metadata"]["textCleaned"] is True
    assert scene["elements"][0]["metadata"]["textCleanedFromAsset"] == asset["id"]
    with Image.open(tmp_path / "assets" / f"{asset['id']}.png") as crop:
        assert crop.size == (140, 130)


def test_whole_chart_owns_cv_and_ocr_without_covering_editable_title(tmp_path: Path, monkeypatch) -> None:
    source_path = tmp_path / "infographic.png"
    with Image.new("RGB", (400, 300), "white") as image:
        draw = ImageDraw.Draw(image)
        draw.rectangle((20, 90, 200, 240), fill="#d1e7ff")
        draw.line((40, 210, 90, 160, 140, 190, 180, 120), fill="#143caa", width=6)
        draw.text((50, 110), "42%", fill="black")
        image.save(source_path)

    plan = {"modules": [
        {"id": "heading", "role": "main_title", "strategy": "editable", "bbox": {"left": 0.02, "top": 0.02, "width": 0.7, "height": 0.18}, "editableIds": ["title"], "confidence": 0.96},
        {"id": "chart", "role": "chart", "strategy": "whole_image", "bbox": {"left": 0.05, "top": 0.3, "width": 0.45, "height": 0.5}, "memberIds": ["chart_label", "wrong_block"], "confidence": 0.92},
    ]}
    vision = validate_json({"reconstructionPlan": plan}, "scene")
    scene = {"canvas": {"width": 400, "height": 300}, "vision": {"aiUsed": True, **vision}, "elements": [
        _element("title", "text", (10, 10, 260, 35), "Editable title"),
        _element("chart_label", "text", (50, 110, 60, 25), "42%"),
        _element("wrong_block", "rectangle", (20, 90, 180, 150)),
        _element("copy_a", "text", (250, 120, 100, 25), "Same text", 0.9),
        _element("copy_b", "text", (251, 121, 100, 25), "Same text", 0.7),
    ]}
    stats = AIReconstructionPlanner().apply(scene, source_path, tmp_path / "assets", "test-project", 1)
    assert stats == {"plannedModules": 2, "wholeImageRegions": 1, "cutoutImages": 1, "nativeShapesPlanned": 0, "plannerSuppressedElements": 2, "plannerDuplicateTexts": 1, "plannerSnappedRegions": 0, "plannerCvVisualRegions": 0}
    by_id = {item["id"]: item for item in scene["elements"]}
    asset = by_id["planner_page_1_region_001"]
    assert asset["metadata"]["reconstructionStrategy"] == "cutout_image"
    assert by_id["title"]["groupId"] == "heading"
    assert not by_id["title"]["metadata"].get("suppressed")
    assert by_id["chart_label"]["metadata"]["textCleanedFromAsset"] == asset["id"]
    assert by_id["wrong_block"]["metadata"]["ownedBy"] == asset["id"]
    assert by_id["copy_b"]["metadata"]["ownedBy"] == "copy_a"
    with Image.open(tmp_path / "assets" / f"{asset['id']}.png") as crop, Image.open(source_path) as original:
        assert crop.mode == "RGB"
        assert crop.tobytes() != original.crop((20, 90, 200, 240)).tobytes()

    base_layout = {"slide": {"width": 400, "height": 300}, "elements": [
        {"id": item["id"], "type": item["type"], "x": item["bbox"]["left"], "y": item["bbox"]["top"], "width": item["bbox"]["width"], "height": item["bbox"]["height"], "zIndex": item["zIndex"], "text": item.get("text"), "style": item.get("style"), "metadata": {}}
        for item in scene["elements"] if not item["id"].startswith("planner_")
    ]}
    layout = ReconstructionPipeline._apply_refined_scene(None, base_layout, scene)
    layout_by_id = {item["id"]: item for item in layout["elements"]}
    assert layout_by_id["chart_label"]["metadata"]["textCleanedFromAsset"] == asset["id"]
    assert layout_by_id[asset["id"]]["metadata"]["textCleaned"] is True
    monkeypatch.setattr(renderer_module, "OUTPUTS_DIR", tmp_path)
    (tmp_path / "test-project" / "assets").mkdir(parents=True)
    (tmp_path / "assets" / f"{asset['id']}.png").replace(tmp_path / "test-project" / "assets" / f"{asset['id']}.png")
    renderer_module.PPTXRenderer().render_project("test-project", [layout])
    slide = Presentation(tmp_path / "test-project" / "editable.pptx").slides[0]
    visible_text = [shape.text for shape in slide.shapes if shape.has_text_frame]
    assert visible_text.count("Editable title") == 1
    assert visible_text.count("Same text") == 1
    assert visible_text.count("42%") == 1
    assert len([shape for shape in slide.shapes if shape.shape_type == 13]) == 1


def test_invalid_plan_is_ignored_and_local_mode_does_not_apply_it(tmp_path: Path) -> None:
    payload = validate_json({"reconstructionPlan": {"modules": [
        {"id": "bad", "strategy": "whole_image", "bbox": {"left": 5, "top": 0, "width": 1, "height": 1}},
    ]}}, "scene")
    assert payload["reconstructionPlan"]["modules"] == []
    source = tmp_path / "source.png"
    Image.new("RGB", (100, 100), "white").save(source)
    scene = {"canvas": {"width": 100, "height": 100}, "vision": {"aiUsed": False, "reconstructionPlan": {"modules": [
        {"id": "unsafe", "strategy": "whole_image", "bbox": {"left": 0, "top": 0, "width": 0.5, "height": 0.5}, "confidence": 1},
    ]}}, "elements": [_element("keep", "text", (10, 10, 50, 20), "Keep me")]}
    stats = AIReconstructionPlanner().apply(scene, source, tmp_path / "assets", "local", 1)
    assert stats["wholeImageRegions"] == 0
    assert scene["elements"][0]["text"] == "Keep me"


def test_overlapping_whole_modules_do_not_claim_the_same_text(tmp_path: Path) -> None:
    source = tmp_path / "source.png"
    Image.new("RGB", (400, 300), "white").save(source)
    scene = {
        "canvas": {"width": 400, "height": 300},
        "vision": {"aiUsed": True, "reconstructionPlan": {"modules": [
            {"id": "upper", "role": "card", "strategy": "whole_image", "bbox": {"left": 0.1, "top": 0.1, "width": 0.7, "height": 0.5}, "confidence": 0.9},
            {"id": "lower", "role": "card", "strategy": "whole_image", "bbox": {"left": 0.1, "top": 0.55, "width": 0.7, "height": 0.35}, "confidence": 0.9},
        ]}},
        "elements": [_element("shared_text", "text", (70, 140, 180, 30), "Shared line")],
    }
    stats = AIReconstructionPlanner().apply(scene, source, tmp_path / "assets", "overlap", 1)
    assert stats["wholeImageRegions"] == 1
    assert len([item for item in scene["elements"] if item["id"].startswith("planner_page_")]) == 1


def test_image_crop_that_slices_a_text_line_is_rejected(tmp_path: Path) -> None:
    source = tmp_path / "source.png"
    Image.new("RGB", (400, 300), "white").save(source)
    scene = {
        "canvas": {"width": 400, "height": 300},
        "vision": {"aiUsed": True, "reconstructionPlan": {"modules": [{
            "id": "visual", "role": "mixed", "strategy": "whole_image",
            "bbox": {"left": 0.1, "top": 0.2, "width": 0.4, "height": 0.4}, "confidence": 0.9,
        }]}},
        "elements": [_element("line", "text", (175, 100, 130, 25), "Important line")],
    }
    stats = AIReconstructionPlanner().apply(scene, source, tmp_path / "assets", "cut", 1)
    assert stats["wholeImageRegions"] == 0
    assert not scene["elements"][0]["metadata"].get("suppressed")


def test_hybrid_chart_snaps_to_cv_figure_and_ignores_hallucinated_members(tmp_path: Path) -> None:
    source = tmp_path / "chart.png"
    Image.new("RGB", (400, 300), "white").save(source)
    scene = {
        "canvas": {"width": 400, "height": 300},
        "regions": [{"type": "figure", "bbox": {"left": 120, "top": 105, "width": 100, "height": 110}, "confidence": 0.9}],
        "vision": {"aiUsed": True, "reconstructionPlan": {"modules": [{
            "id": "panel", "role": "chart_panel", "strategy": "hybrid",
            "bbox": {"left": 0.2, "top": 0.2, "width": 0.5, "height": 0.4},
            "preserveRegions": [{"left": 0.25, "top": 0.3, "width": 0.3, "height": 0.3}],
            "editableIds": ["heading", "chart_label"], "ignoreIds": ["heading"], "confidence": 0.95,
        }]}},
        "elements": [
            _element("heading", "text", (10, 10, 240, 35), "Main heading"),
            _element("chart_label", "text", (140, 140, 50, 25), "Series A"),
            _element("old_crop", "image", (115, 90, 130, 100)),
        ],
    }
    stats = AIReconstructionPlanner().apply(scene, source, tmp_path / "assets", "snap", 1)
    assert stats["wholeImageRegions"] == 1
    assert stats["plannerSnappedRegions"] == 1
    by_id = {item["id"]: item for item in scene["elements"]}
    assert by_id["planner_page_1_region_001"]["bbox"] == {"left": 120, "top": 105, "width": 100, "height": 110}
    assert by_id["chart_label"]["metadata"]["textCleanedFromAsset"] == "planner_page_1_region_001"
    assert by_id["old_crop"]["metadata"]["suppressed"] is True
    assert not by_id["heading"]["metadata"].get("suppressed")


def test_chart_crop_includes_a_nearby_axis_label(tmp_path: Path) -> None:
    source = tmp_path / "chart.png"
    Image.new("RGB", (400, 300), "white").save(source)
    label = _element("axis", "text", (140, 176, 50, 18), "X axis")
    label["metadata"]["rawOCRBBox"] = [140, 176, 190, 194]
    scene = {"canvas": {"width": 400, "height": 300}, "regions": [{"type": "figure", "bbox": {"left": 100, "top": 100, "width": 120, "height": 80}, "confidence": 0.62}], "vision": {"aiUsed": True, "reconstructionPlan": {"modules": [{"id": "panel", "reconstructionStrategy": "mixed_component", "bbox": {"left": 0.2, "top": 0.2, "width": 0.5, "height": 0.6}, "confidence": 0.9}]}}, "elements": [label]}
    stats = AIReconstructionPlanner().apply(scene, source, tmp_path / "assets", "chart", 1)
    assert stats["cutoutImages"] == 1
    asset = next(item for item in scene["elements"] if item.get("type") == "image")
    assert asset["bbox"]["top"] == 100
    assert asset["bbox"]["top"] + asset["bbox"]["height"] >= 194


def test_small_complex_icon_is_extracted_from_mixed_module(tmp_path: Path) -> None:
    source = tmp_path / "icon.png"
    Image.new("RGB", (400, 300), "white").save(source)
    scene = {"canvas": {"width": 400, "height": 300}, "regions": [{"type": "image", "bbox": {"left": 100, "top": 110, "width": 35, "height": 35}, "confidence": 0.68}], "vision": {"aiUsed": True, "reconstructionPlan": {"modules": [{"id": "panel", "reconstructionStrategy": "mixed_component", "bbox": {"left": 0.15, "top": 0.2, "width": 0.45, "height": 0.5}, "confidence": 0.9}]}}, "elements": []}
    stats = AIReconstructionPlanner().apply(scene, source, tmp_path / "assets", "icon", 1)
    assert stats["cutoutImages"] == 1
