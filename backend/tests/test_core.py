from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from fastapi.testclient import TestClient

from app.main import app
from app.services.pptx import PPTXRenderer
from app.services.validation.service import validate_pptx


def make_sample(path: Path) -> None:
    image = Image.new("RGB", (640, 360), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((40, 40, 300, 170), fill="#DCE6F1", outline="#17365D", width=3)
    draw.ellipse((390, 70, 560, 240), fill="#8FBBD9", outline="#17365D", width=3)
    draw.text((75, 85), "Editable title", fill="#17365D")
    image.save(path)


def test_health_and_project_upload(tmp_path: Path) -> None:
    client = TestClient(app)
    assert client.get("/api/health").json()["status"] == "ok"
    project = client.post("/api/projects", json={"name": "test"}).json()
    image_path = tmp_path / "sample.png"
    make_sample(image_path)
    response = client.post("/api/projects/%s/images" % project["id"], files={"files": ("sample.png", image_path.read_bytes(), "image/png")})
    assert response.status_code == 200
    assert response.json()["project"]["imageCount"] == 1


def test_renderer_preserves_editable_object_types(tmp_path: Path) -> None:
    project_id = "test_renderer"
    image_path = tmp_path / "sample.png"
    make_sample(image_path)
    layout = {
        "version": "1.0",
        "slide": {"width": 640, "height": 360},
        "elements": [
            {"id": "bg", "type": "background", "x": 0, "y": 0, "width": 640, "height": 360, "rotation": 0, "zIndex": 0, "src": str(image_path), "style": {}},
            {"id": "text_1", "type": "text", "x": 75, "y": 85, "width": 220, "height": 45, "rotation": 0, "zIndex": 20, "text": "Editable title", "style": {"fontSize": 28, "color": "#17365D"}},
            {"id": "rect_1", "type": "rectangle", "x": 40, "y": 40, "width": 260, "height": 130, "rotation": 0, "zIndex": 10, "style": {"fill": "#DCE6F1", "stroke": "#17365D", "strokeWidth": 2}},
            {"id": "img_1", "type": "image", "x": 390, "y": 70, "width": 170, "height": 170, "rotation": 0, "zIndex": 8, "src": str(image_path), "style": {}},
        ],
    }
    output, report = PPTXRenderer().render_project(project_id, [layout])
    assert output.exists()
    presentation = Presentation(str(output))
    shapes = list(presentation.slides[0].shapes)
    assert any(shape.has_text_frame and "Editable title" in shape.text for shape in shapes)
    assert any(shape.shape_type == 13 for shape in shapes)
    assert report["valid"] is True
    assert validate_pptx(output, [layout])["structural"]["slideCount"] == 1

