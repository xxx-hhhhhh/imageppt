from __future__ import annotations

from pathlib import Path

import cv2
import numpy as np
from pptx import Presentation

from app.services.layout.detectors import detect_text_elements
from app.services.layout.label_detector import detect_label_groups
from app.services.layout.position_refiner import refine_layout
from app.services.layout.typography import estimate_text_style
from app.services.ocr.provider import OCRResult
from app.services.pptx import PPTXRenderer


def _style(text: str, bbox: list[float], width: int = 1000, height: int = 600) -> dict:
    return estimate_text_style(text, bbox, width, height, "#D40D01" if "标题" in text or "标签" in text else "#111111")


def test_typography_has_role_hierarchy() -> None:
    title = estimate_text_style("总体标题", [40, 20, 360, 75], 1000, 600, "#D40D01")
    body = estimate_text_style("正文内容说明", [40, 220, 280, 245], 1000, 600, "#111111")
    label = estimate_text_style("标签标题", [100, 390, 260, 425], 1000, 600, "#D40D01")
    assert title["role"] == "main_title"
    assert body["role"] == "body_text"
    assert label["role"] == "label_text"
    assert title["fontSize"] > body["fontSize"]
    assert title["fontFamily"] == "SimSun"
    assert label["fontFamily"] in {"Microsoft YaHei", "SimHei"}


def test_label_detector_emits_complete_groups(tmp_path: Path) -> None:
    image = np.full((600, 1000, 3), 255, dtype=np.uint8)
    cv2.circle(image, (220, 150), 62, (25, 30, 210), -1)
    cv2.circle(image, (220, 440), 44, (25, 30, 210), -1)
    cv2.imwrite(str(tmp_path / "source.png"), image)
    regions = [
        OCRResult("卡片标题", [165, 285, 275, 320], 0.99, _style("卡片标题", [165, 285, 275, 320])),
        OCRResult("横向标签", [285, 420, 470, 455], 0.99, _style("横向标签", [285, 420, 470, 455])),
        OCRResult("顶部标签", [650, 35, 850, 70], 0.99, _style("顶部标签", [650, 35, 850, 70])),
    ]
    texts = detect_text_elements(regions)
    elements = detect_label_groups(tmp_path / "source.png", 1000, 600, regions, texts, tmp_path / "assets", "regression")
    grouped = [item for item in elements if item.get("groupId")]
    assert grouped
    assert any(item.get("componentType") == "iconCircle" for item in grouped)
    assert any(item.get("componentType") == "titleText" for item in grouped)
    assert any(item.get("componentType") == "container" for item in grouped)
    assert any(item.get("componentType") == "group" for item in grouped)
    assert not any(item.get("componentType") == "iconCircle" and not any(other.get("groupId") == item.get("groupId") and other.get("componentType") in {"titleText", "container"} for other in grouped) for item in grouped)


def test_position_refiner_keeps_source_pixel_coordinates() -> None:
    elements = [{"id": "text", "type": "text", "x": float("nan"), "y": 30, "width": 100, "height": 20, "rotation": 0, "zIndex": 1}]
    refined = refine_layout(elements, 640, 360)
    assert all(np.isfinite(float(item[key])) for item in refined for key in ("x", "y", "width", "height", "rotation"))
    assert refined[0]["x"] >= 0


def test_export_contains_text_shape_and_image(tmp_path: Path) -> None:
    source = tmp_path / "source.png"
    image = np.full((120, 200, 3), 255, dtype=np.uint8)
    cv2.imwrite(str(source), image)
    layout = {"version": "1.1", "slide": {"width": 200, "height": 120}, "elements": [
        {"id": "bg", "type": "background", "x": 0, "y": 0, "width": 200, "height": 120, "rotation": 0, "zIndex": 0, "src": str(source), "style": {}},
        {"id": "container", "type": "roundedRectangle", "x": 20, "y": 20, "width": 160, "height": 80, "rotation": 0, "zIndex": 5, "style": {"fill": "#FBEFED", "stroke": "#F5D7D4"}},
        {"id": "icon", "type": "ellipse", "x": 30, "y": 30, "width": 48, "height": 48, "rotation": 0, "zIndex": 14, "style": {"fill": "#D40D01", "stroke": "#D40D01"}},
        {"id": "title", "type": "text", "x": 88, "y": 40, "width": 82, "height": 26, "rotation": 0, "zIndex": 20, "text": "标签标题", "style": {"fontFamily": "Microsoft YaHei", "fontSize": 20, "fontWeight": 700, "color": "#D40D01"}},
    ]}
    output, report = PPTXRenderer().render_project("regression_layout", [layout])
    assert output.exists()
    shapes = list(Presentation(str(output)).slides[0].shapes)
    assert report["valid"] is True
    assert report["structural"]["textCount"] > 0
    assert report["structural"]["shapeCount"] > 0
    assert report["structural"]["imageCount"] > 0
    assert len(shapes) >= 4
