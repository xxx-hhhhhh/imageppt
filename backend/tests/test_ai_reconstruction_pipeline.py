"""Production acceptance: exclusive ownership replaces erase/critic expectations."""
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest
from app.models.project_store import ProjectStore
from app.services.pptx import renderer
from app.services.reconstruction import pipeline as module
from app.services.reconstruction.exclusive_ownership import (
    audit_raster_text,
    build_exclusive_scene,
)
from app.services.reconstruction.pipeline import (
    AIUnavailableError,
    ReconstructionPipeline,
)
from app.services.visual_qa.analyzer import render_preview
from app.utils import image_io
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont


def fixture_pipeline(monkeypatch, tmp_path, *, ai=True):
    store = ProjectStore(tmp_path / "outputs")
    record = store.create("exclusive-test")
    source = tmp_path / "中文原图.png"
    image = Image.new("RGB", (400, 220), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((20, 45, 380, 170), fill="#EDF4FD")
    draw.ellipse((295, 10, 340, 40), fill="#17365D")
    image.save(source)
    store.add_image(record["id"], {"path": str(source)})
    monkeypatch.setattr(module, "OCRService", lambda _: SimpleNamespace(warnings=[], provider_name="test-ocr", recognize=lambda _: []))
    monkeypatch.setattr(renderer, "OUTPUTS_DIR", store.root)
    pipeline = object.__new__(ReconstructionPipeline)
    pipeline.store = store
    pipeline.segmentation_provider = SimpleNamespace(name="test-seg", warnings=[], segment=lambda *args: [])
    pipeline.scene_analyzer = SimpleNamespace(vision_provider=SimpleNamespace(provider_name="qwen", model_name="test-model"),
        vision_routing={"aiUsed": ai, "usedProvider": "qwen" if ai else "local", "usedModel": "test-model"},
        analyze=lambda *args, **kw: ({"elements": [], "vision": {"reconstructionPlan": {"modules": [
            {"moduleId": "pale-carrier", "role": "module_plate", "bboxPixels": [20,45,381,171]}]}}}, []))
    return pipeline, record["id"], source


def test_pipeline_preserves_pale_plate_and_decoration_without_any_revision(monkeypatch, tmp_path):
    pipeline, project_id, source = fixture_pipeline(monkeypatch, tmp_path)
    slides, _, _ = pipeline.analyze_project(project_id, "maximum")
    layout = slides[0]
    assert layout["metadata"]["automaticRevisionsPaused"]
    assert layout["metadata"]["ownershipAudit"]["unownedPixelCount"] == 0
    assert any(e.get("owner") == "movable_image" for e in layout["elements"])
    assert {e["owner"] for e in layout["elements"]} <= {"movable_image", "native_shape", "editable_text", "intentional_background"}
    root = pipeline.store.root / project_id
    assert np.array_equal(image_io.imread(source), image_io.imread(root / "reconstructed_preview.png"))
    assert (root / "editable.pptx").is_file()


def test_api_failure_never_overwrites_previous_scene_or_assets(monkeypatch, tmp_path):
    pipeline, project_id, _ = fixture_pipeline(monkeypatch, tmp_path)
    pipeline.analyze_project(project_id, "maximum")
    root = pipeline.store.root / project_id
    before = {str(p.relative_to(root)): p.read_bytes() for p in root.glob("assets/*.png")}
    slide_before = (root / "slides/page_1.json").read_bytes()
    pipeline.scene_analyzer.vision_routing["aiUsed"] = False
    with pytest.raises(AIUnavailableError):
        pipeline.analyze_project(project_id, "maximum")
    assert (root / "slides/page_1.json").read_bytes() == slide_before
    assert all((root / name).read_bytes() == value for name, value in before.items())


def test_missing_replacement_asset_prevents_background_write(monkeypatch, tmp_path):
    from app.services.reconstruction import exclusive_ownership as owner
    source = tmp_path / "source.png"
    Image.new("RGB", (120,80), "navy").save(source)
    # Add real content; a flat navy page is legitimate intentional background.
    image = Image.open(source)
    ImageDraw.Draw(image).ellipse((20,20,70,65), fill="red")
    image.save(source)
    original = owner.image_io.imwrite
    monkeypatch.setattr(owner.image_io, "imwrite", lambda path, img, *args: False if "visual" in str(path) else original(path,img,*args))
    with pytest.raises(OSError, match="Replacement asset"):
        build_exclusive_scene(source, {"elements": []}, {}, [], tmp_path / "candidate", "test", 1)
    assert not (tmp_path / "candidate/backgrounds/page_1.png").exists()


def test_publication_failure_restores_existing_scene_ppt_background_and_preview(monkeypatch, tmp_path):
    pipeline, project_id, _ = fixture_pipeline(monkeypatch,tmp_path)
    pipeline.analyze_project(project_id,"maximum")
    root = pipeline.store.root / project_id
    names = ("slides/page_1.json","editable.pptx","backgrounds/page_1.png","reconstructed_preview.png","visual_score.json")
    baseline = {name:(root / name).read_bytes() for name in names}
    original = module.shutil.copy2
    def fail_ppt_publication(source,destination,*args,**kwargs):
        if Path(destination) == root / "editable.pptx":
            raise OSError("simulated publication failure")
        return original(source,destination,*args,**kwargs)
    monkeypatch.setattr(module.shutil,"copy2",fail_ppt_publication)
    with pytest.raises(OSError,match="publication failure"):
        pipeline.analyze_project(project_id,"maximum")
    assert all((root / name).read_bytes() == data for name,data in baseline.items())


def test_measured_native_text_has_no_raster_glyph_duplicate(tmp_path):
    path = Path("C:/Windows/Fonts/arial.ttf")
    if not path.is_file():
        pytest.skip("Font fixture requires Windows Arial")
    source = tmp_path / "plain.png"
    image = Image.new("RGB", (320,100), "white")
    font = ImageFont.truetype(str(path), 28)
    draw = ImageDraw.Draw(image)
    draw.text((25,20), "HELLO", font=font, fill="black")
    image.save(source)
    box = draw.textbbox((25,20), "HELLO", font=font)
    ocr = {"id": "text_001", "type": "text", "text": "HELLO", "confidence": 1,
           "x": box[0], "y": box[1], "width": box[2]-box[0], "height": box[3]-box[1], "style": {}}
    output = tmp_path / "out"
    layout = build_exclusive_scene(source, {"elements": [ocr]}, {}, [], output, "abc", 1)
    assert any(e.get("owner") == "editable_text" for e in layout["elements"])
    assert audit_raster_text(source, layout, output)["rasterNativeDuplicateTextCount"] == 0
    local = deepcopy(layout)
    for e in local["elements"]:
        if e.get("src"):
            e["src"] = str(output / ("backgrounds" if e["type"] == "background" else "assets") / e["src"].split("/")[-1])
    render_preview(output / "backgrounds/page_1.png", local, output / "preview.png")
    assert np.array_equal(image_io.imread(source), image_io.imread(output / "preview.png"))


def test_revision_stop_is_enforced_at_api_and_service(monkeypatch, tmp_path):
    import app.main as api
    from app.services.reconstruction.revision import (
        revise_problem_regions,
        run_revision_loop,
    )
    store = ProjectStore(tmp_path / "outputs")
    record = store.create("pause")
    store.add_image(record["id"], {"path": "not-needed"})
    monkeypatch.setattr(api, "store", store)
    response = TestClient(api.app).post(f"/api/projects/{record['id']}/pages/1/revise")
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "AUTOMATIC_REVISIONS_PAUSED"
    downgrade = TestClient(api.app).post(f"/api/projects/{record['id']}/pages/1/downgrade")
    assert downgrade.status_code == 409
    for call in (revise_problem_regions, run_revision_loop):
        with pytest.raises(ValueError, match="paused"):
            call(store, record["id"], 1)
    assert not (store.root / record["id"] / "slides").exists()


def test_verified_uniform_rectangle_is_a_real_native_shape(tmp_path):
    image = Image.new("RGB", (240,120), "white")
    ImageDraw.Draw(image).rectangle((30,25,169,84),fill="#EDF4FD")
    source = tmp_path / "source.png"
    image.save(source)
    shape = {"id":"shape_001","type":"rectangle","x":30,"y":25,"width":140,"height":60}
    layout = build_exclusive_scene(source,{"elements":[shape]}, {}, [],tmp_path / "out","abc",1)
    assert any(e.get("owner") == "native_shape" for e in layout["elements"])
    # Renderer uses the canonical scene; no raster background-only substitution.
    local = deepcopy(layout)
    for e in local["elements"]:
        if e.get("src"):
            e["src"] = str(tmp_path / "out/backgrounds/page_1.png")
    pptx,_ = renderer.PPTXRenderer().render_project("abc",[local],output_dir=tmp_path / "ppt")
    from pptx import Presentation
    assert any(s.shape_type == 1 for s in Presentation(pptx).slides[0].shapes)


def test_alpha_substrate_cannot_override_a_verified_child_or_claim_its_pixels(tmp_path):
    source = tmp_path / "source.png"
    image = Image.new("RGB",(200,120),"white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((20,20,179,99),fill="#EDF4FD")
    draw.ellipse((60,35,104,79),fill="#17365D")
    image.save(source)
    output = tmp_path / "out"
    (output / "assets").mkdir(parents=True)
    rgb = image_io.imread(source)
    child_mask = np.zeros((45,45),np.uint8)
    import cv2
    cv2.ellipse(child_mask,(22,22),(22,22),0,0,360,255,-1)
    image_io.imwrite(output / "assets/child.png", np.dstack((rgb[35:80,60:105],child_mask)))
    image_io.imwrite(output / "assets/card.png", np.dstack((rgb[20:100,20:180],np.full((80,160),255,np.uint8))))
    segments = [{"id":"child","bbox":{"left":60,"top":35,"width":45,"height":45},"alphaCrop":"/media/assets/abc/child.png"},
                {"id":"card","bbox":{"left":20,"top":20,"width":160,"height":80},"alphaCrop":"/media/assets/abc/card.png"}]
    layout = build_exclusive_scene(source,{"elements":[]},{},segments,output,"abc",1)
    assert any(e.get("metadata",{}).get("edgeSubstratePixels",0)>0 for e in layout["elements"])
    local = deepcopy(layout)
    for e in local["elements"]:
        if e.get("src"):
            e["src"] = str(output / ("backgrounds" if e["type"] == "background" else "assets") / e["src"].split("/")[-1])
    render_preview(output / "backgrounds/page_1.png",local,output / "preview.png")
    from app.services.reconstruction.preservation_checks import inspect_preservation
    check = inspect_preservation(source, layout, output, output / "preview.png")
    # Explicitly measured feather borders may change; non-text interiors and
    # source ownership must not. Exact hard-raster equality defeats antialiasing.
    assert check["unexpectedVisualChangedPixels"] == 0
    assert check["measuredAntialiasEdgePixels"] > 0
    assert np.array_equal(rgb[45:70,70:95],image_io.imread(output / "preview.png")[45:70,70:95])
    count = np.zeros((120,200),np.uint8)
    for node in layout["metadata"]["ownershipAudit"]["nodes"]:
        x1,y1,x2,y2 = node["sourceBBox"]
        count[y1:y2,x1:x2] += image_io.imread(output / node["mask"],-1)[:,:,3]>0
    assert np.all(count == 1)


@pytest.mark.parametrize("surface", ["legacy", "white_objectized"])
def test_pipeline_rejects_erase_first_and_white_objectized_options(monkeypatch, tmp_path, surface):
    pipeline, project_id, _ = fixture_pipeline(monkeypatch, tmp_path)
    monkeypatch.setattr(module, "RECONSTRUCTION_SURFACE_MODE", surface)
    with pytest.raises(ValueError, match="disabled"):
        pipeline.analyze_project(project_id, "maximum")
