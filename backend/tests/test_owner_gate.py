from pathlib import Path

import cv2
import numpy as np
import pytest
from app.services.reconstruction.owner_gate import (
    ensure_visual_owners,
    ownership_evidence,
    tight_text_mask,
)
from app.services.reconstruction.text_erasure import _clean
from app.services.reconstruction.white_objectization import objectize_on_white


def test_unknown_pale_visual_gets_verified_transparent_owner(tmp_path):
    source = np.full((100, 160, 3), 255, np.uint8)
    cv2.circle(source, (65, 50), 25, (248, 249, 254), -1)
    layout = {"elements": []}
    report = ensure_visual_owners(source, layout, tmp_path / "assets", "project", 1)
    assert report["unownedPixelCount"] == 0
    assert report["recoveredPixelCount"] > 1000
    images = layout["elements"]
    assert images and all(item["owner"] == "movable_image" for item in images)
    alpha = cv2.imread(str(tmp_path / "assets" / Path(images[0]["src"]).name), -1)[:, :, 3]
    assert alpha[0, 0] == 0 and alpha.max() == 255


def test_declared_missing_or_blank_image_cannot_claim_source(tmp_path):
    source = np.full((60, 90, 3), 255, np.uint8)
    source[10:50, 10:80] = (40, 80, 140)
    asset = tmp_path / "empty.png"
    cv2.imwrite(str(asset), np.full_like(source, 255))
    layout = {"elements": [{"id": "declared", "type": "image", "x": 0, "y": 0,
                             "width": 90, "height": 60, "src": str(asset), "owner": "movable_image"}]}
    occupied, report = ownership_evidence(source, layout, tmp_path)
    assert not occupied[10:50, 10:80].any()
    assert report["unownedPixelCount"] == 2800
    layout["elements"][0]["src"] = str(tmp_path / "absent.png")
    assert ownership_evidence(source, layout, tmp_path)[1]["missingAssetCount"] == 1


def test_failed_replacement_write_never_clears_background(tmp_path, monkeypatch):
    path, background = tmp_path / "source.png", tmp_path / "background.png"
    source = np.full((40, 70, 3), 255, np.uint8)
    source[5:35, 5:65] = (248, 250, 254)
    cv2.imwrite(str(path), source)
    cv2.imwrite(str(background), source)
    monkeypatch.setattr(cv2, "imwrite", lambda *args: False)
    with pytest.raises(OSError):
        objectize_on_white(path, background, {"elements": []}, tmp_path / "assets", "project", 1)
    assert np.array_equal(cv2.imread(str(background)), source)


def test_bad_geometry_cannot_authorize_deletion(tmp_path):
    source = np.full((50, 80, 3), 255, np.uint8)
    source[5:45, 5:75] = 240
    layout = {"elements": [{"id": "nan", "type": "rectangle", "x": float("nan"), "y": 0,
                             "width": 80, "height": 50, "style": {"fill": "#F0F0F0"}}]}
    assert ownership_evidence(source, layout, tmp_path)[1]["invalidOwnerIds"] == ["nan"]
    assert not ownership_evidence(source, layout, tmp_path)[0][5:45, 5:75].any()


def test_low_confidence_ocr_does_not_claim_card(tmp_path):
    source = np.full((60, 140, 3), 255, np.uint8)
    source[5:55, 5:135] = (244, 250, 252)
    cv2.putText(source, "LABEL", (10, 35), cv2.FONT_HERSHEY_SIMPLEX, .6, (20, 20, 20), 1)
    layout = {"elements": [{"id": "uncertain", "type": "text", "text": "LABEL", "confidence": .1,
                             "x": 5, "y": 10, "width": 110, "height": 30}]}
    assert not ownership_evidence(source, layout, tmp_path)[0][5:55, 5:135].any()


def test_tight_ink_and_external_cleaner_cannot_whiten_card():
    source = np.full((80, 200, 3), (240, 248, 253), np.uint8)
    cv2.putText(source, "TEXT", (45, 48), cv2.FONT_HERSHEY_SIMPLEX, .8, (20, 20, 20), 2)
    region = np.zeros(source.shape[:2], np.uint8)
    region[18:58, 15:175] = 255
    ink = tight_text_mask(source, region, "#141414")
    assert 0 < np.count_nonzero(ink) < np.count_nonzero(region) * .3
    result = _clean(source, region, lambda image, mask: np.full_like(image, 255))
    assert np.array_equal(result[ink == 0], source[ink == 0])


def test_repeated_gate_is_idempotent_and_assets_are_immutable(tmp_path):
    source = np.full((80, 100, 3), 255, np.uint8)
    source[10:70, 10:90] = 248
    layout = {"elements": []}
    first = ensure_visual_owners(source, layout, tmp_path, "project", 1)
    snapshots = {p.name: p.read_bytes() for p in tmp_path.glob("*.png")}
    second = ensure_visual_owners(source, layout, tmp_path, "project", 1)
    assert first["retainedAssetCount"] > 0 and second["retainedAssetCount"] == 0
    assert snapshots == {p.name: p.read_bytes() for p in tmp_path.glob("*.png")}


def test_missing_foreground_replacement_never_authorizes_inpaint(tmp_path):
    from app.services.reconstruction.layered_background import separate_foreground
    background = tmp_path / "background.png"
    source = np.full((80, 100, 3), 255, np.uint8)
    source[20:60, 25:75] = (40, 80, 160)
    cv2.imwrite(str(background), source)
    image = {"id": "missing", "type": "image", "x": 25, "y": 20, "width": 50, "height": 40,
             "src": str(tmp_path / "absent.png"), "metadata": {"preserveWholeAsset": True}}
    assert separate_foreground(background, [image]) == 0
    assert np.array_equal(cv2.imread(str(background)), source)


def test_lama_requires_owner_mask_and_respects_protection():
    from app.services.inpainting.service import InpaintingService
    service = InpaintingService("opencv")
    image = np.full((30, 40, 3), 240, np.uint8)
    image[12:16, 15:22] = 0
    mask = np.zeros(image.shape[:2], np.uint8)
    mask[12:16, 15:22] = 255
    assert np.array_equal(service.clean_array(image, mask), image)
    assert np.array_equal(service.clean_array(image, mask, owner_mask=mask, protected_mask=mask), image)


def test_extremely_wide_real_title_ratio_exports_without_distortion(tmp_path, monkeypatch):
    from app.services.pptx import renderer
    from pptx import Presentation
    monkeypatch.setattr(renderer, "OUTPUTS_DIR", tmp_path)
    path, _ = renderer.PPTXRenderer().render_project("project", [{"slide": {"width": 1588, "height": 115}, "elements": []}])
    deck = Presentation(path)
    assert deck.slide_height >= 914400
    assert abs(deck.slide_width / deck.slide_height - 1588/115) < .0001


@pytest.mark.parametrize("edited", ["User edited text", ""])
def test_ppt_text_uses_canonical_edit_not_stale_ocr_lines(tmp_path, monkeypatch, edited):
    from app.services.pptx import renderer
    from pptx import Presentation
    monkeypatch.setattr(renderer, "OUTPUTS_DIR", tmp_path)
    layout = {"slide": {"width": 800, "height": 450}, "elements": [
        {"id": "edited", "type": "text", "x": 20, "y": 20, "width": 500, "height": 80,
         "text": edited, "lines": [{"text": "Old OCR text"}], "style": {"fontSize": 24}}
    ]}
    path, _ = renderer.PPTXRenderer().render_project("project", [layout])
    assert Presentation(path).slides[0].shapes[0].text == edited
