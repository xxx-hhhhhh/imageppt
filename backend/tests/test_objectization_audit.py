from __future__ import annotations

import json
from pathlib import Path

import cv2
import numpy as np
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app import main as api_main
from app.models.project_store import ProjectStore
from app.services.pptx import renderer
from app.services.reconstruction.objectization_audit import audit_objectization, recover_initial_missing_regions, repair_missing_regions
from app.services.reconstruction.revision import revise_problem_regions
from app.services.reconstruction.white_objectization import objectize_on_white
from app.services.visual_qa.analyzer import render_preview, run_visual_qa


FIXTURE = Path(__file__).parent / "fixtures" / "complex_modules.png"


def test_audit_reports_large_colored_visual_mismatch(tmp_path):
    source = np.full((200, 400, 3), 255, np.uint8)
    preview = source.copy()
    source[20:180, 20:380] = (40, 110, 190)
    preview[20:180, 20:380] = (190, 110, 40)
    for name, image in (("source.png", source), ("background.png", np.full_like(source, 255)),
                        ("preview.png", preview)):
        cv2.imwrite(str(tmp_path / name), image)

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", {"elements": []})

    assert report["missingVisualPixels"] == 0
    assert report["visualMismatchPixels"] > 50000
    assert any(issue["problem"] == "visualContentMismatch" and issue["pixelArea"] > 50000
               for issue in report["issues"])


def test_audit_catches_flat_dark_surface_lost_to_white(tmp_path):
    source = np.full((120, 200, 3), (52, 31, 21), np.uint8)
    cv2.putText(source, "A", (30, 60), cv2.FONT_HERSHEY_SIMPLEX, 1, (245, 245, 245), 2)
    for name, image in (("source.png", source), ("background.png", np.full_like(source, 255)),
                        ("preview.png", np.full_like(source, 255))):
        cv2.imwrite(str(tmp_path / name), image)
    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", {"elements": []})
    assert report["pageSurfaceMismatchPixels"] > 20000
    assert any(issue["problem"] == "pageSurfaceLost" for issue in report["issues"])


def test_audit_catches_two_level_wide_support_strip_erasure(tmp_path):
    source = np.full((240, 420, 3), 250, np.uint8)
    cv2.rectangle(source, (25, 75), (395, 110), (248, 248, 248), -1)
    for name, image in (("source.png", source), ("background.png", np.full_like(source, 255)),
                        ("preview.png", np.full_like(source, 255))):
        cv2.imwrite(str(tmp_path / name), image)

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", {"elements": []})

    assert report["missingVisualPixels"] > 10000
    assert any(issue["problem"] in {"missingBackplate", "missingVisualObject"} for issue in report["issues"])


def test_audit_sees_colored_visual_lost_inside_ocr_box(tmp_path):
    source = np.full((160, 300, 3), 255, np.uint8)
    cv2.circle(source, (60, 75), 16, (50, 170, 40), -1)
    cv2.putText(source, "GO", (92, 83), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (20, 20, 20), 2)
    preview = np.full_like(source, 255)
    cv2.putText(preview, "GO", (94, 83), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (20, 20, 20), 2)
    for name, image in (("source.png", source), ("background.png", np.full_like(source, 255)),
                        ("preview.png", preview)):
        cv2.imwrite(str(tmp_path / name), image)
    layout = {"elements": [{"id": "label", "type": "text", "x": 88, "y": 55,
                            "width": 80, "height": 35, "text": "GO",
                            "metadata": {"rawOCRBBox": [35, 50, 170, 95]}}]}

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", layout)

    assert report["missingVisualPixels"] > 500
    assert any(issue["problem"] == "missingVisualObject" and issue["bbox"][0] <= 60 <= issue["bbox"][2]
               for issue in report["issues"])
    created = repair_missing_regions(tmp_path / "source.png", layout, report["issues"],
                                     tmp_path / "assets", "demo", 1,
                                     preview_path=tmp_path / "preview.png")
    assert len(created) == 1
    asset = cv2.imread(str(tmp_path / "assets" / Path(created[0]["src"]).name), cv2.IMREAD_UNCHANGED)
    assert asset.shape[2] == 4
    assert asset[75 - int(created[0]["y"]), 60 - int(created[0]["x"]), 3] == 255
    assert not (created[0]["x"] <= 110 < created[0]["x"] + created[0]["width"])


def test_audit_does_not_confuse_recolored_editable_glyphs_with_lost_visuals(tmp_path):
    source = np.full((160, 300, 3), 255, np.uint8)
    preview = source.copy()
    cv2.putText(source, "TITLE", (40, 80), cv2.FONT_HERSHEY_SIMPLEX, 1, (20, 20, 220), 2)
    cv2.putText(preview, "TITLE", (48, 80), cv2.FONT_HERSHEY_SIMPLEX, 1, (20, 20, 220), 2)
    for name, image in (("source.png", source), ("background.png", np.full_like(source, 255)),
                        ("preview.png", preview)):
        cv2.imwrite(str(tmp_path / name), image)
    layout = {"elements": [{"id": "title", "type": "text", "x": 40, "y": 45,
                            "width": 150, "height": 45, "text": "TITLE",
                            "style": {"color": "#DC1414"},
                            "metadata": {"rawOCRBBox": [35, 45, 195, 90]}}]}

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", layout)

    assert report["missingVisualObjects"] == 0


def test_initial_recovery_restores_icon_inside_text_box(tmp_path, monkeypatch):
    monkeypatch.setattr(renderer, "OUTPUTS_DIR", tmp_path)
    root = tmp_path / "demo"
    (root / "backgrounds").mkdir(parents=True)
    source = np.full((160, 300, 3), 255, np.uint8)
    cv2.circle(source, (60, 75), 16, (50, 170, 40), -1)
    cv2.putText(source, "GO", (92, 83), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (20, 20, 20), 2)
    source_path = root / "source.png"
    background_path = root / "backgrounds" / "page_1.png"
    preview_path = root / "preview.png"
    cv2.imwrite(str(source_path), source)
    cv2.imwrite(str(background_path), np.full_like(source, 255))
    layout = {"slide": {"width": 300, "height": 160}, "elements": [
        {"id": "label", "type": "text", "x": 88, "y": 55, "width": 80, "height": 35,
         "zIndex": 20, "text": "GO", "style": {"color": "#141414", "fontSize": 25},
         "metadata": {"rawOCRBBox": [35, 50, 170, 95]}}]}
    render_preview(background_path, layout, preview_path)
    before = audit_objectization(source_path, background_path, preview_path, layout)
    assert before["missingVisualObjects"] >= 1

    recovered = recover_initial_missing_regions(source_path, background_path, preview_path,
                                                layout, root / "assets", "demo", 1)

    after = audit_objectization(source_path, background_path, preview_path, layout)
    assert recovered == 1
    assert after["missingVisualObjects"] == 0
    assert cv2.imread(str(preview_path))[75, 60].tolist() == [50, 170, 40]
    assert len([item for item in layout["elements"] if item["type"] == "text"]) == 1


def test_audit_flags_opaque_slide_screenshot_but_not_irregular_cutout(tmp_path):
    source = np.full((180, 300, 3), 255, np.uint8)
    cv2.rectangle(source, (30, 25), (100, 95), (35, 95, 180), -1)
    cv2.rectangle(source, (180, 50), (260, 135), (40, 160, 70), -1)
    (tmp_path / "assets").mkdir()
    cv2.imwrite(str(tmp_path / "source.png"), source)
    cv2.imwrite(str(tmp_path / "background.png"), np.full_like(source, 255))
    cv2.imwrite(str(tmp_path / "preview.png"), source)
    cv2.imwrite(str(tmp_path / "assets" / "screenshot.png"), source)
    opaque = {"elements": [{"id": "screenshot", "type": "image", "x": 0, "y": 0,
                            "width": 300, "height": 180, "src": "/media/assets/demo/screenshot.png"}]}

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", opaque)
    assert report["monolithicPageImageCount"] == 1
    assert any(issue["problem"] == "monolithicPageImage" for issue in report["issues"])

    alpha = np.zeros(source.shape[:2], np.uint8)
    polygon = np.array([[5, 25], [100, 5], [295, 30], [280, 165], [35, 175]], np.int32)
    cv2.fillPoly(alpha, [polygon], 255)
    cutout = np.dstack((source, alpha))
    cv2.imwrite(str(tmp_path / "assets" / "cutout.png"), cutout)
    irregular = {"elements": [{"id": "cutout", "type": "image", "x": 0, "y": 0,
                               "width": 300, "height": 180, "src": "/media/assets/demo/cutout.png"}]}
    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", irregular)
    assert report["monolithicPageImageCount"] == 0


def test_audit_flags_tiny_solid_corner_mark_lost_on_large_slide(tmp_path):
    source = np.full((900, 1600, 3), 255, np.uint8)
    source[120:124, 130:134] = (35, 85, 205)
    source[360, 420] = (90, 90, 90)
    for name, image in (("source.png", source), ("background.png", np.full_like(source, 255)),
                        ("preview.png", np.full_like(source, 255))):
        cv2.imwrite(str(tmp_path / name), image)

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", {"elements": []})

    assert report["missingVisualObjects"] == 1
    assert report["issues"][-1]["bbox"][0] <= 130


def test_large_bounded_visual_loss_is_reported_and_recovered_as_shape(tmp_path, monkeypatch):
    monkeypatch.setattr(renderer, "OUTPUTS_DIR", tmp_path)
    root = tmp_path / "demo"
    (root / "backgrounds").mkdir(parents=True)
    source = np.full((200, 400, 3), 255, np.uint8)
    cv2.rectangle(source, (20, 20), (379, 179), (50, 110, 190), -1)
    source_path = root / "source.png"
    background_path = root / "backgrounds" / "page_1.png"
    preview_path = root / "preview.png"
    cv2.imwrite(str(source_path), source)
    cv2.imwrite(str(background_path), np.full_like(source, 255))
    cv2.imwrite(str(preview_path), np.full_like(source, 255))
    layout = {"slide": {"width": 400, "height": 200}, "elements": []}

    before = audit_objectization(source_path, background_path, preview_path, layout)
    assert before["missingVisualPixels"] > 50000
    assert any(issue["problem"] == "largeVisualLoss" for issue in before["issues"])
    recovered = recover_initial_missing_regions(source_path, background_path, preview_path,
                                                layout, root / "assets", "demo", 1)

    assert recovered == 1
    assert layout["elements"][0]["type"] == "rectangle"
    assert audit_objectization(source_path, background_path, preview_path, layout)["missingVisualPixels"] == 0
    assert np.all(cv2.imread(str(background_path)) == 255)


def test_large_textured_visual_loss_recovers_as_movable_image(tmp_path, monkeypatch):
    monkeypatch.setattr(renderer, "OUTPUTS_DIR", tmp_path)
    root = tmp_path / "demo"
    (root / "backgrounds").mkdir(parents=True)
    source = np.full((200, 400, 3), 255, np.uint8)
    for y in range(20, 180):
        source[y, 20:380] = (50 + y // 4, 110, 190)
    source_path = root / "source.png"
    background_path = root / "backgrounds" / "page_1.png"
    preview_path = root / "preview.png"
    cv2.imwrite(str(source_path), source)
    cv2.imwrite(str(background_path), np.full_like(source, 255))
    cv2.imwrite(str(preview_path), np.full_like(source, 255))
    layout = {"slide": {"width": 400, "height": 200}, "elements": []}

    recovered = recover_initial_missing_regions(source_path, background_path, preview_path,
                                                layout, root / "assets", "demo", 1)

    assert recovered == 1
    assert layout["elements"][0]["type"] == "image"
    assert (root / "assets" / Path(layout["elements"][0]["src"]).name).is_file()
    assert audit_objectization(source_path, background_path, preview_path, layout)["missingVisualPixels"] == 0


def test_audit_does_not_mark_preserved_white_badge_detail_as_missing(tmp_path):
    source = np.full((150, 240, 3), 255, np.uint8)
    cv2.circle(source, (95, 75), 40, (30, 85, 195), -1)
    cv2.line(source, (75, 75), (115, 75), (255, 255, 255), 6)
    cv2.line(source, (95, 55), (95, 95), (255, 255, 255), 6)
    for name, image in (("source.png", source), ("background.png", np.full_like(source, 255)), ("preview.png", source)):
        cv2.imwrite(str(tmp_path / name), image)

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", {"elements": []})

    assert report["missingVisualPixels"] == 0
    assert report["missingVisualObjects"] == 0
    assert report["retainedVisualCoverage"] == 1.0


def test_first_pass_recovers_missing_small_visual_as_movable_asset(tmp_path, monkeypatch):
    monkeypatch.setattr(renderer, "OUTPUTS_DIR", tmp_path)
    source = np.full((160, 260, 3), 255, np.uint8)
    cv2.rectangle(source, (90, 65), (112, 88), (35, 105, 190), -1)
    for name, image in (("source.png", source), ("background.png", np.full_like(source, 255)),
                        ("preview.png", np.full_like(source, 255))):
        cv2.imwrite(str(tmp_path / name), image)
    layout = {"slide": {"width": 260, "height": 160}, "elements": []}
    assets = tmp_path / "demo" / "assets"

    recovered = recover_initial_missing_regions(tmp_path / "source.png", tmp_path / "background.png",
                                                tmp_path / "preview.png", layout, assets, "demo", 1)

    assert recovered == 1
    assert len(layout["elements"]) == 1
    assert layout["elements"][0]["type"] == "image"
    assert layout["elements"][0]["src"].startswith("/media/assets/demo/initial_page_1_")
    assert (assets / Path(layout["elements"][0]["src"]).name).is_file()
    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", layout)
    assert report["missingVisualObjects"] == 0
    assert report["missingVisualPixels"] < 20


def test_first_pass_discards_recovery_that_does_not_render_better(tmp_path, monkeypatch):
    source = np.full((160, 260, 3), 255, np.uint8)
    cv2.rectangle(source, (90, 65), (112, 88), (35, 105, 190), -1)
    for name, image in (("source.png", source), ("background.png", np.full_like(source, 255)),
                        ("preview.png", np.full_like(source, 255))):
        cv2.imwrite(str(tmp_path / name), image)
    from app.services.visual_qa import analyzer
    monkeypatch.setattr(analyzer, "render_preview", lambda _background, _layout, output: cv2.imwrite(
        str(output), np.full_like(source, 255)))
    layout = {"slide": {"width": 260, "height": 160}, "elements": []}
    assets = tmp_path / "assets"
    original_preview = (tmp_path / "preview.png").read_bytes()

    recovered = recover_initial_missing_regions(tmp_path / "source.png", tmp_path / "background.png",
                                                tmp_path / "preview.png", layout, assets, "demo", 1)

    assert recovered == 0
    assert layout["elements"] == []
    assert (tmp_path / "preview.png").read_bytes() == original_preview
    assert not list(assets.glob("*.png"))


def test_audit_detects_missing_pale_plate_and_writes_owner_debug(tmp_path):
    source = np.full((180, 300, 3), 255, np.uint8)
    cv2.rectangle(source, (30, 40), (210, 112), (246, 248, 250), -1)
    cv2.imwrite(str(tmp_path / "source.png"), source)
    cv2.imwrite(str(tmp_path / "background.png"), np.full_like(source, 255))
    cv2.imwrite(str(tmp_path / "preview.png"), np.full_like(source, 255))
    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png", tmp_path / "preview.png", {"elements": []}, tmp_path / "debug.png")
    assert report["whiteBackground"] is True
    assert report["missingBackplates"] == 1
    assert report["missingVisualObjects"] == 0
    assert report["issues"][0]["bbox"][0] <= 30
    assert (tmp_path / "debug.png").is_file()


def test_sparse_pale_chart_line_is_a_visual_object_not_solid_backplate(tmp_path):
    source = np.full((130, 230, 3), 255, np.uint8)
    cv2.line(source, (30, 85), (190, 45), (225, 230, 235), 2)
    for name, image in (("source.png", source), ("background.png", np.full_like(source, 255)),
                        ("preview.png", np.full_like(source, 255))):
        cv2.imwrite(str(tmp_path / name), image)

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", {"elements": []})

    assert report["missingVisualObjects"] >= 1
    assert report["missingBackplates"] == 0


def test_audit_detects_neutral_plate_missing_from_gradient_page(tmp_path):
    source = np.empty((180, 300, 3), np.uint8)
    for y in range(180):
        shade = round(247 + 7 * y / 179)
        source[y, :] = (shade, shade, shade)
    cv2.rectangle(source, (55, 45), (220, 125), (242, 242, 242), -1)
    cv2.imwrite(str(tmp_path / "source.png"), source)
    cv2.imwrite(str(tmp_path / "background.png"), np.full_like(source, 255))
    cv2.imwrite(str(tmp_path / "preview.png"), np.full_like(source, 255))

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png", tmp_path / "preview.png", {"elements": []})

    assert report["missingVisualPixels"] > 10000
    assert report["retainedVisualCoverage"] < 0.1
    assert report["missingBackplates"] + report["missingVisualObjects"] >= 1


def test_audit_counts_visual_recolored_instead_of_erased(tmp_path):
    source = np.full((150, 250, 3), 255, np.uint8)
    cv2.rectangle(source, (35, 35), (175, 105), (70, 95, 185), -1)
    preview = source.copy()
    preview[35:106, 35:176] = (190, 210, 235)
    cv2.imwrite(str(tmp_path / "source.png"), source)
    cv2.imwrite(str(tmp_path / "background.png"), np.full_like(source, 255))
    cv2.imwrite(str(tmp_path / "preview.png"), preview)

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png", tmp_path / "preview.png", {"elements": []})

    assert report["missingVisualPixels"] == 0
    assert report["visualMismatchPixels"] > 9000
    assert report["salientVisualPixels"] >= report["visualMismatchPixels"]


def test_suppressed_badge_text_does_not_hide_erased_internal_symbol(tmp_path):
    source = np.full((150, 240, 3), 255, np.uint8)
    cv2.circle(source, (95, 75), 40, (30, 45, 190), -1)
    cv2.putText(source, "AI", (77, 84), cv2.FONT_HERSHEY_SIMPLEX, 0.55, (255, 255, 255), 2)
    preview = source.copy()
    cv2.rectangle(preview, (76, 65), (111, 88), (255, 255, 255), -1)
    cv2.imwrite(str(tmp_path / "source.png"), source)
    cv2.imwrite(str(tmp_path / "background.png"), np.full_like(source, 255))
    cv2.imwrite(str(tmp_path / "preview.png"), preview)
    layout = {"elements": [
        {"id": "badge", "type": "image", "x": 55, "y": 35, "width": 80, "height": 80,
         "src": "badge.png", "metadata": {"wholeBadgeAsset": True}},
        {"id": "letters", "type": "text", "x": 76, "y": 65, "width": 36, "height": 23,
         "text": "AI", "metadata": {"rawOCRBBox": [76, 65, 112, 88], "suppressed": True, "ownedBy": "badge"}},
    ]}

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png", tmp_path / "preview.png", layout)

    assert report["missingVisualPixels"] > 100
    assert any(issue["problem"] in {"missingVisualObject", "visualContentMismatch"} for issue in report["issues"])


@pytest.mark.parametrize("text_metadata,text_box", [
    ({"rawOCRBBox": [25, 25, 90, 75], "suppressed": True, "ownedBy": "visual"}, (25, 25, 65, 50)),
    ({"rawOCRBBox": [25, 25, 65, 40]}, (25, 25, 120, 80)),
])
def test_text_box_cannot_hide_missing_visual_outside_active_ocr_ink(tmp_path, text_metadata, text_box):
    source = np.full((120, 180, 3), 255, np.uint8)
    cv2.rectangle(source, (75, 50), (105, 85), (25, 80, 180), -1)
    cv2.imwrite(str(tmp_path / "source.png"), source)
    white = np.full_like(source, 255)
    cv2.imwrite(str(tmp_path / "background.png"), white)
    cv2.imwrite(str(tmp_path / "preview.png"), white)
    layout = {"elements": [{"id": "ocr", "type": "text", "text": "Label", "x": text_box[0], "y": text_box[1],
                            "width": text_box[2], "height": text_box[3], "metadata": text_metadata}]}

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", layout)

    assert report["missingVisualPixels"] > 600
    assert report["missingVisualObjects"] >= 1


def test_missing_visual_repair_masks_only_lost_pixels_not_neighboring_text(tmp_path):
    source = np.full((120, 220, 3), 255, np.uint8)
    source[30:100, 20:200] = (247, 246, 245)
    cv2.rectangle(source, (42, 50), (75, 82), (25, 85, 185), -1)
    cv2.putText(source, "LABEL", (117, 70), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (20, 45, 90), 1)
    preview = source.copy()
    preview[50:83, 42:76] = 255
    cv2.imwrite(str(tmp_path / "source.png"), source)
    cv2.imwrite(str(tmp_path / "preview.png"), preview)
    layout = {"elements": [{"id": "label", "type": "text", "text": "LABEL", "x": 112, "y": 50,
                            "width": 85, "height": 28, "zIndex": 10,
                            "metadata": {"rawOCRBBox": [115, 54, 177, 73]}}]}
    created = repair_missing_regions(tmp_path / "source.png", layout,
                                     [{"problem": "missingVisualObject", "bbox": [20, 30, 200, 100]}],
                                     tmp_path / "assets", "test", 0, preview_path=tmp_path / "preview.png")

    assert len(created) == 1
    item = created[0]
    assert item["type"] == "image"
    assert 40 <= item["x"] <= 44 and item["x"] + item["width"] <= 78
    assert item["metadata"]["sourceMaskPixels"] > 900
    cutout = cv2.imread(str(tmp_path / "assets" / Path(item["src"]).name), cv2.IMREAD_UNCHANGED)
    assert cutout.shape[2] == 4
    assert np.count_nonzero(cutout[:, :, 3]) == item["metadata"]["sourceMaskPixels"]


def test_small_font_baseline_shift_is_not_a_missing_visual(tmp_path):
    source = np.full((90, 180, 3), 255, np.uint8)
    preview = source.copy()
    cv2.putText(source, "LABEL", (20, 42), cv2.FONT_HERSHEY_SIMPLEX, 0.65, (25, 50, 105), 1)
    cv2.putText(preview, "LABEL", (20, 44), cv2.FONT_HERSHEY_SIMPLEX, 0.65, (25, 50, 105), 1)
    cv2.imwrite(str(tmp_path / "source.png"), source)
    cv2.imwrite(str(tmp_path / "preview.png"), preview)
    cv2.imwrite(str(tmp_path / "background.png"), np.full_like(source, 255))
    layout = {"elements": [{"id": "label", "type": "text", "text": "LABEL", "x": 20, "y": 20,
                            "width": 70, "height": 28,
                            "metadata": {"rawOCRBBox": [20, 27, 87, 43]}}]}

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", layout)

    assert report["missingVisualObjects"] == 0
    assert report["visualMismatchRegions"] == 0


def test_audit_reports_contiguous_missing_visual_region(tmp_path):
    source = np.full((120, 180, 3), 255, np.uint8)
    cv2.rectangle(source, (60, 40), (79, 59), (20, 90, 180), -1)
    cv2.imwrite(str(tmp_path / "source.png"), source)
    cv2.imwrite(str(tmp_path / "background.png"), np.full_like(source, 255))
    cv2.imwrite(str(tmp_path / "preview.png"), np.full_like(source, 255))

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", {"elements": []})

    assert report["largestMissingVisualRegion"] >= 380


def test_large_slide_reports_all_missing_small_icons(tmp_path):
    source = np.full((1000, 1000, 3), 255, np.uint8)
    for row in range(7):
        for column in range(10):
            x, y = 30 + column * 28, 30 + row * 28
            cv2.rectangle(source, (x, y), (x + 7, y + 7), (30, 90, 170), -1)
    white = np.full_like(source, 255)
    cv2.imwrite(str(tmp_path / "source.png"), source)
    cv2.imwrite(str(tmp_path / "background.png"), white)
    cv2.imwrite(str(tmp_path / "preview.png"), white)

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png",
                                 tmp_path / "preview.png", {"elements": []})

    assert report["missingVisualObjects"] == 70
    assert len([issue for issue in report["issues"] if issue["problem"] == "missingVisualObject"]) == 70


def test_audit_flags_image_owner_that_renders_blank(tmp_path):
    source = np.full((150, 240, 3), 255, np.uint8)
    cv2.circle(source, (70, 70), 22, (40, 80, 180), -1)
    cv2.imwrite(str(tmp_path / "source.png"), source)
    cv2.imwrite(str(tmp_path / "background.png"), np.full_like(source, 255))
    cv2.imwrite(str(tmp_path / "preview.png"), np.full_like(source, 255))
    layout = {"elements": [{"id": "badge", "type": "image", "x": 45, "y": 45, "width": 50, "height": 50, "src": "broken.png"}]}
    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png", tmp_path / "preview.png", layout)
    assert report["blankVisualOwners"] == 1
    assert any(item["problem"] == "blankVisualOwner" and item["elementId"] == "badge" for item in report["issues"])


def test_sparse_transparent_asset_is_not_falsely_reported_blank(tmp_path):
    source = np.full((150, 240, 3), 255, np.uint8)
    source[60:72, 80:91] = (35, 70, 190)
    asset_dir = tmp_path / "assets"
    asset_dir.mkdir()
    asset = np.zeros((50, 60, 4), np.uint8)
    asset[:, :, :3] = 255
    asset[20:32, 20:31, :3] = (35, 70, 190)
    asset[20:32, 20:31, 3] = 255
    cv2.imwrite(str(asset_dir / "sparse.png"), asset)
    for name, image in (("source.png", source), ("background.png", np.full_like(source, 255)), ("preview.png", source)):
        cv2.imwrite(str(tmp_path / name), image)
    layout = {"elements": [{"id": "sparse", "type": "image", "x": 60, "y": 40, "width": 60, "height": 50,
                             "src": "/media/assets/demo/sparse.png"}]}

    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png", tmp_path / "preview.png", layout)

    assert report["blankVisualOwners"] == 0


def test_transparent_image_bbox_does_not_hide_missing_visual_content(tmp_path):
    source = np.full((150, 300, 3), 255, np.uint8)
    cv2.rectangle(source, (30, 45), (75, 100), (40, 90, 180), -1)
    cv2.rectangle(source, (175, 45), (220, 100), (50, 120, 70), -1)
    background = np.full_like(source, 255)
    preview = background.copy()
    preview[45:101, 30:76] = source[45:101, 30:76]
    asset_dir = tmp_path / "assets"
    asset_dir.mkdir()
    asset = np.zeros((80, 220, 4), np.uint8)
    asset[:, :, :3] = 255
    asset[0:56, 0:46, 3] = 255
    cv2.imwrite(str(asset_dir / "partial.png"), asset)
    for name, image in (("source.png", source), ("background.png", background), ("preview.png", preview)):
        cv2.imwrite(str(tmp_path / name), image)
    layout = {"elements": [{"id": "partial", "type": "image", "x": 30, "y": 45, "width": 220, "height": 80, "src": "/media/assets/demo/partial.png"}]}
    report = audit_objectization(tmp_path / "source.png", tmp_path / "background.png", tmp_path / "preview.png", layout)
    assert report["missingVisualObjects"] >= 1
    assert any(issue["bbox"][0] <= 175 <= issue["bbox"][2] for issue in report["issues"] if issue["problem"] == "missingVisualObject")


def test_initial_recovery_restores_pale_gaps_inside_movable_card(tmp_path, monkeypatch):
    monkeypatch.setattr(renderer, "OUTPUTS_DIR", tmp_path)
    project_id = "a" * 32
    root = tmp_path / project_id
    (root / "assets").mkdir(parents=True)
    source = np.full((180, 300, 3), 255, np.uint8)
    source[35:145, 35:265] = (252, 252, 252)
    alpha = np.zeros((110, 230), np.uint8)
    alpha[:, :] = 255
    alpha[45:90, 90:185] = 0
    asset = np.dstack((source[35:145, 35:265], alpha))
    cv2.imwrite(str(root / "assets" / "card.png"), asset)
    source_path, background, preview = root / "source.png", root / "background.png", root / "preview.png"
    cv2.imwrite(str(source_path), source)
    cv2.imwrite(str(background), np.full_like(source, 255))
    layout = {"slide": {"width": 300, "height": 180}, "elements": [
        {"id": "card", "type": "image", "x": 35, "y": 35, "width": 230, "height": 110, "zIndex": 1,
         "src": f"/media/assets/{project_id}/card.png"},
        {"id": "label", "type": "text", "text": "Label", "x": 115, "y": 70, "width": 65, "height": 20,
         "zIndex": 2, "style": {"fontSize": 17, "color": "#502D23"}},
    ]}
    render_preview(background, layout, preview)
    before = cv2.imread(str(preview))
    assert np.all(before[112, 170] == 255)
    assert audit_objectization(source_path, background, preview, layout)["paleAssetGapPixels"] > 1000

    recovered = recover_initial_missing_regions(source_path, background, preview, layout, root / "assets", project_id, 1)

    assert recovered >= 1
    after = cv2.imread(str(preview))
    assert np.max(np.abs(after[112, 170].astype(int) - source[112, 170].astype(int))) < 3
    assert any((item.get("metadata") or {}).get("qaIssue") == "paleAssetGap" for item in layout["elements"])
    assert any((item.get("metadata") or {}).get("qaIssue") == "paleTextSupportGap" for item in layout["elements"])
    assert audit_objectization(source_path, background, preview, layout)["paleAssetGapPixels"] < 200
    assert next(item for item in layout["elements"] if item["id"] == "label")["type"] == "text"


def test_residual_ribbon_absorbs_its_pale_support_as_one_movable_asset(tmp_path, monkeypatch):
    from app.services.reconstruction.objectization_audit import _recover_pale_asset_gaps

    monkeypatch.setattr(renderer, "OUTPUTS_DIR", tmp_path)
    project_id = "b" * 32
    root = tmp_path / project_id
    assets = root / "assets"
    assets.mkdir(parents=True)
    source = np.full((130, 260, 3), 255, np.uint8)
    source[35:105, 25:235] = (246, 245, 244)
    source[60:90, 25:235] = (30, 50, 190)
    source_path = root / "source.png"
    cv2.imwrite(str(source_path), source)
    alpha = np.zeros((70, 210), np.uint8)
    alpha[25:55] = 255
    original_path = assets / "ribbon.png"
    cv2.imwrite(str(original_path), np.dstack((source[35:105, 25:235], alpha)))
    original_bytes = original_path.read_bytes()
    layout = {"slide": {"width": 260, "height": 130}, "elements": [
        {"id": "ribbon", "type": "image", "x": 25, "y": 35, "width": 210, "height": 70,
         "zIndex": 1, "src": f"/media/assets/{project_id}/ribbon.png",
         "metadata": {"reconstructionStrategySource": "residual_detection", "layerRole": "residual"}},
    ]}
    background = root / "background.png"
    preview = root / "preview.png"
    cv2.imwrite(str(background), np.full_like(source, 255))
    render_preview(background, layout, preview)
    created = _recover_pale_asset_gaps(source_path, preview, layout, assets, project_id, 1)

    assert len(created) == 1
    assert len(layout["elements"]) == 1
    assert created[0]["id"] == "ribbon"
    assert layout["elements"][0]["src"] != f"/media/assets/{project_id}/ribbon.png"
    assert original_path.read_bytes() == original_bytes
    merged = cv2.imread(str(assets / Path(layout["elements"][0]["src"]).name), cv2.IMREAD_UNCHANGED)
    assert merged[5, 5, 3] == 255
    assert merged[35, 5, 3] == 255
    render_preview(background, layout, preview)
    assert np.max(np.abs(cv2.imread(str(preview)).astype(np.int16) - source.astype(np.int16))) < 3
    assert _recover_pale_asset_gaps(source_path, preview, layout, assets, project_id, 1) == []


def test_text_support_asset_erases_original_glyphs(tmp_path):
    from app.services.reconstruction.objectization_audit import _recover_text_support_gaps

    source = np.full((130, 240, 3), 255, np.uint8)
    source[25:105, 25:215] = (246, 247, 251)
    cv2.putText(source, "CARD", (75, 74), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (20, 20, 30), 2)
    missing = np.zeros(source.shape[:2], bool)
    missing[48:85, 65:165] = True
    asset_dir = tmp_path / "assets"
    asset_dir.mkdir()
    layout = {"elements": [
        {"id": "card", "type": "image", "x": 25, "y": 25, "width": 190, "height": 80},
        {"id": "label", "type": "text", "text": "CARD", "x": 65, "y": 48, "width": 100, "height": 37},
    ]}

    created = _recover_text_support_gaps(source, missing, layout, asset_dir, "fixture", 1)

    assert len(created) == 1
    repaired = cv2.imread(str(asset_dir / "initial_page_1_text_support_001.png"), cv2.IMREAD_UNCHANGED)
    assert repaired is not None and repaired.shape[2] == 4
    assert np.min(repaired[:, :, :3]) > 200
    assert created[0]["metadata"]["editableTextIds"] == ["label"]


def test_revision_repairs_only_missing_local_plate(tmp_path):
    store = ProjectStore(tmp_path)
    project_id = store.create("plate-revision")["id"]
    root = tmp_path / project_id
    (root / "backgrounds").mkdir()
    source = np.full((180, 300, 3), 255, np.uint8)
    cv2.rectangle(source, (30, 40), (210, 112), (246, 248, 250), -1)
    cv2.imwrite(str(root / "source.png"), source)
    cv2.imwrite(str(root / "backgrounds" / "page_1.png"), np.full_like(source, 255))
    layout = {"slide": {"width": 300, "height": 180}, "elements": []}
    store.save_slide(project_id, 1, layout)
    render_preview(root / "backgrounds" / "page_1.png", layout, root / "reconstructed_preview.png")
    score = run_visual_qa(root / "source.png", root / "reconstructed_preview.png", root, layout)
    (root / "visual_score.json").write_text(json.dumps(score), encoding="utf-8")
    result = revise_problem_regions(store, project_id, 1)
    assert result["accepted"] is True
    assert any(issue["problem"] == "missingBackplate" for issue in result["issuesBefore"])
    assert not any(issue["problem"] == "missingBackplate" for issue in result["issuesAfter"])
    assert len(store.get_slide(project_id, 1)["elements"]) == 1
    assert np.all(cv2.imread(str(root / "backgrounds" / "page_1.png")) == 255)
    assert cv2.imread(str(root / "reconstructed_preview.png"))[60, 60, 0] < 255


def test_revision_recovers_neutral_plate_on_gradient_page(tmp_path):
    store = ProjectStore(tmp_path)
    project_id = store.create("gradient-plate-revision")["id"]
    root = tmp_path / project_id
    (root / "backgrounds").mkdir()
    source = np.empty((180, 300, 3), np.uint8)
    for y in range(180):
        shade = round(247 + 7 * y / 179)
        source[y, :] = (shade, shade, shade)
    cv2.rectangle(source, (55, 45), (220, 125), (242, 242, 242), -1)
    cv2.imwrite(str(root / "source.png"), source)
    cv2.imwrite(str(root / "backgrounds" / "page_1.png"), np.full_like(source, 255))
    layout = {"slide": {"width": 300, "height": 180}, "elements": []}
    store.save_slide(project_id, 1, layout)
    render_preview(root / "backgrounds" / "page_1.png", layout, root / "reconstructed_preview.png")
    score = run_visual_qa(root / "source.png", root / "reconstructed_preview.png", root, layout)
    (root / "visual_score.json").write_text(json.dumps(score), encoding="utf-8")

    result = revise_problem_regions(store, project_id, 1)

    assert result["accepted"] is True
    assert any(issue["problem"] in {"missingBackplate", "missingVisualObject"} for issue in result["issuesBefore"])
    assert np.max(np.abs(cv2.imread(str(root / "reconstructed_preview.png"))[80, 100].astype(int) - source[80, 100].astype(int))) < 15
    assert any(item["type"] in {"rectangle", "image"} and item["width"] < 200 for item in store.get_slide(project_id, 1)["elements"])


def test_revision_repairs_wrong_colored_local_plate_without_removing_other_visuals(tmp_path, monkeypatch):
    monkeypatch.setattr(renderer, "OUTPUTS_DIR", tmp_path)
    store = ProjectStore(tmp_path)
    project_id = store.create("mismatch-revision")["id"]
    root = tmp_path / project_id
    (root / "backgrounds").mkdir()
    source = np.full((180, 300, 3), 255, np.uint8)
    cv2.rectangle(source, (18, 25), (270, 145), (80, 130, 190), -1)
    cv2.rectangle(source, (45, 48), (180, 112), (242, 246, 250), -1)
    cv2.imwrite(str(root / "source.png"), source)
    cv2.imwrite(str(root / "backgrounds" / "page_1.png"), np.full_like(source, 255))
    layout = {"slide": {"width": 300, "height": 180}, "elements": [
        {"id": "outer", "type": "rectangle", "x": 18, "y": 25, "width": 253, "height": 121, "zIndex": 1,
         "style": {"fill": "#BE8250", "stroke": "#BE8250", "strokeWidth": 0, "opacity": 1}},
        {"id": "wrong_plate", "type": "rectangle", "x": 45, "y": 48, "width": 136, "height": 65, "zIndex": 2,
         "style": {"fill": "#BE8250", "stroke": "#BE8250", "strokeWidth": 0, "opacity": 1}},
    ]}
    store.save_slide(project_id, 1, layout)
    render_preview(root / "backgrounds" / "page_1.png", layout, root / "reconstructed_preview.png")
    before = cv2.imread(str(root / "reconstructed_preview.png"))
    assert np.max(np.abs(before[70, 70].astype(int) - source[70, 70].astype(int))) > 48
    report = audit_objectization(root / "source.png", root / "backgrounds" / "page_1.png", root / "reconstructed_preview.png", layout)
    assert report["visualMismatchRegions"] >= 1
    score = run_visual_qa(root / "source.png", root / "reconstructed_preview.png", root, layout)
    (root / "visual_score.json").write_text(json.dumps(score), encoding="utf-8")
    result = revise_problem_regions(store, project_id, 1)
    assert result["accepted"] is True
    assert not any(issue["problem"] == "visualContentMismatch" for issue in result["issuesAfter"])
    after = cv2.imread(str(root / "reconstructed_preview.png"))
    assert np.max(np.abs(after[70, 70].astype(int) - source[70, 70].astype(int))) < 12
    assert np.max(np.abs(after[35, 35].astype(int) - source[35, 35].astype(int))) < 12
    revised = store.get_slide(project_id, 1)["elements"]
    assert len([item for item in revised if item["type"] == "rectangle" and not (item.get("metadata") or {}).get("suppressed")]) == 1
    assert len([item for item in revised if item["type"] == "image" and not (item.get("metadata") or {}).get("suppressed")]) == 1
    assert result["missingAssetCount"] == 0


def test_revision_recovers_complex_unowned_visual_as_movable_image(tmp_path, monkeypatch):
    monkeypatch.setattr(renderer, "OUTPUTS_DIR", tmp_path)
    store = ProjectStore(tmp_path)
    project_id = store.create("image-revision")["id"]
    root = tmp_path / project_id
    (root / "backgrounds").mkdir()
    source = np.full((180, 300, 3), 255, np.uint8)
    cv2.circle(source, (80, 75), 27, (30, 90, 185), -1, cv2.LINE_AA)
    cv2.circle(source, (80, 75), 8, (220, 220, 40), -1, cv2.LINE_AA)
    cv2.imwrite(str(root / "source.png"), source)
    cv2.imwrite(str(root / "backgrounds" / "page_1.png"), np.full_like(source, 255))
    layout = {"slide": {"width": 300, "height": 180}, "elements": []}
    store.save_slide(project_id, 1, layout)
    render_preview(root / "backgrounds" / "page_1.png", layout, root / "reconstructed_preview.png")
    score = run_visual_qa(root / "source.png", root / "reconstructed_preview.png", root, layout)
    (root / "visual_score.json").write_text(json.dumps(score), encoding="utf-8")
    result = revise_problem_regions(store, project_id, 1)
    assert result["accepted"] is True
    images = [item for item in store.get_slide(project_id, 1)["elements"] if item["type"] == "image"]
    assert len(images) == 1 and images[0]["src"].startswith(f"/media/assets/{project_id}/")
    assert (root / "assets" / Path(images[0]["src"]).name).is_file()
    assert np.all(cv2.imread(str(root / "backgrounds" / "page_1.png")) == 255)
    assert cv2.imread(str(root / "reconstructed_preview.png"))[75, 80, 0] < 255
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    renderer.PPTXRenderer()._add_element(slide, images[0], 0.02, 0.02)
    assert len(slide.shapes) == 1


def test_revision_recovers_missing_thin_decorative_strip(tmp_path, monkeypatch):
    monkeypatch.setattr(renderer, "OUTPUTS_DIR", tmp_path)
    store = ProjectStore(tmp_path)
    project_id = store.create("strip-revision")["id"]
    root = tmp_path / project_id
    (root / "backgrounds").mkdir()
    source = np.full((180, 320, 3), 255, np.uint8)
    source[70, 45:245] = (45, 90, 185)
    cv2.imwrite(str(root / "source.png"), source)
    cv2.imwrite(str(root / "backgrounds" / "page_1.png"), np.full_like(source, 255))
    layout = {"slide": {"width": 320, "height": 180}, "elements": []}
    store.save_slide(project_id, 1, layout)
    render_preview(root / "backgrounds" / "page_1.png", layout, root / "reconstructed_preview.png")
    score = run_visual_qa(root / "source.png", root / "reconstructed_preview.png", root, layout)
    (root / "visual_score.json").write_text(json.dumps(score), encoding="utf-8")

    result = revise_problem_regions(store, project_id, 1)

    assert result["accepted"] is True
    assert any(issue["problem"] == "missingVisualObject" for issue in result["issuesBefore"])
    assert not any(issue["problem"] == "missingVisualObject" for issue in result["issuesAfter"])
    images = [item for item in store.get_slide(project_id, 1)["elements"] if item["type"] == "image"]
    assert len(images) == 1 and images[0]["height"] <= 3
    assert (root / "assets" / Path(images[0]["src"]).name).is_file()
    actual = cv2.imread(str(root / "reconstructed_preview.png"))
    assert np.max(np.abs(actual[70, 100].astype(int) - source[70, 100].astype(int))) < 10


def test_debug_artifact_is_retrievable_from_existing_project_route(tmp_path, monkeypatch):
    store = ProjectStore(tmp_path)
    project_id = store.create("debug-audit")["id"]
    root = tmp_path / project_id
    cv2.imwrite(str(root / "objectization_debug.png"), np.full((20, 30, 3), 255, np.uint8))
    (root / "objectization_audit.json").write_text(json.dumps({"whiteBackground": True}), encoding="utf-8")
    monkeypatch.setattr(api_main, "store", store)
    monkeypatch.setattr(api_main, "OUTPUTS_DIR", tmp_path)
    client = TestClient(api_main.app)
    assert client.get(f"/api/projects/{project_id}/artifacts/objectization_debug.png").status_code == 200
    assert client.get(f"/api/projects/{project_id}/artifacts/objectization_audit.json").json()["whiteBackground"] is True


@pytest.mark.parametrize("variant", ["original", "lightened", "small_badge"])
def test_complex_fixture_white_objectized_end_to_end(tmp_path, variant):
    source = cv2.imread(str(FIXTURE))
    assert source is not None
    if variant == "lightened":
        source = cv2.addWeighted(source, 0.65, np.full_like(source, 255), 0.35, 0)
    elif variant == "small_badge":
        cv2.circle(source, (570, 70), 22, (35, 100, 190), -1, cv2.LINE_AA)
        cv2.circle(source, (570, 70), 7, (255, 255, 255), -1, cv2.LINE_AA)
    source_path = tmp_path / "source.png"
    background = tmp_path / "background.png"
    preview = tmp_path / "preview.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": source.shape[1], "height": source.shape[0]}, "elements": []}
    stats = objectize_on_white(source_path, background, layout, tmp_path / "assets", "fixture", 1)
    assert stats["whiteObjectAssets"] + stats["whiteObjectShapes"] > 0
    for item in layout["elements"]:
        if item.get("type") == "image":
            item["src"] = str(tmp_path / "assets" / Path(item["src"]).name)
    render_preview(background, layout, preview)
    report = audit_objectization(source_path, background, preview, layout, tmp_path / "debug.png")
    assert report["whiteBackground"] is True
    assert report["backgroundResidualRegions"] == 0
    assert report["visualMismatchRegions"] == 0
    assert all(item["owner"] in {"editable_text", "movable_image", "native_shape"} for item in report["ownerRegions"])
    assert (tmp_path / "debug.png").is_file()
