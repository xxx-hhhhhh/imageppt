"""Local paper / OCR alignment regressions, independent of uploaded slides."""
from itertools import pairwise
from pathlib import Path

import cv2
import numpy as np
import pytest
from app.services.pptx.renderer import PPTXRenderer
from app.services.reconstruction.editable_text import (
    measured_surface,
    measured_text_surface,
    restore_text_candidate,
    without_untranscribed_bullet,
)
from app.services.reconstruction.mixed_text import restore_text_candidates
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation


def test_local_ring_avoids_accent_without_relaxing_surface_quality():
    yy,xx=np.mgrid[:48,:360]
    paper=np.stack((230+xx*.02,238+yy*.1,np.full(xx.shape,246)),axis=-1).astype(np.uint8)
    patch=paper.copy()
    patch[-4:]=(180,240,150)  # A neighbouring underline, not a white background.
    assert measured_surface(patch) is None
    surface,evidence=measured_text_surface(patch)
    assert evidence['backgroundSamplingInset']>0
    assert evidence['backgroundError95']<=7
    assert evidence['backgroundInlierRatio']>=.7
    assert np.max(np.abs(surface[6:-6].astype(float)-paper[6:-6]))<2
    assert np.array_equal(patch[-1,0],(180,240,150))


def test_local_ring_does_not_authorize_a_textured_visual():
    patch=np.random.default_rng(13).integers(0,256,(60,300,3),dtype=np.uint8)
    assert measured_text_surface(patch) is None


def test_disconnected_first_glyph_does_not_hide_list_marker():
    ink=np.zeros((40,110),np.uint8)
    cv2.circle(ink,(9,20),4,1,-1)
    # First body glyph has separate upper/lower strokes, neither > .6*28.
    ink[9:16,28:47]=1
    ink[18:31,28:48]=1
    ink[9:31,53:73]=1
    body,removed=without_untranscribed_bullet(ink.astype(bool),'实验内容',28)
    assert removed and not body[:,:20].any()
    assert np.array_equal(body[:,20:],ink[:,20:].astype(bool))
    retained,removed=without_untranscribed_bullet(ink.astype(bool),'•实验内容',28)
    assert not removed and np.array_equal(retained,ink)


def test_wide_glyph_mixed_banner_preserves_characters_and_underline(tmp_path):
    font_path=Path('C:/Windows/Fonts/msyhbd.ttc')
    if not font_path.is_file():
        pytest.skip('Windows installed-font fixture')
    image=Image.new('RGB',(1250,120),'#EAF5FC')
    draw=ImageDraw.Draw(image)
    parts=[('一项说明：',34,21),('这里保留“完整文字”，也保留字号变化后的全部正文。',26,30)]
    cursor=24
    boxes=[]
    for text,size,y in parts:
        font=ImageFont.truetype(str(font_path),size)
        draw.text((cursor,y),text,font=font,fill='#153D83')
        boxes.append(draw.textbbox((cursor,y),text,font=font))
        cursor+=round(draw.textlength(text,font=font))+15
    x,y,right,bottom=min(b[0] for b in boxes),min(b[1] for b in boxes),max(b[2] for b in boxes),max(b[3] for b in boxes)
    draw.rectangle((x,bottom+3,right,bottom+6),fill='#83E2CE')
    source=cv2.cvtColor(np.asarray(image),cv2.COLOR_RGB2BGR)
    original=source.copy()
    item={'id':'banner','type':'text','text':''.join(p[0] for p in parts),'confidence':.99,
          'x':x,'y':y,'width':right-x,'height':bottom-y,
          'style':{'color':'#153D83','fontFamily':'Microsoft YaHei'}}
    diagnostics={}
    candidates=restore_text_candidates(source,item,diagnostics=diagnostics)
    assert ''.join(e['text'] for e,_,_ in candidates)==item['text'],diagnostics
    assert len(candidates)>1
    assert np.array_equal(source,original)
    mask=np.logical_or.reduce([mask for _,mask,_ in candidates])
    assert not mask[bottom+3:bottom+7].any()
    spans=[e['metadata']['sourceCharRange'] for e,_,_ in candidates]
    assert spans[0][0]==0 and spans[-1][1]==len(item['text'])
    assert all(a[1]==b[0] for a,b in pairwise(spans))
    layout={'slide':{'width':1250,'height':120},'elements':[e for e,_,_ in candidates]}
    pptx,validation=PPTXRenderer().render_project('banner',[layout],output_dir=tmp_path)
    assert validation['valid']
    shapes=Presentation(pptx).slides[0].shapes
    assert all(shape.has_text_frame for shape in shapes)
    assert ''.join(shape.text for shape in shapes)==item['text']


def test_foreground_uses_solid_ink_not_pale_fringe_of_approximate_ocr_color():
    font_path=Path('C:/Windows/Fonts/msyhbd.ttc')
    if not font_path.is_file():
        pytest.skip('Windows installed-font fixture')
    image=Image.new('RGB',(400,100),'#EAF5FC')
    draw=ImageDraw.Draw(image)
    font=ImageFont.truetype(str(font_path),30)
    text='强调文字颜色'
    draw.text((24,20),text,font=font,fill='#0054C7')
    x,y,right,bottom=draw.textbbox((24,20),text,font=font)
    item={'id':'color','text':text,'confidence':.99,'x':x,'y':y,
          'width':right-x,'height':bottom-y,'style':{'color':'#032A78'}}
    source=cv2.cvtColor(np.asarray(image),cv2.COLOR_RGB2BGR)
    candidate=restore_text_candidate(source,item)
    assert candidate is not None
    actual=np.array(list(bytes.fromhex(candidate[0]['style']['color'][1:])))
    assert np.max(np.abs(actual.astype(float)-[0,84,199]))<10


def test_adjacent_runs_use_one_verified_paper_field():
    font_path=Path('C:/Windows/Fonts/msyhbd.ttc')
    if not font_path.is_file():
        pytest.skip('Windows installed-font fixture')
    yy,xx=np.mgrid[:100,:300]
    paper=np.stack((230+xx*.03,245+yy*.01,np.full(xx.shape,250)),axis=-1).astype(np.uint8)
    image=Image.fromarray(cv2.cvtColor(paper,cv2.COLOR_BGR2RGB))
    draw=ImageDraw.Draw(image)
    font=ImageFont.truetype(str(font_path),26)
    draw.text((20,20),'相邻文字',font=font,fill='#17365D')
    source=cv2.cvtColor(np.asarray(image),cv2.COLOR_RGB2BGR)
    candidates=[]
    for index,text in enumerate(('相邻','文字')):
        x,y,right,bottom=draw.textbbox((20+52*index,20),text,font=font)
        item={'id':f'run{index}','text':text,'confidence':.99,'x':x,'y':y,
              'width':right-x,'height':bottom-y,'style':{'color':'#17365D'}}
        candidate=restore_text_candidate(source,item,paper_reference=(0,0,paper,{'backgroundError95':0}))
        assert candidate is not None
        assert candidate[0]['metadata']['textErasureEvidence']['sharedOcrLinePaper']
        candidates.append(candidate)
    overlap=candidates[0][1]&candidates[1][1]
    assert overlap.any()
    assert np.array_equal(candidates[0][2][overlap],candidates[1][2][overlap])
