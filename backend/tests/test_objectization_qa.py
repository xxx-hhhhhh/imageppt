from __future__ import annotations

from pathlib import Path

import cv2
import numpy as np
from pptx import Presentation

from app.services.pptx.renderer import PPTXRenderer
from app.services.reconstruction.objectization_qa import repair_objectized_modules
from app.services.reconstruction.white_objectization import objectize_on_white
from app.services.visual_qa.analyzer import render_preview


FIXTURE = Path(__file__).parent / "fixtures" / "complex_modules.png"


def test_objectization_qa_restores_missing_plate_and_repairs_square_badge(tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.pptx.renderer.OUTPUTS_DIR", tmp_path)
    project_id = "complex-modules"
    asset_dir = tmp_path / project_id / "assets"
    asset_dir.mkdir(parents=True)
    source = cv2.imread(str(FIXTURE))
    assert source is not None and source.shape[:2] == (360, 640)
    # Simulate an upstream whole-image badge crop with an opaque square exterior.
    badge = source[72:136, 60:124]
    cv2.imwrite(str(asset_dir / "badge.png"), badge)
    layout = {"slide": {"width": 640, "height": 360}, "elements": [
        {"id": "badge", "type": "image", "x": 60, "y": 72, "width": 64, "height": 64, "zIndex": 12, "groupId": "module_a", "role": "badge", "src": f"/media/assets/{project_id}/badge.png", "metadata": {"reconstructionStrategy": "local_image"}},
        {"id": "badge_geometry", "type": "ellipse", "x": 60, "y": 72, "width": 64, "height": 64, "groupId": "module_a", "metadata": {"suppressed": True}},
        {"id": "title_a", "type": "text", "x": 137, "y": 89, "width": 145, "height": 27, "zIndex": 20, "groupId": "module_a", "text": "MODULE A", "style": {"fontSize": 19, "color": "#1F273E"}},
        {"id": "body_a", "type": "text", "x": 61, "y": 145, "width": 165, "height": 28, "zIndex": 20, "groupId": "module_a", "text": "Editable text", "style": {"fontSize": 18, "color": "#4B5267"}},
        {"id": "title_b", "type": "text", "x": 363, "y": 94, "width": 160, "height": 28, "zIndex": 20, "groupId": "module_b", "text": "MODULE B", "style": {"fontSize": 20, "color": "#273044"}},
        {"id": "body_b", "type": "text", "x": 371, "y": 213, "width": 120, "height": 26, "zIndex": 20, "groupId": "module_b", "text": "Details", "metadata": {"rawOCRBBox": [371, 213, 491, 239]}, "style": {"fontSize": 18, "color": "#495466"}},
    ]}
    background = tmp_path / "background.png"
    objectize_on_white(FIXTURE, background, layout, asset_dir, project_id, 1)
    # Exercise the QA recovery path after an incomplete module split.
    layout["elements"] = [item for item in layout["elements"] if not (item.get("groupId") == "module_a" and (item.get("metadata") or {}).get("layerRole") == "container")]
    before_path = tmp_path / "before.png"
    render_preview(background, layout, before_path)
    report = repair_objectized_modules(FIXTURE, layout, asset_dir, project_id)
    after_path = tmp_path / "after.png"
    render_preview(background, layout, after_path)

    assert report["checkedModules"] >= 2
    assert report["missingBackplates"] >= 1
    assert report["recoveredBackplates"] + report["reboundBackplates"] >= 1
    assert report["squareCutouts"] == report["repairedCutouts"] == 1
    plate = next(item for item in layout["elements"] if item.get("groupId") == "module_a" and (item.get("metadata") or {}).get("layerRole") == "container")
    assert plate["zIndex"] < 12
    assert "badge" in plate["metadata"]["moduleMemberIds"]
    badge_item = next(item for item in layout["elements"] if item.get("id") == "badge")
    repaired = cv2.imread(str(asset_dir / Path(badge_item["src"]).name), cv2.IMREAD_UNCHANGED)
    assert repaired.shape[2] == 4 and repaired[0, 0, 3] == 0 and repaired[32, 32, 3] == 255
    assert cv2.imread(str(asset_dir / "badge.png"), cv2.IMREAD_UNCHANGED).shape[2] == 3
    assert all(item.get("type") == "text" for item in layout["elements"] if item.get("id", "").startswith(("title_", "body_")))
    before = cv2.imread(str(before_path)).astype(np.int16)
    after = cv2.imread(str(after_path)).astype(np.int16)
    # Sample the unobstructed plate corner: recovering it must improve fidelity.
    original = source.astype(np.int16)
    assert np.mean(np.abs(after[59:70, 43:55] - original[59:70, 43:55])) <= np.mean(np.abs(before[59:70, 43:55] - original[59:70, 43:55]))
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    PPTXRenderer()._add_element(slide, plate, 0.02, 0.02)
    PPTXRenderer()._add_element(slide, badge_item, 0.02, 0.02)
    assert len(slide.shapes) == 2  # Both remain movable PPT objects.


def test_objectization_qa_does_not_turn_page_background_into_a_plate(tmp_path):
    source = np.full((240, 400, 3), 246, np.uint8)
    path = tmp_path / "plain.png"
    cv2.imwrite(str(path), source)
    layout = {"slide": {"width": 400, "height": 240}, "elements": [{"id": "text", "type": "text", "x": 90, "y": 80, "width": 100, "height": 30, "groupId": "module", "text": "Hello"}]}
    report = repair_objectized_modules(path, layout, tmp_path, "plain")
    assert report["recoveredBackplates"] == 0
    assert len(layout["elements"]) == 1


def test_textured_local_plate_falls_back_to_image_plus_editable_text(tmp_path):
    source = np.full((220, 380, 3), 250, np.uint8)
    source[35:145, 42:255] = (232, 237, 243)
    for x in range(42, 255, 7):
        source[35:145, x:x + 3] = (212, 225, 238)
    cv2.circle(source, (75, 78), 13, (30, 100, 190), -1)
    path = tmp_path / "textured.png"
    cv2.imwrite(str(path), source)
    layout = {"slide": {"width": 380, "height": 220}, "elements": [
        {"id": "icon", "type": "ellipse", "x": 62, "y": 65, "width": 26, "height": 26, "zIndex": 10, "groupId": "card", "style": {"fill": "#BE641E"}},
        {"id": "title", "type": "text", "x": 105, "y": 62, "width": 90, "height": 25, "zIndex": 20, "groupId": "card", "text": "Title", "metadata": {"rawOCRBBox": [105, 62, 195, 87]}},
    ]}
    report = repair_objectized_modules(path, layout, tmp_path, "textured")
    assert report["moduleImageFallbacks"] == 1
    plate = next(item for item in layout["elements"] if (item.get("metadata") or {}).get("layerRole") == "container")
    assert plate["type"] == "image" and plate["groupId"] == "card"
    assert plate["metadata"]["moduleMemberIds"] == ["icon", "title"]
    assert plate["metadata"]["foregroundCleaned"] is True
    assert not layout["elements"][0].get("metadata", {}).get("suppressed")
    clean_plate = cv2.imread(str(tmp_path / Path(plate["src"]).name), cv2.IMREAD_UNCHANGED)
    assert np.max(np.abs(clean_plate[78 - plate["y"], 75 - plate["x"], :3].astype(int)
                         - np.array([30, 100, 190]))) > 50
    assert layout["elements"][1]["type"] == "text" and not layout["elements"][1]["metadata"].get("suppressed")


def test_textured_backplate_keeps_image_icon_independently_movable(tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.pptx.renderer.OUTPUTS_DIR", tmp_path)
    project_id = "textured-module"
    asset_dir = tmp_path / project_id / "assets"
    asset_dir.mkdir(parents=True)
    source = np.full((220, 380, 3), 250, np.uint8)
    source[35:145, 42:255] = (232, 237, 243)
    for x in range(42, 255, 7):
        source[35:145, x:x + 3] = (212, 225, 238)
    cv2.circle(source, (75, 78), 13, (30, 100, 190), -1)
    source_path = tmp_path / "textured-icon.png"
    cv2.imwrite(str(source_path), source)
    alpha = np.zeros((30, 30), np.uint8)
    cv2.circle(alpha, (15, 15), 13, 255, -1)
    icon = np.dstack((source[63:93, 60:90], alpha))
    cv2.imwrite(str(asset_dir / "icon.png"), icon)
    layout = {"elements": [
        {"id": "icon", "type": "image", "x": 60, "y": 63, "width": 30, "height": 30,
         "zIndex": 10, "groupId": "card", "src": f"/media/assets/{project_id}/icon.png"},
    ]}

    report = repair_objectized_modules(source_path, layout, asset_dir, project_id)

    assert report["moduleImageFallbacks"] == 1
    assert not layout["elements"][0].get("metadata", {}).get("suppressed")
    plate = next(item for item in layout["elements"] if (item.get("metadata") or {}).get("layerRole") == "container")
    assert plate["metadata"]["foregroundCleaned"] is True
    cleaned = cv2.imread(str(asset_dir / Path(plate["src"]).name), cv2.IMREAD_UNCHANGED)
    assert np.max(np.abs(cleaned[78 - plate["y"], 75 - plate["x"], :3].astype(int)
                         - np.array([30, 100, 190]))) > 50
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for item in layout["elements"]:
        PPTXRenderer()._add_element(slide, item, 0.02, 0.02)
    assert len(slide.shapes) == 2 and all(shape.shape_type == 13 for shape in slide.shapes)


def test_round_cutout_quality_replaces_opaque_square_asset(tmp_path):
    source = np.full((100, 100, 3), 255, np.uint8)
    cv2.circle(source, (50, 50), 26, (40, 90, 180), -1)
    source_path = tmp_path / "source.png"
    cv2.imwrite(str(source_path), source)
    asset_dir = tmp_path / "assets"
    asset_dir.mkdir()
    cv2.imwrite(str(asset_dir / "badge.png"), source[20:80, 20:80])
    layout = {"elements": [{"id": "badge", "type": "image", "x": 20, "y": 20, "width": 60, "height": 60,
                            "src": "/media/assets/demo/badge.png", "metadata": {"reconstructionStrategySource": "round_contour"}}]}
    report = repair_objectized_modules(source_path, layout, asset_dir, "demo")
    assert report["roundCutoutsChecked"] == 1
    assert report["roundCutoutIssues"] == 0
    assert report["repairedCutouts"] == 1
    repaired = cv2.imread(str(asset_dir / Path(layout["elements"][0]["src"]).name), cv2.IMREAD_UNCHANGED)
    assert repaired.shape[2] == 4 and repaired[0, 0, 3] == 0 and repaired[30, 30, 3] == 255
