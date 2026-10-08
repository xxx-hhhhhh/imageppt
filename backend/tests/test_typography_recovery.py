"""Installed-font and mixed-size regressions, independent of any source slide."""
from pathlib import Path

import numpy as np
import pytest
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.oxml.ns import qn

from app.services.pptx.renderer import PPTXRenderer
from app.services.reconstruction.editable_text import restore_text_candidate
from app.services.reconstruction.mixed_text import mixed_spans, restore_text_candidates
from app.utils.image_io import imread


def face(name, size):
    path=Path('C:/Windows/Fonts')/name
    if not path.is_file():
        pytest.skip('Windows installed-font fixture')
    return ImageFont.truetype(str(path),size)


@pytest.mark.parametrize('filename,family,weight',[
    ('simsun.ttc','SimSun',400),('msyhbd.ttc','Microsoft YaHei',700),
    ('simkai.ttf','KaiTi',400),
])
def test_font_ranking_preserves_real_glyph_family(tmp_path,filename,family,weight):
    image=Image.new('RGB',(440,110),'#EDF3F8')
    draw=ImageDraw.Draw(image)
    font=face(filename,32)
    text='字形匹配与编辑'
    draw.text((20,20),text,font=font,fill='#17365D')
    box=draw.textbbox((20,20),text,font=font)
    path=tmp_path/'font.png'
    image.save(path)
    item={'id':'ocr','text':text,'type':'text','confidence':.99,
          'x':box[0],'y':box[1],'width':box[2]-box[0],'height':box[3]-box[1],
          'style':{'color':'#17365D','fontFamily':'Microsoft YaHei'}}
    candidate=restore_text_candidate(imread(path),item)
    assert candidate is not None
    assert candidate[0]['style']['fontFamily']==family
    assert candidate[0]['style']['fontWeight']==weight
    assert candidate[0]['metadata']['fontMatchingVersion']==2


def test_mixed_size_preserves_every_ocr_character_and_tail(tmp_path):
    image=Image.new('RGB',(580,130),'#FCF2E4')
    draw=ImageDraw.Draw(image)
    chunks=[('9','times.ttf',54,40),('项','simsun.ttc',30,54),
            ('，','simsun.ttc',30,54),('服务教师','simsun.ttc',30,54),
            ('204','times.ttf',42,47),('人','simsun.ttc',30,54)]
    cursor=25
    boxes=[]
    for text,filename,size,y in chunks:
        font=face(filename,size)
        draw.text((cursor,y),text,font=font,fill='#17365D')
        boxes.append(draw.textbbox((cursor,y),text,font=font))
        cursor+=round(draw.textlength(text,font=font))+5
    text=''.join(c[0] for c in chunks)
    box=(min(b[0] for b in boxes),min(b[1] for b in boxes),max(b[2] for b in boxes),max(b[3] for b in boxes))
    path=tmp_path/'mixed.png'
    image.save(path)
    item={'id':'ocr','text':text,'type':'text','confidence':.99,'x':box[0],'y':box[1],
          'width':box[2]-box[0],'height':box[3]-box[1], 'style':{'color':'#17365D'}}
    source=imread(path)
    spans=mixed_spans(source,item)
    assert ''.join(s['text'] for s in spans)==text
    candidates=restore_text_candidates(source,item)
    assert ''.join(c[0]['text'] for c in candidates)==text
    assert max(c[0]['style']['fontSize'] for c in candidates)>min(c[0]['style']['fontSize'] for c in candidates)*1.3
    assert all(c[0]['metadata']['sourceOcrIds']==['ocr'] for c in candidates)
    # Masks must be exclusive: a partial span never erases another one.
    masks=np.stack([c[1] for c in candidates])
    assert masks.sum(axis=0).max()==1


def test_ppt_sets_east_asian_font_and_pixel_tracking(tmp_path):
    item={'id':'t','type':'text','text':'中文字体，完整正文','x':20,'y':25,'width':400,'height':60,
          'style':{'fontFamily':'SimSun','fontSize':32,'fontWeight':700,'letterSpacing':1.25},
          'metadata':{'exclusiveFontVerified':True}}
    layout={'slide':{'width':800,'height':450},'elements':[item]}
    path,_=PPTXRenderer().render_project('font',[layout],output_dir=tmp_path)
    shape=Presentation(path).slides[0].shapes[0]
    run=shape.text_frame.paragraphs[0].runs[0]
    assert run.text==item['text'] and run.font.bold
    for tag in ('latin','ea','cs'):
        assert run.font._element.find(qn('a:'+tag)).get('typeface')=='SimSun'
    assert int(run.font._element.get('spc'))==round(1.25*13.333/800*72*100)


def test_manual_family_edit_overrides_old_measurement_file():
    from app.services.typography.font_raster import paint_measured_text
    face('simsun.ttc',32)
    style={'fontFamily':'SimSun','fontFile':'msyhbd.ttc','fontSize':32,'fontWeight':400,'color':'#17365D'}
    item={'text':'字体编辑','x':10,'y':10,'style':style}
    first=Image.new('RGBA',(240,100))
    paint_measured_text(first,item)
    style['fontFile']='simsun.ttc'
    same=Image.new('RGBA',(240,100))
    paint_measured_text(same,item)
    assert np.array_equal(np.asarray(first),np.asarray(same))
    style['fontFamily']='Microsoft YaHei'
    changed=Image.new('RGBA',(240,100))
    paint_measured_text(changed,item)
    assert not np.array_equal(np.asarray(first),np.asarray(changed))
