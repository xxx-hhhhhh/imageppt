from __future__ import annotations

import copy
from pathlib import Path

import cv2
import numpy as np
import pytest
from app.services.inpainting.service import InpaintingService
from app.services.ocr.provider import OCRResult
from app.services.reconstruction.object_first import (
    ObjectFirstBuilder,
    important_mask,
    read_image,
    write_image,
)
from app.services.reconstruction.revisions import RevisionManager
from app.services.scene.ownership import OwnershipLedger, canonicalize, resolve_asset
from app.services.segmentation.segmentation_provider import OpenCVSegmentationProvider
from app.services.visual_qa.preservation import revision_gate
from PIL import Image, ImageDraw


def build(tmp_path: Path, regions=None, scene=None, segments=True):
    source = tmp_path / "源图.png"
    image = Image.new("RGB", (480, 280), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((20, 60, 210, 200), fill="#F3F7FA")
    draw.rectangle((235, 65, 435, 100), fill="#104D90")
    draw.ellipse((260, 125, 360, 225), fill="#C00000")
    draw.line((280, 175, 340, 175), fill="white", width=8)
    draw.line((310, 145, 310, 205), fill="white", width=8)
    draw.text((28, 80), "Editable text", fill="#111111")
    image.save(source)
    regions = regions if regions is not None else [OCRResult("Editable text", [27, 78, 110, 99], 0.99, {"fontSize": 12, "color": "#111111"})]
    output = tmp_path / "outputs"
    asset_dir = output / "test" / "assets"
    provider = OpenCVSegmentationProvider()
    segmentation = provider.segment(source, asset_dir, "test") if segments else []
    layout = ObjectFirstBuilder(output).build(source, {}, scene or {}, regions, segmentation, "test", 1, output / "test" / "backgrounds" / "page.png", InpaintingService("opencv"))
    return source, output, layout


def test_no_owner_no_deletion(tmp_path):
    ledger = OwnershipLedger((10, 10), tmp_path)
    requested = np.full((10, 10), 255, np.uint8)
    assert not ledger.authorize(requested).any()
    with pytest.raises(ValueError, match="no reliable replacement"):
        ledger.claim({"id": "missing", "owner": "movable_image", "x": 0, "y": 0, "width": 10, "height": 10, "src": "missing.png"}, requested)
    with pytest.raises(ValueError):
        ledger.claim({"id": "ignored", "owner": "intentional_ignore", "x": 0, "y": 0, "width": 10, "height": 10}, requested)


@pytest.mark.parametrize("segments", [True, False])
def test_every_visual_pixel_has_a_live_owner_and_pale_plate_survives(tmp_path, segments):
    source, output, layout = build(tmp_path, segments=segments)
    covered = np.zeros((280, 480), bool)
    for item in layout["elements"]:
        mask = resolve_asset(item.get("mask"), output)
        if mask:
            covered |= cv2.imdecode(np.fromfile(mask, dtype=np.uint8), cv2.IMREAD_GRAYSCALE) > 0
    assert np.all(covered[important_mask(read_image(source)) > 0])
    assert layout["metadata"]["unownedVisualPixels"] == 0
    assert any(item["owner"] == "native_shape" and item["style"]["fill"] == "#F3F7FA" for item in layout["elements"])
    manager = RevisionManager(output, "test", 1)
    background = resolve_asset(layout["backgroundUrl"], output)
    _, qa = manager.snapshot(layout, source, background, "committed")
    assert qa["missingAssetCount"] == 0
    assert qa["visualAreaPreserved"] > 0.97
    assert qa["objectExtractionCoverage"] == 1


def test_badge_is_transparent_real_silhouette_not_a_white_square(tmp_path):
    _, output, layout = build(tmp_path)
    image = next(item for item in layout["elements"] if item["owner"] == "movable_image" and 250 < item["x"] < 270 and item["y"] > 110)
    with Image.open(resolve_asset(image["src"], output)) as asset:
        rgba = np.asarray(asset.convert("RGBA"))
    assert rgba[0, 0, 3] == 0
    assert rgba[rgba.shape[0]//2, rgba.shape[1]//2, 3] == 255
    assert rgba[rgba.shape[0]//2, rgba.shape[1]//2, :3].min() > 240  # white artwork retained


def test_vlm_cannot_rewrite_or_suppress_plain_ocr(tmp_path):
    _, _, layout = build(tmp_path, scene={"elements": [{"id": "text_001", "text": "hallucinated", "role": "card_title", "groupId": "module", "metadata": {"suppressRender": True, "ownedBy": "nonexistent", "reconstructionStrategy": "group"}}]})
    text = next(item for item in layout["elements"] if item["type"] == "text")
    assert text["text"] == "Editable text"
    assert text["owner"] == "editable_text"
    assert text["role"] == "card_title"
    assert text["metadata"].get("ownedBy") is None


def test_low_confidence_text_is_preserved_as_movable_visual(tmp_path):
    _, _, layout = build(tmp_path, regions=[OCRResult("uncertain", [27, 78, 110, 99], 0.2, {})])
    assert not any(item["type"] == "text" for item in layout["elements"])
    assert layout["metadata"]["unownedVisualPixels"] == 0


def test_protected_pixels_survive_a_misbehaving_inpaint_model(tmp_path):
    source, _, _layout = build(tmp_path)
    source_image = read_image(source)
    ledger = OwnershipLedger(source_image.shape[:2], tmp_path)
    mask = np.zeros(source_image.shape[:2], np.uint8)
    mask[82:88, 30:80] = 255
    item = {"id": "text", "type": "text", "owner": "editable_text", "x": 27, "y": 78, "width": 83, "height": 21, "text": "editable"}
    ledger.claim(item, mask)
    service = InpaintingService("opencv")
    class BleedingModel:
        name = "lama-test-double"
        def inpaint(self, image, mask, output):
            write_image(output, np.full(source_image.shape, 255, np.uint8))
    service.provider = BleedingModel()
    destination = tmp_path / "patched.png"
    service.restore_owned_text(source, mask, cv2.bitwise_not(mask), ledger, destination)
    patched = read_image(destination)
    assert np.array_equal(patched[mask == 0], source_image[mask == 0])
    assert service.last_strategies[0]["passes"] == 1


def test_legacy_bbox_cleanup_without_ownership_is_denied(tmp_path):
    source, _, _ = build(tmp_path)
    target = tmp_path / "legacy.png"
    service = InpaintingService("opencv")
    service.restore_background(source, [], target)
    assert source.read_bytes() == target.read_bytes()
    assert service.reclean_background(target, [[0, 0, 400, 250]]) == 0


def test_two_rounds_preserve_assets_background_scene_and_qa(tmp_path):
    source, output, layout = build(tmp_path)
    background = resolve_asset(layout["backgroundUrl"], output)
    manager = RevisionManager(output, "test", 1)
    _, qa = manager.snapshot(layout, source, background, "committed")
    baseline_elements = copy.deepcopy(layout["elements"])
    background_bytes = background.read_bytes()
    for _ in range(2):
        layout, new_qa, decision = manager.revise(layout, source, background, {"issues": []}, qa)
        assert decision["accepted"] is False
        assert "no_measured_improvement" in decision["reasons"]
        assert layout["elements"] == baseline_elements
        assert new_qa["visualAreaPreserved"] == qa["visualAreaPreserved"]
        assert background.read_bytes() == background_bytes
    for revision in manager.root.glob("r_*"):
        assert all((revision / name).exists() for name in ("scene.json", "assets", "background.png", "preview.png", "qa.json", "decision.json"))


@pytest.mark.parametrize("key,value", [("visualAreaPreserved", 0.5), ("missingAssetCount", 1), ("missingVisualCount", 1), ("imageAssetCount", 1), ("backgroundWhiteArea", 1.0)])
def test_revision_gate_rejects_destructive_regression(key, value):
    baseline = {"overall": 0.8, "visualAreaPreserved": 0.98, "imageAssetCount": 10, "backgroundWhiteArea": 0.8}
    candidate = {**baseline, "overall": 0.99, key: value}
    assert revision_gate(baseline, candidate, {"elements": []}, {"elements": []})[0] is False


def test_manual_edit_refreshes_derived_bbox_without_second_graph(tmp_path):
    _, _, layout = build(tmp_path)
    item = next(item for item in layout["elements"] if item["type"] == "text")
    item["x"] = 52
    canonicalize(layout)
    assert item["bbox"]["left"] == 52
    assert "nodes" not in layout
    item["text"] = "manually edited"
    canonicalize(layout)
    assert "\n".join(line["text"] for line in item["lines"]) == "manually edited"


def test_export_uses_current_text_not_stale_ocr_lines(tmp_path, monkeypatch):
    from app.services.pptx import renderer as renderer_module
    from pptx import Presentation
    monkeypatch.setattr(renderer_module, "OUTPUTS_DIR", tmp_path)
    layout = {"slide": {"width": 480, "height": 280}, "elements": [{"id": "edit", "type": "text", "owner": "editable_text", "text": "Edited now", "lines": [{"text": "Stale OCR"}], "x": 10, "y": 10, "width": 300, "height": 50, "style": {"fontSize": 24}}]}
    path, report = renderer_module.PPTXRenderer().render_project("edit", [layout])
    assert Presentation(path).slides[0].shapes[0].text == "Edited now"
    assert report["valid"] is True


def test_validation_counts_reopened_text_instead_of_assuming_coverage(tmp_path):
    from app.services.validation.service import validate_pptx
    from pptx import Presentation
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    path = tmp_path / "missing-text.pptx"
    deck.save(path)
    layout = {"slide": {"width": 480, "height": 280}, "elements": [{"id": "lost", "type": "text", "text": "must be native", "x": 10, "y": 10, "width": 100, "height": 20, "rotation": 0}]}
    report = validate_pptx(path, [layout])
    assert report["valid"] is False
    assert report["ocrCoverage"][0]["exportedText"] == 0
    assert report["ocrCoverage"][0]["ratio"] == 0


def test_missing_owned_image_cancels_export_without_silent_omission(tmp_path):
    from app.services.pptx.renderer import PPTXRenderer
    layout = {"slide": {"width": 480, "height": 280}, "elements": [{"id": "visual", "type": "image", "owner": "movable_image", "src": str(tmp_path / "missing.png"), "x": 10, "y": 10, "width": 100, "height": 80}]}
    with pytest.raises(ValueError, match="Required scene asset missing"):
        PPTXRenderer().render_project("unused", [layout])
