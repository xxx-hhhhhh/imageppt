from __future__ import annotations

from pathlib import Path

import cv2
import numpy as np
from pptx import Presentation

from app.services.reconstruction.white_objectization import layer_objectized_elements, objectize_on_white
from app.services.reconstruction.objectization_qa import repair_objectized_modules
from app.services.reconstruction.objectization_audit import audit_objectization
from app.services.reconstruction.text_erasure import erase_editable_text_sources
from app.services.pptx.renderer import PPTXRenderer
from app.services.visual_qa.analyzer import render_preview


def test_dark_flat_page_is_a_movable_shape_over_white_base(tmp_path):
    source = np.full((180, 300, 3), (52, 31, 21), np.uint8)
    cv2.putText(source, "Dark", (35, 75), cv2.FONT_HERSHEY_SIMPLEX, 1, (245, 245, 245), 2)
    source_path = tmp_path / "source.png"
    background_path = tmp_path / "backgrounds" / "page_1.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 300, "height": 180}, "elements": [
        {"id": "title", "type": "text", "x": 35, "y": 45, "width": 100, "height": 45,
         "text": "Dark", "zIndex": 1, "style": {"color": "#f5f5f5", "fontSize": 32},
         "metadata": {"rawOCRBBox": [35, 45, 135, 90]}}]}
    objectize_on_white(source_path, background_path, layout, tmp_path / "assets", "demo", 1)
    assert np.all(cv2.imread(str(background_path)) == 255)
    surfaces = [item for item in layout["elements"] if (item.get("metadata") or {}).get("pageSurface")]
    assert len(surfaces) == 1
    assert surfaces[0]["type"] == "rectangle"
    assert surfaces[0]["style"]["fill"].lower() == "#151f34"
    assert surfaces[0]["zIndex"] < layout["elements"][0]["zIndex"]
    assert layout["elements"][0]["type"] == "text"
    preview_path = tmp_path / "preview.png"
    render_preview(background_path, layout, preview_path)
    preview = cv2.imread(str(preview_path))
    assert np.array_equal(preview[150, 150], np.array([52, 31, 21]))
    assert audit_objectization(source_path, background_path, preview_path, layout)["pageSurfaceMismatchPixels"] < 1000


def test_wide_pale_support_survives_white_objectization(tmp_path, monkeypatch):
    from app.services.pptx import renderer
    monkeypatch.setattr(renderer, "OUTPUTS_DIR", tmp_path)
    source = np.full((240, 420, 3), 250, np.uint8)
    cv2.rectangle(source, (25, 75), (395, 110), (248, 248, 248), -1)
    source_path = tmp_path / "source.png"
    background_path = tmp_path / "backgrounds" / "page_1.png"
    preview_path = tmp_path / "preview.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 420, "height": 240}, "elements": []}

    objectize_on_white(source_path, background_path, layout, tmp_path / "demo" / "assets", "demo", 1)
    render_preview(background_path, layout, preview_path)
    preview = cv2.imread(str(preview_path))

    assert np.array_equal(preview[90, 200], source[90, 200])
    assert any(item["type"] in {"rectangle", "roundedRectangle", "image"}
               and item["x"] <= 25 and item["x"] + item["width"] >= 395 for item in layout["elements"])
    assert not audit_objectization(source_path, background_path, preview_path, layout)["issues"]
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for item in layout["elements"]:
        PPTXRenderer()._add_element(slide, item, 0.02, 0.02)
    assert len(slide.shapes) == len(layout["elements"])


def test_large_bounded_panel_is_native_and_keeps_inner_visual_movable(tmp_path, monkeypatch):
    from app.services.pptx import renderer
    monkeypatch.setattr(renderer, "OUTPUTS_DIR", tmp_path)
    source = np.full((240, 420, 3), 255, np.uint8)
    cv2.rectangle(source, (5, 5), (414, 234), (65, 105, 160), -1)
    cv2.rectangle(source, (200, 120), (390, 220), (35, 55, 80), -1)
    cv2.circle(source, (300, 145), 18, (30, 210, 230), -1)
    cv2.putText(source, "PANEL", (80, 100), cv2.FONT_HERSHEY_SIMPLEX, 1, (255, 255, 255), 2)
    source_path = tmp_path / "source.png"
    background_path = tmp_path / "backgrounds" / "page_1.png"
    preview_path = tmp_path / "preview.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 420, "height": 240}, "elements": [
        {"id": "title", "type": "text", "x": 75, "y": 65, "width": 150, "height": 45,
         "zIndex": 20, "text": "PANEL", "style": {"color": "#FFFFFF", "fontSize": 32},
         "metadata": {"rawOCRBBox": [75, 65, 225, 110]}}]}

    objectize_on_white(source_path, background_path, layout, tmp_path / "demo" / "assets", "demo", 1)
    render_preview(background_path, layout, preview_path)
    preview = cv2.imread(str(preview_path))
    panels = [item for item in layout["elements"] if (item.get("metadata") or {}).get("reconstructionStrategySource") == "large_inset_panel"]
    assert len(panels) == 1
    assert panels[0]["type"] == "rectangle"
    assert panels[0]["x"] == 5 and panels[0]["y"] == 5
    assert panels[0]["width"] == 410 and panels[0]["height"] == 230
    assert any(item["type"] == "image" and item["x"] <= 300 <= item["x"] + item["width"] for item in layout["elements"])
    assert all(item["width"] * item["height"] < 420 * 240 * 0.5
               for item in layout["elements"] if item["type"] == "image")
    assert np.array_equal(preview[50, 300], source[50, 300])
    assert np.array_equal(preview[200, 300], source[200, 300])
    assert np.array_equal(preview[145, 300], source[145, 300])
    assert np.all(cv2.imread(str(background_path)) == 255)
    audit = audit_objectization(source_path, background_path, preview_path, layout)
    assert audit["missingVisualPixels"] < 100


def test_white_surface_extracts_unowned_visual_and_keeps_text_editable(tmp_path):
    source = np.full((240, 400, 3), 255, np.uint8)
    cv2.rectangle(source, (30, 40), (105, 115), (30, 70, 180), -1)
    cv2.putText(source, "Title", (180, 75), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 0, 0), 2)
    source_path = tmp_path / "source.png"
    background_path = tmp_path / "backgrounds" / "page_1.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 400, "height": 240}, "elements": [{"id": "title", "type": "text", "x": 175, "y": 45, "width": 100, "height": 40, "text": "Title", "metadata": {"rawOCRBBox": [175, 45, 275, 85]}}]}

    stats = objectize_on_white(source_path, background_path, layout, tmp_path / "assets", "demo", 1)

    assert stats["whiteObjectAssets"] + stats["whiteContainerShapes"] >= 1
    assert np.all(cv2.imread(str(background_path)) == 255)
    assert layout["metadata"]["reconstructionSurfaceMode"] == "white_objectized"
    assert layout["elements"][0]["type"] == "text"
    visuals = [item for item in layout["elements"] if item["type"] != "text"]
    assert any(item["x"] <= 30 and item["x"] + item["width"] >= 105 for item in visuals)
    assert all(not (item["x"] >= 170 and item["x"] < 280) for item in visuals)


def test_text_backplates_and_cards_export_as_independent_shapes(tmp_path):
    source = np.full((300, 500, 3), 255, np.uint8)
    cv2.rectangle(source, (25, 30), (190, 80), (225, 155, 90), -1)
    cv2.putText(source, "LABEL", (45, 65), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (25, 25, 25), 2)
    cv2.rectangle(source, (230, 35), (460, 245), (245, 238, 225), -1)
    cv2.putText(source, "CARD", (265, 105), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (35, 35, 35), 2)
    source_path = tmp_path / "source.png"
    background_path = tmp_path / "backgrounds" / "page_1.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 500, "height": 300}, "elements": [
        {"id": "label_text", "type": "text", "x": 40, "y": 40, "width": 115, "height": 35, "zIndex": 20, "text": "LABEL", "metadata": {"rawOCRBBox": [40, 40, 155, 75]}},
        {"id": "card_text", "type": "text", "x": 260, "y": 75, "width": 105, "height": 42, "zIndex": 20, "text": "CARD", "metadata": {"rawOCRBBox": [260, 75, 365, 117]}},
    ]}

    stats = objectize_on_white(source_path, background_path, layout, tmp_path / "assets", "demo", 1)

    shapes = [item for item in layout["elements"] if item.get("metadata", {}).get("layerRole") == "container"]
    assert stats["whiteContainerShapes"] >= 2
    assert any(item["x"] <= 25 and item["width"] >= 160 for item in shapes)
    assert any(item["x"] <= 230 and item["width"] >= 225 for item in shapes)
    assert all(item["zIndex"] < 20 for item in shapes)
    assert np.all(cv2.imread(str(background_path)) == 255)
    preview = tmp_path / "preview.png"
    render_preview(background_path, layout, preview)
    pixels = cv2.imread(str(preview))
    assert np.max(np.abs(pixels[38, 32].astype(int) - source[38, 32].astype(int))) < 8
    assert np.all(pixels[10, 10] == 255)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for item in shapes:
        PPTXRenderer()._add_element(slide, item, 0.02, 0.02)
    assert len(slide.shapes) == len(shapes)
    assert all(shape.shape_type == 1 for shape in slide.shapes)  # native PowerPoint shapes


def test_pale_module_backplate_survives_with_foreground_and_group(tmp_path):
    source = np.full((240, 400, 3), 250, np.uint8)
    cv2.rectangle(source, (48, 42), (226, 112), (242, 244, 246), -1)
    cv2.circle(source, (76, 77), 13, (180, 80, 40), -1)
    cv2.putText(source, "ITEM", (101, 85), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (30, 30, 30), 2)
    source_path = tmp_path / "pale.png"
    background_path = tmp_path / "backgrounds" / "page_1.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 400, "height": 240}, "elements": [
        {"id": "module_icon", "type": "ellipse", "x": 63, "y": 64, "width": 26, "height": 26, "zIndex": 10, "groupId": "module_a", "style": {"fill": "#2850B4"}},
        {"id": "module_title", "type": "text", "x": 99, "y": 62, "width": 91, "height": 29, "zIndex": 20, "groupId": "module_a", "text": "ITEM", "metadata": {"rawOCRBBox": [99, 62, 190, 91]}},
    ]}
    stats = objectize_on_white(source_path, background_path, layout, tmp_path / "assets", "pale-module", 1)
    plates = [item for item in layout["elements"] if item.get("metadata", {}).get("layerRole") == "container"]
    assert stats["whiteContainerShapes"] >= 1
    plate = next(item for item in plates if item["x"] <= 48 and item["x"] + item["width"] >= 226)
    assert plate["groupId"] == "module_a"
    assert set(plate["metadata"]["moduleMemberIds"]) == {"module_icon", "module_title"}
    assert plate["zIndex"] < 10
    assert plate["type"] in {"rectangle", "roundedRectangle", "image"}
    assert np.all(cv2.imread(str(background_path)) == 255)
    preview = tmp_path / "preview.png"
    render_preview(background_path, layout, preview)
    pixels = cv2.imread(str(preview))
    assert np.max(np.abs(pixels[48, 51].astype(int) - source[48, 51].astype(int))) < 8
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    PPTXRenderer()._add_element(slide, plate, 0.02, 0.02)
    assert len(slide.shapes) == 1


def test_page_wide_pale_background_is_not_a_local_backplate(tmp_path):
    source = np.full((240, 400, 3), 245, np.uint8)
    source_path = tmp_path / "page.png"
    background_path = tmp_path / "backgrounds" / "page_1.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 400, "height": 240}, "elements": [
        {"id": "title", "type": "text", "x": 80, "y": 60, "width": 120, "height": 30, "text": "Title", "metadata": {"rawOCRBBox": [80, 60, 200, 90]}}
    ]}
    objectize_on_white(source_path, background_path, layout, tmp_path / "assets", "page-bg", 1)
    assert not [item for item in layout["elements"] if item.get("metadata", {}).get("layerRole") == "container"]


def test_colored_title_surface_keeps_white_inner_frame_as_movable_detail(tmp_path):
    source = np.full((240, 420, 3), 255, np.uint8)
    cv2.rectangle(source, (42, 48), (352, 142), (165, 86, 38), -1)
    cv2.rectangle(source, (54, 60), (340, 130), (255, 255, 255), 4)
    cv2.putText(source, "TITLE", (130, 105), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (255, 255, 255), 2)
    source_path = tmp_path / "source.png"
    background_path = tmp_path / "backgrounds" / "page_1.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 420, "height": 240}, "elements": [
        {"id": "title", "type": "text", "x": 125, "y": 78, "width": 110, "height": 35, "zIndex": 20, "groupId": "header", "text": "TITLE", "style": {"color": "#FFFFFF", "fontSize": 24}, "metadata": {"rawOCRBBox": [125, 78, 235, 113]}}
    ]}
    stats = objectize_on_white(source_path, background_path, layout, tmp_path / "assets", "inner-frame", 1)
    parents = [item for item in layout["elements"] if (item.get("metadata") or {}).get("layerRole") == "container" and item.get("type") in {"rectangle", "roundedRectangle"}]
    details = [item for item in layout["elements"] if (item.get("metadata") or {}).get("layerRole") == "container_detail"]
    assert parents and details and stats["whiteInternalDetails"] >= 1
    parent = next(item for item in parents if item["x"] <= 42 and item["x"] + item["width"] >= 352)
    detail = next(item for item in details if item["x"] <= 54 and item["x"] + item["width"] >= 340)
    assert parent["groupId"] == detail["groupId"] == "header"
    assert detail["id"] in parent["metadata"]["moduleMemberIds"]
    assert parent["zIndex"] < detail["zIndex"] < layout["elements"][0]["zIndex"]
    assert np.all(cv2.imread(str(background_path)) == 255)
    if detail["type"] == "image":
        asset = cv2.imread(str(tmp_path / "assets" / Path(detail["src"]).name), cv2.IMREAD_UNCHANGED)
        assert asset.shape[2] == 4 and np.count_nonzero(asset[:, :, 3]) > 0
        assert asset[asset.shape[0] // 2, asset.shape[1] // 2, 3] == 0
        detail["src"] = str(tmp_path / "assets" / Path(detail["src"]).name)
    preview = tmp_path / "preview.png"
    render_preview(background_path, layout, preview)
    pixels = cv2.imread(str(preview))
    assert np.max(np.abs(pixels[60, 100].astype(int) - source[60, 100].astype(int))) < 12
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for item in (parent, detail, layout["elements"][0]):
        PPTXRenderer()._add_element(slide, item, 0.02, 0.02)
    assert len(slide.shapes) == 3
    assert any(shape.has_text_frame and "TITLE" in shape.text for shape in slide.shapes)


def test_outline_card_becomes_movable_container(tmp_path):
    source = np.full((240, 400, 3), 255, np.uint8)
    cv2.rectangle(source, (45, 35), (330, 190), (120, 65, 35), 3)
    cv2.putText(source, "NOTE", (80, 100), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (20, 20, 20), 2)
    source_path = tmp_path / "source.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 400, "height": 240}, "elements": [{"id": "note", "type": "text", "x": 78, "y": 72, "width": 105, "height": 38, "zIndex": 20, "text": "NOTE", "metadata": {"rawOCRBBox": [78, 72, 183, 110]}}]}

    stats = objectize_on_white(source_path, tmp_path / "backgrounds" / "page_1.png", layout, tmp_path / "assets", "demo", 1)

    containers = [item for item in layout["elements"] if item.get("metadata", {}).get("layerRole") == "container"]
    assert stats["whiteContainerShapes"] >= 1
    assert any(item["x"] <= 46 and item["width"] >= 280 and item["zIndex"] < 20 for item in containers)


def test_textured_bordered_panel_uses_independent_text_clean_image(tmp_path):
    source = np.full((220, 380, 3), 255, np.uint8)
    for x in range(42, 338):
        source[32:184, x] = (85 + (x % 90), 145, 215)
    cv2.rectangle(source, (40, 30), (339, 185), (28, 52, 100), 3)
    cv2.putText(source, "INFO", (95, 105), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (10, 10, 10), 2)
    source_path = tmp_path / "source.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 380, "height": 220}, "elements": [{"id": "info", "type": "text", "x": 90, "y": 76, "width": 95, "height": 40, "zIndex": 20, "text": "INFO", "metadata": {"rawOCRBBox": [90, 76, 185, 116]}}]}

    stats = objectize_on_white(source_path, tmp_path / "backgrounds" / "page_1.png", layout, tmp_path / "assets", "demo", 1)

    assets = [item for item in layout["elements"] if item.get("metadata", {}).get("layerRole") == "container" and item["type"] == "image"]
    assert stats["whiteContainerShapes"] >= 1
    assert assets and assets[0]["zIndex"] < 20
    assert assets[0]["metadata"]["textCleaned"] is True
    assert assets[0]["metadata"]["editableTextIds"] == ["info"]
    asset_path = tmp_path / "assets" / f"{assets[0]['id']}.png"
    assert asset_path.exists()
    before = asset_path.read_bytes()
    assets[0]["src"] = str(asset_path)
    from app.services.reconstruction.text_erasure import erase_editable_text_sources
    cleanup = erase_editable_text_sources(tmp_path / "backgrounds" / "page_1.png", layout)
    assert cleanup["assetTextErased"] == 0
    assert asset_path.read_bytes() == before


def test_text_cleanup_preserves_colored_header_above_white_panel():
    from app.services.reconstruction.white_objectization import _clean_container_text

    source = np.full((64, 220, 3), 255, np.uint8)
    for x in range(220):
        source[:32, x] = (180 + x // 12, 110 + x // 18, 30 + x // 20)
    cv2.putText(source, "HEADING", (35, 29), cv2.FONT_HERSHEY_SIMPLEX, 0.62, (255, 255, 255), 2)
    mask = np.zeros(source.shape[:2], np.uint8)
    mask[12:43, 31:180] = 255  # OCR box crosses the header edge.

    cleaned = _clean_container_text(source, mask)

    assert np.max(np.abs(cleaned[21, 95].astype(int) - source[21, 20].astype(int))) < 30
    assert np.min(cleaned[39, 95]) >= 245
    assert np.max(np.abs(cleaned[31, 95].astype(int) - source[31, 95].astype(int))) < 10


def test_small_contour_within_long_ocr_line_stays_owned_by_editable_text():
    from app.services.reconstruction.white_objectization import _is_text_glyph_candidate

    line = {"id": "heading", "type": "text", "x": 40, "y": 20, "width": 420, "height": 70,
            "text": "Heading", "metadata": {"rawOCRBBox": [40, 20, 460, 90]}}
    assert _is_text_glyph_candidate((100, 29, 146, 78), [line], 500, 160)
    assert not _is_text_glyph_candidate((480, 29, 526, 78), [line], 600, 160)
    assert not _is_text_glyph_candidate((55, 10, 455, 110), [line], 500, 160)


def test_container_cleanup_uses_raw_ocr_bounds_instead_of_expanded_textbox(tmp_path):
    source = np.full((210, 390, 3), 255, np.uint8)
    cv2.rectangle(source, (35, 25), (350, 180), (85, 55, 35), 3)
    cv2.rectangle(source, (42, 48), (343, 92), (175, 80, 20), -1)
    cv2.putText(source, "TITLE", (100, 76), cv2.FONT_HERSHEY_SIMPLEX, 0.65, (255, 255, 255), 2)
    source_path = tmp_path / "source.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 390, "height": 210}, "elements": [
        {"id": "title", "type": "text", "x": 75, "y": 42, "width": 210, "height": 72,
         "text": "TITLE", "metadata": {"rawOCRBBox": [96, 54, 200, 80]}}
    ]}
    objectize_on_white(source_path, tmp_path / "background.png", layout, tmp_path / "assets", "raw-ocr", 1)
    assets = [item for item in layout["elements"] if item.get("metadata", {}).get("layerRole") == "container" and item["type"] == "image"]
    assert assets
    panel = next(item for item in assets if item["x"] <= 42 and item["x"] + item["width"] >= 343)
    image = cv2.imread(str(tmp_path / "assets" / Path(panel["src"]).name), cv2.IMREAD_UNCHANGED)
    pixel = image[86 - panel["y"], 150 - panel["x"], :3]
    assert np.max(np.abs(pixel.astype(int) - source[86, 150].astype(int))) < 10


def test_faint_text_shadow_is_not_exported_as_visual_asset(tmp_path):
    from app.services.reconstruction.white_objectization import _is_text_shadow_residual

    asset_dir = tmp_path / "assets"
    asset_dir.mkdir()
    image = np.full((35, 80, 4), 250, np.uint8)
    image[:, :, 3] = 0
    image[8:29:3, 12:68:3, 3] = 255
    path = asset_dir / "shadow.png"
    cv2.imwrite(str(path), image)
    asset = {"id": "shadow", "type": "image", "x": 30, "y": 25, "width": 80, "height": 35, "src": f"/media/assets/demo/{path.name}"}
    title = {"id": "title", "type": "text", "x": 25, "y": 20, "width": 120, "height": 45,
             "style": {"color": "#17365D"}, "metadata": {"rawOCRBBox": [25, 20, 145, 65]}}

    assert _is_text_shadow_residual(asset, [title], asset_dir, 200, 100)
    image[8:29:3, 12:68:3, :3] = (60, 120, 190)
    cv2.imwrite(str(path), image)
    assert not _is_text_shadow_residual(asset, [title], asset_dir, 200, 100)
    image[:, :, :3] = 250
    image[:, :, 3] = 255
    cv2.imwrite(str(path), image)
    assert not _is_text_shadow_residual(asset, [title], asset_dir, 200, 100)


def test_residual_icon_inside_flat_card_is_not_claimed_by_card(tmp_path):
    source = np.full((220, 360, 3), 255, np.uint8)
    cv2.rectangle(source, (30, 30), (320, 180), (230, 220, 205), -1)
    cv2.circle(source, (105, 105), 20, (180, 55, 25), -1)
    source_path = tmp_path / "source.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 360, "height": 220}, "elements": []}

    stats = objectize_on_white(source_path, tmp_path / "backgrounds" / "page_1.png", layout, tmp_path / "assets", "demo", 1)

    assert stats["whiteContainerShapes"] >= 1
    assert stats["whiteObjectAssets"] >= 1
    residual = next(item for item in layout["elements"] if item.get("metadata", {}).get("layerRole") == "residual" and 80 <= item["x"] <= 110)
    container = next(item for item in layout["elements"] if item.get("metadata", {}).get("layerRole") == "container")
    assert residual["zIndex"] > container["zIndex"]
    residual["src"] = str(tmp_path / "assets" / f"{residual['id']}.png")
    preview = tmp_path / "preview.png"
    render_preview(tmp_path / "backgrounds" / "page_1.png", layout, preview)
    assert np.max(np.abs(cv2.imread(str(preview))[105, 105].astype(int) - source[105, 105].astype(int))) < 8


def test_ocr_box_does_not_erase_its_pale_supporting_strip(tmp_path):
    source = np.full((150, 320, 3), 255, np.uint8)
    cv2.rectangle(source, (45, 52), (265, 92), (249, 250, 249), -1)
    cv2.putText(source, "LABEL", (85, 80), cv2.FONT_HERSHEY_SIMPLEX, 0.65, (25, 25, 25), 2)
    source_path = tmp_path / "source.png"
    background_path = tmp_path / "backgrounds" / "page.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 320, "height": 150}, "elements": [
        {"id": "label", "type": "text", "x": 40, "y": 47, "width": 230, "height": 50,
         "text": "LABEL", "zIndex": 20, "metadata": {"rawOCRBBox": [40, 47, 270, 97]}}
    ]}
    stats = objectize_on_white(source_path, background_path, layout, tmp_path / "assets", "pale", 1)
    visuals = [item for item in layout["elements"] if item["type"] != "text"]
    assert visuals and stats["whiteObjectAssets"] + stats["whiteObjectShapes"] > 0
    assert any(item["x"] <= 50 and item["x"] + item["width"] >= 260 for item in visuals)
    assert np.all(cv2.imread(str(background_path)) == 255)
    for item in visuals:
        if item["type"] == "image":
            item["src"] = str(tmp_path / "assets" / Path(item["src"]).name)
    preview = tmp_path / "preview.png"
    render_preview(background_path, layout, preview)
    assert np.max(np.abs(cv2.imread(str(preview))[55, 55].astype(int) - source[55, 55].astype(int))) < 8


def test_unplanned_complex_local_region_is_movable_fallback(tmp_path):
    source = np.full((170, 300, 3), 255, np.uint8)
    for row in range(40, 120):
        source[row, 35:140] = (235 + row % 3, 240, 244)
    cv2.circle(source, (88, 80), 19, (80, 100, 190), -1)
    source_path = tmp_path / "source.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 300, "height": 170}, "elements": []}
    stats = objectize_on_white(source_path, tmp_path / "backgrounds" / "page.png", layout, tmp_path / "assets", "fallback", 1)
    assert stats["whiteObjectAssets"] + stats["whiteObjectShapes"] >= 2
    assert any(item["x"] <= 88 <= item["x"] + item["width"] for item in layout["elements"])


def test_round_badge_and_pale_support_remain_separate_movable_objects(tmp_path):
    source = np.full((220, 320, 3), 255, np.uint8)
    cv2.rectangle(source, (38, 95), (210, 177), (246, 248, 250), -1)
    cv2.circle(source, (84, 92), 30, (35, 92, 185), -1, cv2.LINE_AA)
    cv2.circle(source, (84, 92), 29, (255, 255, 255), 2, cv2.LINE_AA)
    cv2.circle(source, (84, 92), 9, (230, 220, 55), -1, cv2.LINE_AA)
    source_path = tmp_path / "source.png"
    background = tmp_path / "background.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 320, "height": 220}, "elements": []}
    objectize_on_white(source_path, background, layout, tmp_path / "assets", "round", 1)
    badges = [item for item in layout["elements"] if item["type"] == "image" and item["x"] <= 84 <= item["x"] + item["width"] and item["y"] <= 70]
    assert badges, [(item["type"], item["x"], item["y"], item["width"], item["height"]) for item in layout["elements"]]
    badge = badges[0]
    image = cv2.imread(str(tmp_path / "assets" / Path(badge["src"]).name), cv2.IMREAD_UNCHANGED)
    assert image.shape[2] == 4
    assert image[0, 0, 3] < 32 and image[image.shape[0] // 2, image.shape[1] // 2, 3] > 220, [(item["type"], item["x"], item["y"], item["width"], item["height"], item.get("metadata", {}).get("layerRole")) for item in layout["elements"]]
    assert any(item is not badge and item["x"] <= 45 and item["x"] + item["width"] >= 200 and item["y"] <= 110 <= item["y"] + item["height"] for item in layout["elements"]), [(item["type"], item["x"], item["y"], item["width"], item["height"], item.get("metadata", {}).get("layerRole")) for item in layout["elements"]]
    assert np.all(cv2.imread(str(background)) == 255)
    report = repair_objectized_modules(source_path, layout, tmp_path / "assets", "round")
    assert report["roundCutoutsChecked"] >= 1 and report["roundCutoutIssues"] == 0
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for item in layout["elements"]:
        if item["type"] == "image":
            item["src"] = str(tmp_path / "assets" / Path(item["src"]).name)
        PPTXRenderer()._add_element(slide, item, 0.02, 0.02)
    assert len(slide.shapes) == len(layout["elements"])
    preview = tmp_path / "preview.png"
    render_preview(background, layout, preview)
    pixels = cv2.imread(str(preview))
    assert np.max(np.abs(pixels[92, 84].astype(int) - source[92, 84].astype(int))) < 12
    assert np.max(np.abs(pixels[130, 170].astype(int) - source[130, 170].astype(int))) < 12


def test_large_map_silk_and_pale_card_survive_white_objectization(tmp_path):
    source = np.full((400, 700, 3), 255, np.uint8)
    map_contour = np.array([[45, 55], [300, 18], [650, 48], [675, 170], [620, 270], [490, 300], [270, 320], [65, 275]], np.int32)
    cv2.fillPoly(source, [map_contour], (105, 145, 175))
    cv2.rectangle(source, (450, 75), (590, 145), (245, 248, 250), -1)
    silk = np.array([[80, 350], [270, 340], [460, 358], [625, 335], [590, 383], [310, 375]], np.int32)
    cv2.fillPoly(source, [silk], (35, 50, 195))
    path = tmp_path / "source.png"
    background = tmp_path / "background.png"
    cv2.imwrite(str(path), source)
    layout = {"slide": {"width": 700, "height": 400}, "elements": []}
    stats = objectize_on_white(path, background, layout, tmp_path / "assets", "complex", 1)
    assert stats["whiteObjectAssets"] >= 2
    assert any(item["type"] == "image" and item["width"] > 500 for item in layout["elements"])
    assert any(item["y"] >= 330 and item["type"] == "image" for item in layout["elements"])
    assert np.all(cv2.imread(str(background)) == 255)
    for item in layout["elements"]:
        if item["type"] == "image":
            item["src"] = str(tmp_path / "assets" / Path(item["src"]).name)
    preview = tmp_path / "preview.png"
    render_preview(background, layout, preview)
    pixels = cv2.imread(str(preview))
    for y, x in ((150, 170), (370, 420), (100, 500)):
        assert np.max(np.abs(pixels[y, x].astype(int) - source[y, x].astype(int))) < 15
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for item in layout["elements"]:
        PPTXRenderer()._add_element(slide, item, 0.01, 0.01)
    assert len(slide.shapes) == len(layout["elements"])
    assert sum(shape.shape_type == 13 for shape in slide.shapes) >= 2


def test_full_slide_detailed_map_is_movable_instead_of_discarded_or_fragmented(tmp_path):
    height, width = 180, 320
    source = np.full((height, width, 3), (85, 145, 170), np.uint8)
    for y in range(0, height, 18):
        for x in range(0, width, 20):
            color = (60 + (x * 3 + y) % 75, 95 + (x + y * 2) % 80, 120 + (x * 2 + y) % 90)
            cv2.rectangle(source, (x, y), (min(width - 1, x + 19), min(height - 1, y + 17)), color, -1)
    cv2.line(source, (0, 25), (width - 1, 150), (15, 35, 45), 3)
    path = tmp_path / "source.png"
    background = tmp_path / "background.png"
    cv2.imwrite(str(path), source)
    layout = {"slide": {"width": width, "height": height}, "elements": []}

    objectize_on_white(path, background, layout, tmp_path / "assets", "detailed-map", 1)

    images = [item for item in layout["elements"] if item["type"] == "image"]
    assert images
    assert len(layout["elements"]) < 10
    assert all(item["src"].startswith("/media/assets/") for item in images)
    for item in images:
        item["src"] = str(tmp_path / "assets" / Path(item["src"]).name)
    preview = tmp_path / "preview.png"
    render_preview(background, layout, preview)
    reconstructed = cv2.imread(str(preview))
    assert np.mean(np.abs(reconstructed.astype(np.int16) - source.astype(np.int16))) < 5
    assert np.all(cv2.imread(str(background)) == 255)


def test_transparent_existing_image_does_not_claim_separate_visual_in_its_bbox(tmp_path):
    source = np.full((180, 320, 3), 255, np.uint8)
    first = np.array([[35, 45], [95, 35], [110, 105], [55, 125]], np.int32)
    second = np.array([[190, 55], [260, 35], [275, 110], [210, 125]], np.int32)
    cv2.fillPoly(source, [first], (50, 105, 185))
    cv2.fillPoly(source, [second], (80, 145, 70))
    source_path = tmp_path / "source.png"
    background = tmp_path / "background.png"
    asset_dir = tmp_path / "assets"
    asset_dir.mkdir()
    cv2.imwrite(str(source_path), source)
    existing = np.zeros((120, 270, 4), np.uint8)
    existing[:, :, :3] = source[20:140, 25:295]
    mask = np.zeros(source.shape[:2], np.uint8)
    cv2.fillPoly(mask, [first], 255)
    existing[:, :, 3] = mask[20:140, 25:295]
    existing_path = asset_dir / "existing.png"
    cv2.imwrite(str(existing_path), existing)
    layout = {"slide": {"width": 320, "height": 180}, "elements": [
        {"id": "existing", "type": "image", "x": 25, "y": 20, "width": 270, "height": 120,
         "zIndex": 2, "src": str(existing_path), "metadata": {"reconstructionStrategy": "cutout_image"}}
    ]}

    objectize_on_white(source_path, background, layout, asset_dir, "transparent", 1)

    recovered = [item for item in layout["elements"] if item["id"] != "existing" and item["type"] == "image" and item["x"] <= 230 <= item["x"] + item["width"]]
    assert recovered, [(item["id"], item["type"], item["x"], item["width"]) for item in layout["elements"]]
    assert all(item["width"] < 120 for item in recovered)
    for item in recovered:
        item["src"] = str(asset_dir / Path(item["src"]).name)
    preview = tmp_path / "preview.png"
    render_preview(background, layout, preview)
    pixels = cv2.imread(str(preview))
    for y, x in ((75, 75), (70, 235)):
        assert np.max(np.abs(pixels[y, x].astype(int) - source[y, x].astype(int))) < 10
    assert np.all(cv2.imread(str(background)) == 255)


def test_thin_decorative_strips_are_movable_and_detected_by_qa(tmp_path):
    source = np.full((180, 320, 3), 255, np.uint8)
    source[42, 30:230] = (50, 95, 185)
    source[65:155, 275] = (40, 150, 80)
    source_path = tmp_path / "source.png"
    background = tmp_path / "background.png"
    preview = tmp_path / "preview.png"
    cv2.imwrite(str(source_path), source)
    cv2.imwrite(str(background), np.full_like(source, 255))
    missing = audit_objectization(source_path, background, background, {"elements": []})
    assert missing["missingVisualObjects"] == 2
    layout = {"slide": {"width": 320, "height": 180}, "elements": []}

    stats = objectize_on_white(source_path, background, layout, tmp_path / "assets", "strips", 1)

    strips = [item for item in layout["elements"] if item["type"] == "image"]
    assert stats["whiteObjectAssets"] >= 2 and len(strips) >= 2
    assert any(item["height"] <= 5 and item["width"] >= 190 for item in strips)
    assert any(item["width"] <= 5 and item["height"] >= 85 for item in strips)
    for item in strips:
        item["src"] = str(tmp_path / "assets" / Path(item["src"]).name)
    render_preview(background, layout, preview)
    report = audit_objectization(source_path, background, preview, layout)
    assert report["missingVisualObjects"] == 0
    pixels = cv2.imread(str(preview))
    assert np.max(np.abs(pixels[42, 100].astype(int) - source[42, 100].astype(int))) < 10
    assert np.max(np.abs(pixels[100, 275].astype(int) - source[100, 275].astype(int))) < 10
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for item in strips:
        PPTXRenderer()._add_element(slide, item, 0.02, 0.02)
    assert len(slide.shapes) == len(strips)


def test_group_and_line_bounds_do_not_hide_unrelated_artwork(tmp_path):
    source = np.full((200, 340, 3), 255, np.uint8)
    cv2.line(source, (35, 80), (300, 80), (50, 95, 185), 2)
    icon = np.array([[140, 118], [195, 108], [210, 160], [155, 175]], np.int32)
    cv2.fillPoly(source, [icon], (70, 145, 40))
    source_path = tmp_path / "source.png"
    background = tmp_path / "background.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 340, "height": 200}, "elements": [
        {"id": "module_group", "type": "group", "x": 20, "y": 30, "width": 300, "height": 150,
         "metadata": {"reconstructionStrategy": "group"}},
        {"id": "rule", "type": "line", "x": 35, "y": 35, "width": 265, "height": 90,
         "style": {"stroke": "#B95F32", "strokeWidth": 2},
         "metadata": {"reconstructionStrategy": "native_shape"}},
    ]}

    objectize_on_white(source_path, background, layout, tmp_path / "assets", "group", 1)

    recovered = [item for item in layout["elements"] if item["type"] == "image" and item["x"] <= 175 <= item["x"] + item["width"]]
    assert recovered
    assert all(item["width"] < 100 for item in recovered)
    for item in recovered:
        item["src"] = str(tmp_path / "assets" / Path(item["src"]).name)
    preview = tmp_path / "preview.png"
    render_preview(background, layout, preview)
    assert np.max(np.abs(cv2.imread(str(preview))[140, 170].astype(int) - source[140, 170].astype(int))) < 10
    assert np.all(cv2.imread(str(background)) == 255)


def test_native_white_card_stays_below_chart_detail_asset(tmp_path):
    source = np.full((120, 220, 3), 255, np.uint8)
    cv2.line(source, (50, 80), (170, 45), (45, 75, 130), 4)
    alpha = np.zeros(source.shape[:2], np.uint8)
    cv2.line(alpha, (50, 80), (170, 45), 255, 4)
    asset = tmp_path / "chart.png"
    cv2.imwrite(str(asset), np.dstack((source, alpha)))
    layout = {"slide": {"width": 220, "height": 120}, "elements": [
        {"id": "chart_detail", "type": "image", "x": 0, "y": 0, "width": 220, "height": 120,
         "zIndex": 1, "src": str(asset), "metadata": {"layerRole": "residual", "reconstructionStrategy": "cutout_image"}},
        {"id": "card", "type": "roundedRectangle", "x": 35, "y": 30, "width": 150, "height": 65,
         "zIndex": 10, "style": {"fill": "#FEFEFE", "stroke": "#FEFEFE"},
         "metadata": {"reconstructionStrategy": "native_shape"}},
    ]}

    background = tmp_path / "background.png"
    cv2.imwrite(str(background), np.full_like(source, 255))
    source_path = tmp_path / "source.png"
    cv2.imwrite(str(source_path), source)
    before_preview = tmp_path / "before.png"
    render_preview(background, layout, before_preview)
    before = audit_objectization(source_path, background, before_preview, layout)
    assert before["coveredMissingVisualPixels"] > 100
    assert before["missingVisualObjects"] >= 1

    layer_objectized_elements(layout["elements"])

    assert layout["elements"][0]["zIndex"] > layout["elements"][1]["zIndex"]
    preview = tmp_path / "preview.png"
    render_preview(background, layout, preview)
    rendered = cv2.imread(str(preview))
    assert np.max(np.abs(rendered[65, 105].astype(np.int16) - source[65, 105].astype(np.int16))) < 5
    after = audit_objectization(source_path, background, preview, layout)
    assert after["coveredMissingVisualPixels"] < before["coveredMissingVisualPixels"]


def test_planned_icon_image_stays_above_native_card_fill(tmp_path):
    source = np.full((100, 160, 3), 255, np.uint8)
    cv2.circle(source, (80, 50), 20, (10, 175, 35), -1)
    icon = tmp_path / "icon.png"
    cv2.imwrite(str(icon), source[30:71, 60:101])
    layout = {"slide": {"width": 160, "height": 100}, "elements": [
        {"id": "planned_icon", "type": "image", "x": 60, "y": 30, "width": 41, "height": 41,
         "zIndex": 2, "src": str(icon), "metadata": {"reconstructionStrategy": "local_image"}},
        {"id": "card", "type": "roundedRectangle", "x": 25, "y": 15, "width": 110, "height": 70,
         "zIndex": 10, "style": {"fill": "#FFFFFF", "stroke": "#FFFFFF"},
         "metadata": {"reconstructionStrategy": "native_shape"}},
    ]}

    layer_objectized_elements(layout["elements"])

    assert layout["elements"][0]["zIndex"] > layout["elements"][1]["zIndex"]
    background = tmp_path / "background.png"
    cv2.imwrite(str(background), np.full_like(source, 255))
    preview = tmp_path / "preview.png"
    render_preview(background, layout, preview)
    assert np.array_equal(cv2.imread(str(preview))[50, 80], source[50, 80])


def test_logical_group_does_not_block_round_asset_classification(tmp_path):
    source = np.full((160, 260, 3), 255, np.uint8)
    cv2.circle(source, (90, 80), 27, (45, 95, 185), -1)
    source_path = tmp_path / "source.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": 260, "height": 160}, "elements": [
        {"id": "logical_group", "type": "group", "x": 30, "y": 20, "width": 190, "height": 125,
         "metadata": {"reconstructionStrategy": "group"}}
    ]}

    objectize_on_white(source_path, tmp_path / "background.png", layout, tmp_path / "assets", "group", 1)

    badges = [item for item in layout["elements"] if (item.get("metadata") or {}).get("reconstructionStrategySource") == "round_contour"]
    assert len(badges) == 1
    asset = cv2.imread(str(tmp_path / "assets" / Path(badges[0]["src"]).name), cv2.IMREAD_UNCHANGED)
    assert asset.shape[2] == 4 and asset[0, 0, 3] < 32 and asset[asset.shape[0] // 2, asset.shape[1] // 2, 3] > 220


def test_gradient_page_keeps_card_surfaces_and_ribbon_after_text_cleanup(tmp_path):
    height, width = 240, 420
    source = np.empty((height, width, 3), np.uint8)
    for y in range(height):
        shade = round(247 + 7 * y / (height - 1))
        source[y, :] = (shade, min(255, shade + 1), min(255, shade + 3))
    cv2.rectangle(source, (30, 40), (175, 150), (244, 244, 253), -1)
    cv2.rectangle(source, (220, 40), (375, 150), (242, 245, 252), -1)
    cv2.putText(source, "CARD", (62, 102), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (30, 30, 30), 2)
    ribbon = np.array([[15, 185], [140, 175], [280, 192], [405, 177], [405, 239], [15, 239]], np.int32)
    cv2.fillPoly(source, [ribbon], (30, 45, 190))
    cv2.putText(source, "MOVE", (170, 217), cv2.FONT_HERSHEY_SIMPLEX, 0.65, (220, 230, 250), 2)
    source_path = tmp_path / "source.png"
    background = tmp_path / "background.png"
    cv2.imwrite(str(source_path), source)
    layout = {"slide": {"width": width, "height": height}, "elements": [
        {"id": "card_text", "type": "text", "x": 58, "y": 78, "width": 95, "height": 34, "zIndex": 20,
         "text": "CARD", "metadata": {"rawOCRBBox": [58, 78, 153, 112]}},
        {"id": "ribbon_text", "type": "text", "x": 166, "y": 193, "width": 95, "height": 32, "zIndex": 20,
         "text": "MOVE", "metadata": {"rawOCRBBox": [166, 193, 261, 225]}},
    ]}
    cv2.imwrite(str(background), np.full_like(source, 255))
    before = audit_objectization(source_path, background, background, layout)
    assert before["retainedVisualCoverage"] < 0.4
    assert before["missingVisualObjects"] + before["missingBackplates"] >= 2

    objectize_on_white(source_path, background, layout, tmp_path / "assets", "gradient", 1)
    visuals = [item for item in layout["elements"] if item["type"] in {"image", "rectangle", "roundedRectangle"}]
    assert any(item["x"] <= 40 and item["x"] + item["width"] >= 160 for item in visuals)
    assert any(item["x"] <= 230 and item["x"] + item["width"] >= 360 for item in visuals)
    assert any(item["y"] >= 170 and item["width"] >= 350 for item in visuals)
    assert all(item["width"] * item["height"] < width * height * 0.5 for item in visuals if item["type"] == "image")
    for item in visuals:
        if item["type"] == "image":
            item["src"] = str(tmp_path / "assets" / Path(item["src"]).name)
    erase_editable_text_sources(background, layout, clean_background=False, project_root=tmp_path)
    preview = tmp_path / "preview.png"
    render_preview(background, layout, preview)
    pixels = cv2.imread(str(preview))
    assert np.all(cv2.imread(str(background)) == 255)
    assert np.max(np.abs(pixels[60, 45].astype(int) - source[60, 45].astype(int))) < 15
    assert np.max(np.abs(pixels[205, 100].astype(int) - source[205, 100].astype(int))) < 20
    report = audit_objectization(source_path, background, preview, layout)
    assert report["retainedVisualCoverage"] > 0.9
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for item in visuals:
        PPTXRenderer()._add_element(slide, item, 0.02, 0.02)
    assert len(slide.shapes) == len(visuals)
