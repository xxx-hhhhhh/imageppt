import cv2
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

from app.services.pptx.renderer import PPTXRenderer, font_px_to_pt
from app.services.reconstruction.exclusive_ownership import box_of, build_exclusive_scene
from app.services.reconstruction.mixed_text import restore_text_candidates
from app.services.segmentation.segmentation_provider import OpenCVSegmentationProvider
from app.services.validation.reconstruction_quality import reconstruction_quality
from app.services.validation.service import validate_pptx
from app.services.visual_qa.analyzer import _similarity
from app.utils import image_io


def test_ordinary_latin_line_is_not_split_into_glyph_height_spans(tmp_path):
    image = Image.new('RGB',(500,100),'#F4F9FB')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf',24)
    text = 'Better Earth through Geospatial AI'
    draw.text((20,20),text,font=font,fill='#17365D')
    x,y,r,b = draw.textbbox((20,20),text,font=font)
    source = cv2.cvtColor(np.array(image),cv2.COLOR_RGB2BGR)
    item = {'id':'line','type':'text','text':text,'confidence':.99,'x':x,'y':y,'width':r-x,'height':b-y,
            'style':{'color':'#17365D','fontFamily':'Arial'}}
    candidates = restore_text_candidates(source,item)
    assert len(candidates)==1
    assert candidates[0][0]['text']==text


def test_partial_mixed_line_is_atomic(monkeypatch):
    import app.services.reconstruction.mixed_text as mixed
    item = {'id':'a','text':'whole'}
    monkeypatch.setattr(mixed,'mixed_spans',lambda *_:[{'id':'b','text':'w'},{'id':'c','text':'hole'}])
    def restore(source,entry,*args):
        if entry['id']=='b':
            return {'text':'w','metadata':{}},np.ones((2,2),bool),source
        return None
    monkeypatch.setattr(mixed,'restore_text_candidate',restore)
    assert mixed.restore_text_candidates(np.zeros((2,2,3),np.uint8),item)==[]


def test_normalized_box_clips_right_edge_without_becoming_one_pixel():
    assert box_of({'bbox':{'left':.94,'top':.3,'width':.12,'height':.2}},1000,500)==(940,150,1000,250)


def test_nested_modules_have_distinct_replacement_owners(tmp_path):
    source = np.full((120,200,3),245,np.uint8)
    source[20:70,20:70]=(0,100,220)
    source[20:70,100:150]=(220,100,0)
    path=tmp_path/'source.png'
    image_io.imwrite(path,source)
    scene={'vision':{'reconstructionPlan':{'modules':[
        {'moduleId':'parent','bbox':[10,10,170,90],'children':[
            {'moduleId':'left','bbox':[20,20,70,70]}, {'moduleId':'right','bbox':[100,20,150,70]}]}]}}}
    layout=build_exclusive_scene(path,{'elements':[]},scene,[],tmp_path/'out','test',1)
    images=[e for e in layout['elements'] if e.get('owner')=='movable_image']
    assert {'left','right','parent'}<={e.get('groupId') for e in images}


def test_nested_contours_propose_objects_inside_cards(tmp_path):
    source=np.full((400,600,3),255,np.uint8)
    cv2.rectangle(source,(40,40),(550,350),(0,0,0),2)
    cv2.circle(source,(160,160),40,(30,90,150),-1)
    path=tmp_path/'source.png'
    image_io.imwrite(path,source)
    proposals=OpenCVSegmentationProvider().segment(path)
    assert any(p['bbox']['width']<150 and p['bbox']['height']<150 for p in proposals)


def test_segmentation_excludes_text_prompts_but_keeps_visuals():
    provider=OpenCVSegmentationProvider()
    provider.text_regions=[{'x':20,'y':20,'width':150,'height':30,'confidence':.99}]
    assert provider._text_dominated(20,20,150,30)
    assert not provider._text_dominated(10,10,300,200)


def test_text_on_flat_card_allows_native_shape_after_glyph_ownership(tmp_path):
    image=Image.new('RGB',(360,120),'white')
    draw=ImageDraw.Draw(image)
    draw.rectangle((10,10,339,99),fill='#DAEDFC')
    font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',24)
    draw.text((30,30),'原生卡片',font=font,fill='#17365D')
    x,y,r,b=draw.textbbox((30,30),'原生卡片',font=font)
    text={'id':'text','type':'text','text':'原生卡片','confidence':.99,'x':x,'y':y,'width':r-x,'height':b-y,
          'style':{'fontFamily':'Microsoft YaHei','color':'#17365D'}}
    rectangle={'id':'card','type':'rectangle','x':10,'y':10,'width':330,'height':90}
    path=tmp_path/'source.png'
    image.save(path)
    layout=build_exclusive_scene(path,{'elements':[text,rectangle]}, {},[],tmp_path/'out','test',1)
    assert any(e.get('owner')=='editable_text' for e in layout['elements'])
    assert any(e.get('owner')=='native_shape' for e in layout['elements'])


def test_quality_does_not_confuse_pixels_and_editability():
    layout={'slide':{'width':100,'height':100},'elements':[], 'metadata':{'ownershipAudit':{
        'textDecisions':[{'id':'a','text':'AB','sourceCharacterCount':2,'editableCharacterCount':1,'reason':'partial'}],
        'nodes':[{'owner':'movable_image','sourcePixelCount':7000}]}}}
    score=reconstruction_quality(layout,{'missingAssetCount':0,'rasterNativeDuplicateTextCount':0},{})
    assert score['editableTextCoverage']==0 and score['editableCharacterCoverage']==.5
    assert score['fragmentedEditableLineCount']==1
    assert score['objectExtractionCoverage'] is None and score['missingVisualCount'] is None
    assert score['qualityStatus']=='partial_editable'


def test_export_validation_reopens_actual_text_and_detects_loss(tmp_path):
    element={'id':'t','type':'text','text':'Editable body','x':10,'y':10,'width':200,'height':30,'style':{'fontSize':20}}
    layout={'slide':{'width':400,'height':200},'elements':[element]}
    path,report=PPTXRenderer().render_project('test',[layout],output_dir=tmp_path)
    assert report['valid'] and report['ocrCoverage'][0]['ratio']==1
    deck=Presentation(path)
    deck.slides[0].shapes[0].text='Lost body'
    deck.save(path)
    report=validate_pptx(path,[layout])
    assert not report['valid'] and report['ocrCoverage'][0]['ratio']==0
    assert report['ocrCoverage'][0]['missingText']==['Editable body']


def test_small_source_font_is_not_silently_inflated():
    assert font_px_to_pt(7,13.333/1920)<4


def test_ssim_penalizes_blur_not_just_average_brightness():
    source=np.full((100,200,3),255,np.uint8)
    cv2.putText(source,'DETAIL',(10,65),cv2.FONT_HERSHEY_SIMPLEX,1.5,(0,0,0),2)
    blur=cv2.GaussianBlur(source,(15,15),3)
    assert _similarity(source,source)[2]==1
    assert _similarity(source,blur)[2]<.9
