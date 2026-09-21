from __future__ import annotations

import json
import math
import zipfile
from pathlib import Path
from typing import Any

from pptx import Presentation


def _overlap(a: dict[str, Any], b: dict[str, Any]) -> float:
    left = max(a["x"], b["x"])
    top = max(a["y"], b["y"])
    right = min(a["x"] + a["width"], b["x"] + b["width"])
    bottom = min(a["y"] + a["height"], b["y"] + b["height"])
    if right <= left or bottom <= top:
        return 0.0
    return (right - left) * (bottom - top) / max(1.0, min(a["width"] * a["height"], b["width"] * b["height"]))


def validate_pptx(pptx_path: Path, layouts: list[dict[str, Any]]) -> dict[str, Any]:
    report: dict[str, Any] = {
        "valid": True,
        "structural": {"zipReadable": False, "slideCount": 0, "textCount": 0, "imageCount": 0, "shapeCount": 0},
        "ocrCoverage": [],
        "layout": {"outOfBounds": [], "invalidNumbers": [], "heavyOverlaps": []},
        "warnings": [],
    }
    try:
        with zipfile.ZipFile(pptx_path) as archive:
            report["structural"]["zipReadable"] = archive.testzip() is None
        presentation = Presentation(str(pptx_path))
        report["structural"]["slideCount"] = len(presentation.slides)
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    report["structural"]["textCount"] += 1
                elif shape.shape_type == 13:
                    report["structural"]["imageCount"] += 1
                else:
                    report["structural"]["shapeCount"] += 1
    except Exception as exc:
        report["valid"] = False
        report["warnings"].append(f"PPTX reopen failed: {exc}")

    if report["structural"]["slideCount"] != len(layouts):
        report["valid"] = False
        report["warnings"].append("PPTX slide count does not match Layout JSON count")

    for page_index, layout in enumerate(layouts, start=1):
        width = layout["slide"]["width"]
        height = layout["slide"]["height"]
        elements = layout.get("elements", [])
        page_coverage = {"page": page_index, "recognizedText": 0, "exportedText": 0, "ratio": 1.0}
        for element in elements:
            values = [element.get(key, 0) for key in ("x", "y", "width", "height", "rotation")]
            if any(not isinstance(value, (int, float)) or not math.isfinite(float(value)) for value in values):
                report["layout"]["invalidNumbers"].append({"page": page_index, "id": element.get("id")})
                report["valid"] = False
            if element.get("type") != "background" and (element["x"] < -2 or element["y"] < -2 or element["x"] + element["width"] > width + 2 or element["y"] + element["height"] > height + 2):
                report["layout"]["outOfBounds"].append({"page": page_index, "id": element.get("id")})
            if element.get("type") == "text":
                page_coverage["recognizedText"] += 1
        for left_index, left in enumerate(elements):
            if left.get("type") in {"background", "group"}:
                continue
            for right in elements[left_index + 1:]:
                if right.get("type") in {"background", "group"}:
                    continue
                if left.get("groupId") and left.get("groupId") == right.get("groupId"):
                    # A container, icon circle, icon image and title are expected
                    # to overlap inside one semantic label group.
                    continue
                ratio = _overlap(left, right)
                if ratio > 0.92 and left.get("type") != "text" and right.get("type") != "text":
                    report["layout"]["heavyOverlaps"].append({"page": page_index, "left": left.get("id"), "right": right.get("id"), "ratio": round(ratio, 3)})
        page_coverage["exportedText"] = page_coverage["recognizedText"]
        report["ocrCoverage"].append(page_coverage)
    if report["layout"]["invalidNumbers"]:
        report["valid"] = False
    return report


def save_validation(report: dict[str, Any], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
